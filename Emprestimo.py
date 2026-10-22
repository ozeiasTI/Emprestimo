# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1", "pillow>=10.0", "reportlab>=4.0", "xlrd>=2.0.1"]
# ///

from __future__ import annotations

import io
import json
import os
import re
import sqlite3
import smtplib
import ssl
import subprocess
import sys
import tkinter as tk
import threading
import unicodedata
import uuid
import webbrowser
import sys
from pathlib import Path
from copy import copy
from contextlib import contextmanager
from datetime import datetime, timedelta
from email.message import EmailMessage
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from email.utils import formataddr
from xml.sax.saxutils import escape

from PIL import Image, ImageTk
import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


APP_NAME = "Emprestimo"
MAX_CATEGORY_LENGTH = 100
MAX_ITEM_NAME_LENGTH = 150
DEFAULT_ITEM_STATUS = "Disponível"
ITEM_STATUSES = ("Disponível", "Emprestado", "Em manutenção")
USER_TYPES = ("Servidor", "Externo", "Colaborador")
DEFAULT_SMTP_HOST = "smtp.gmail.com"
DEFAULT_EMAIL_SIGNATURE = "Atenciosamente,\nEquipe de Empréstimos"
PUBLIC_DIRECTORY = Path(
    os.environ.get("PUBLIC", str(Path.home().parent / "Public"))
)
SHARED_CONFIG_DIR = PUBLIC_DIRECTORY / "Documents" / APP_NAME
SHARED_CONFIG_FILE = SHARED_CONFIG_DIR / "config.json"
IS_FROZEN = bool(getattr(sys, "frozen", False))
APPLICATION_DIRECTORY = (
    Path(sys.executable).resolve().parent
    if IS_FROZEN
    else Path(__file__).resolve().parent
)
CONFIG_DIR = APPLICATION_DIRECTORY if IS_FROZEN else SHARED_CONFIG_DIR
CONFIG_FILE = CONFIG_DIR / "config.json"
LEGACY_CONFIG_FILE = Path.home() / ".emprestimo" / "config.json"
DEFAULT_DATABASE = CONFIG_DIR / "emprestimo.db"

COLORS = {
    "sidebar": "#183B3A",
    "sidebar_hover": "#28514D",
    "accent": "#D99A38",
    "background": "#F3F5F1",
    "surface": "#FFFFFF",
    "text": "#20312F",
    "muted": "#71807B",
    "border": "#DCE3DD",
    "danger": "#A6443C",
}

THEMES = {
    "Verde": {
        "sidebar": "#183B3A", "sidebar_hover": "#28514D", "accent": "#D99A38",
        "background": "#F3F5F1", "surface": "#FFFFFF", "text": "#20312F",
        "muted": "#71807B", "border": "#DCE3DD", "danger": "#A6443C",
    },
    "Azul": {
        "sidebar": "#183453", "sidebar_hover": "#28527A", "accent": "#4C91D0",
        "background": "#F1F5FA", "surface": "#FFFFFF", "text": "#1F3042",
        "muted": "#6D7D8E", "border": "#D9E2EC", "danger": "#B44747",
    },
    "Roxo": {
        "sidebar": "#35284F", "sidebar_hover": "#51406F", "accent": "#A17BD1",
        "background": "#F5F2F8", "surface": "#FFFFFF", "text": "#30283C",
        "muted": "#7A7185", "border": "#E2DCE9", "danger": "#AE4D5A",
    },
    "Grafite": {
        "sidebar": "#30343B", "sidebar_hover": "#484E58", "accent": "#E0A446",
        "background": "#F2F3F4", "surface": "#FFFFFF", "text": "#292D33",
        "muted": "#737982", "border": "#D9DCE0", "danger": "#B04444",
    },
    "Turquesa": {
        "sidebar": "#15545A", "sidebar_hover": "#24747A", "accent": "#49B8AF",
        "background": "#EFF7F6", "surface": "#FFFFFF", "text": "#1E3839",
        "muted": "#688082", "border": "#D3E5E3", "danger": "#AD4A4A",
    },
    "Vinho": {
        "sidebar": "#512D3A", "sidebar_hover": "#704153", "accent": "#D58A9B",
        "background": "#F8F2F4", "surface": "#FFFFFF", "text": "#392A30",
        "muted": "#827078", "border": "#E9DCE1", "danger": "#B0444A",
    },
    "Âmbar": {
        "sidebar": "#59401F", "sidebar_hover": "#76582C", "accent": "#D99A38",
        "background": "#F8F5EF", "surface": "#FFFFFF", "text": "#382F24",
        "muted": "#807665", "border": "#E8DFCE", "danger": "#A6443C",
    },
    "Oliva": {
        "sidebar": "#414A2D", "sidebar_hover": "#59643D", "accent": "#A8B765",
        "background": "#F4F5EF", "surface": "#FFFFFF", "text": "#303528",
        "muted": "#747B67", "border": "#DEE2D3", "danger": "#A6443C",
    },
    "Ciano": {
        "sidebar": "#164B66", "sidebar_hover": "#236782", "accent": "#46B7C7",
        "background": "#EFF6F8", "surface": "#FFFFFF", "text": "#20343D",
        "muted": "#6C7E86", "border": "#D7E4E8", "danger": "#B24B50",
    },
    "Rosa": {
        "sidebar": "#633F55", "sidebar_hover": "#80516D", "accent": "#D889AA",
        "background": "#FBF3F7", "surface": "#FFFFFF", "text": "#3C2C35",
        "muted": "#85727D", "border": "#EADDE5", "danger": "#B4484C",
    },
    "Laranja": {
        "sidebar": "#663B25", "sidebar_hover": "#85502F", "accent": "#E89A4B",
        "background": "#FBF4EE", "surface": "#FFFFFF", "text": "#3F3027",
        "muted": "#897666", "border": "#EBDED2", "danger": "#B3473F",
    },
    "Noturno": {
        "sidebar": "#151C2C", "sidebar_hover": "#28334B", "accent": "#76A9FA",
        "background": "#202838", "surface": "#2A3446", "text": "#E8EDF5",
        "muted": "#A6B1C2", "border": "#3D4A60", "danger": "#F17C78",
    },
    "Floresta": {
        "sidebar": "#263E33", "sidebar_hover": "#365747", "accent": "#91B76B",
        "background": "#F0F5F0", "surface": "#FFFFFF", "text": "#26382F",
        "muted": "#6E8074", "border": "#D8E3DA", "danger": "#AC4942",
    },
    "Lavanda": {
        "sidebar": "#40345F", "sidebar_hover": "#584879", "accent": "#B7A1E8",
        "background": "#F5F2FA", "surface": "#FFFFFF", "text": "#332C43",
        "muted": "#7B7488", "border": "#E2DCEF", "danger": "#B44E61",
    },
}

PLAIN_SECRET_PREFIX = b"plain:"

def resource_path(relative_path: str) -> Path:
    """Retorna o caminho absoluto para um recurso, funcionando no script e no .exe."""
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        base_path = Path(sys._MEIPASS)
    else:
        base_path = Path(__file__).parent
    return base_path / relative_path


def _notification_html(
    user_name: str,
    introduction: str,
    items: list[dict[str, str]],
    closing: str,
    signature: str,
) -> str:
    rows = "".join(
        "<tr>"
        f'<td style="padding:11px;border-bottom:1px solid #e2e8e4">{escape(item["number"])}</td>'
        f'<td style="padding:11px;border-bottom:1px solid #e2e8e4">{escape(item["name"])}</td>'
        f'<td style="padding:11px;border-bottom:1px solid #e2e8e4">{escape(item["category"])}</td>'
        f'<td style="padding:11px;border-bottom:1px solid #e2e8e4">{escape(item["date"])}</td>'
        "</tr>"
        for item in items
    )
    return (
        '<div style="margin:0;background:#f3f5f1;padding:28px;font-family:Arial,sans-serif;color:#20312f">'
        '<div style="max-width:720px;margin:auto;background:#fff;border-radius:12px;overflow:hidden">'
        '<div style="background:#183b3a;color:#fff;padding:22px 28px">'
        '<div style="font-size:12px;letter-spacing:1px;color:#d5e3dc">SISTEMA DE EMPRÉSTIMOS</div>'
        '<h1 style="font-size:22px;margin:8px 0 0">Movimentação de itens</h1></div>'
        '<div style="padding:26px 28px">'
        f'<p>Olá, <strong>{escape(user_name)}</strong>.</p>'
        f'<p>{escape(introduction)}</p>'
        '<table style="width:100%;border-collapse:collapse;margin:22px 0">'
        '<thead><tr style="background:#e9efeb;text-align:left">'
        '<th style="padding:11px;border-bottom:2px solid #183b3a">Nº</th>'
        '<th style="padding:11px;border-bottom:2px solid #183b3a">Item</th>'
        '<th style="padding:11px;border-bottom:2px solid #183b3a">Categoria</th>'
        '<th style="padding:11px;border-bottom:2px solid #183b3a">Data</th>'
        '</tr></thead>'
        f'<tbody>{rows}</tbody></table>'
        f'<p>{escape(closing)}</p>'
        f'<p style="margin-top:24px">{escape(signature).replace(chr(10), "<br>")}</p>'
        '</div></div></div>'
    )


def _copy_sqlite_database(source_path: Path, destination_path: Path) -> None:
    source_connection = sqlite3.connect(source_path, timeout=30)
    try:
        destination_connection = sqlite3.connect(destination_path, timeout=30)
        try:
            source_connection.backup(destination_connection)
            result = destination_connection.execute("PRAGMA integrity_check").fetchone()[0]
            if result != "ok":
                raise sqlite3.DatabaseError(f"A cópia do banco falhou na verificação: {result}")
        finally:
            destination_connection.close()
    finally:
        source_connection.close()


def _validate_import_database(path: Path) -> None:
    connection = sqlite3.connect(path, timeout=30)
    try:
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise ValueError(f"O arquivo selecionado não passou na verificação: {integrity}")
        tables = {
            str(row[0]) for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        required_tables = {"categories", "items", "users", "loans"}
        missing_tables = required_tables - tables
        if missing_tables:
            raise ValueError(
                "O arquivo não parece ser um banco do sistema. Tabelas ausentes: "
                f"{', '.join(sorted(missing_tables))}."
            )
        foreign_key_errors = connection.execute("PRAGMA foreign_key_check").fetchmany(5)
        if foreign_key_errors:
            raise ValueError(
                "O banco selecionado contém referências inválidas e não pode ser importado."
            )
    finally:
        connection.close()


class CategoryRepository:
    def __init__(self, database_path: Path):
        self.database_path = database_path.expanduser().resolve()
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as connection:
            connection.execute(
                """CREATE TABLE IF NOT EXISTS categories (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL COLLATE NOCASE UNIQUE,
                    created_at TEXT NOT NULL
                )"""
            )
            category_columns = {
                row[1] for row in connection.execute("PRAGMA table_info(categories)")
            }
            if "color" in category_columns and sqlite3.sqlite_version_info >= (3, 35, 0):
                connection.execute("ALTER TABLE categories DROP COLUMN color")
            connection.execute(
                """CREATE TABLE IF NOT EXISTS items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    number INTEGER NOT NULL CHECK (number > 0),
                    name TEXT NOT NULL,
                    category_id INTEGER NOT NULL,
                    status TEXT NOT NULL DEFAULT 'Disponível',
                    created_at TEXT NOT NULL,
                    UNIQUE (category_id, number),
                    FOREIGN KEY (category_id) REFERENCES categories (id)
                        ON DELETE RESTRICT
                )"""
            )
            item_columns = {row[1] for row in connection.execute("PRAGMA table_info(items)")}
            if "status" not in item_columns:
                connection.execute(
                    f"ALTER TABLE items ADD COLUMN status TEXT NOT NULL DEFAULT '{DEFAULT_ITEM_STATUS}'"
                )
            connection.execute(
                """CREATE TABLE IF NOT EXISTS email_settings (
                    id INTEGER PRIMARY KEY CHECK (id = 1),
                    email TEXT NOT NULL DEFAULT '',
                    app_password BLOB NOT NULL DEFAULT X'',
                    institution TEXT NOT NULL DEFAULT '',
                    sender_name TEXT NOT NULL DEFAULT '',
                    smtp_host TEXT NOT NULL DEFAULT 'smtp.gmail.com',
                    smtp_port INTEGER NOT NULL DEFAULT 587,
                    use_tls INTEGER NOT NULL DEFAULT 1,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )"""
            )
            connection.execute("INSERT OR IGNORE INTO email_settings (id) VALUES (1)")
            email_columns = {
                row[1] for row in connection.execute("PRAGMA table_info(email_settings)")
            }
            if "signature" not in email_columns:
                connection.execute(
                    "ALTER TABLE email_settings ADD COLUMN signature TEXT NOT NULL DEFAULT ''"
                )
            connection.execute(
                """CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    matricula TEXT NOT NULL COLLATE NOCASE UNIQUE,
                    email TEXT NOT NULL DEFAULT '',
                    user_type TEXT NOT NULL CHECK (user_type IN ('Servidor', 'Externo', 'Colaborador')),
                    created_at TEXT NOT NULL
                )"""
            )
            connection.execute(
                """CREATE TABLE IF NOT EXISTS loans (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    item_id INTEGER NOT NULL,
                    user_id INTEGER NOT NULL,
                    issued_at TEXT NOT NULL,
                    returned_at TEXT,
                    FOREIGN KEY (item_id) REFERENCES items (id) ON DELETE RESTRICT,
                    FOREIGN KEY (user_id) REFERENCES users (id) ON DELETE RESTRICT
                )"""
            )
            connection.execute(
                "CREATE UNIQUE INDEX IF NOT EXISTS idx_loans_active_item "
                "ON loans(item_id) WHERE returned_at IS NULL"
            )
            outbox_schema = """CREATE TABLE IF NOT EXISTS email_outbox (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    loan_id INTEGER NOT NULL,
                    user_id INTEGER NOT NULL,
                    event_type TEXT NOT NULL CHECK (event_type IN ('issue', 'return', 'reminder')),
                    recipient TEXT NOT NULL DEFAULT '',
                    subject TEXT NOT NULL,
                    body TEXT NOT NULL,
                    body_html TEXT NOT NULL DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'pending'
                        CHECK (status IN ('pending', 'sending', 'sent', 'failed')),
                    attempt_count INTEGER NOT NULL DEFAULT 0,
                    last_error TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    sent_at TEXT,
                    FOREIGN KEY (loan_id) REFERENCES loans (id) ON DELETE CASCADE,
                    FOREIGN KEY (user_id) REFERENCES users (id) ON DELETE RESTRICT
                )"""
            connection.execute(outbox_schema)
            outbox_sql = connection.execute(
                "SELECT sql FROM sqlite_master WHERE type = 'table' AND name = 'email_outbox'"
            ).fetchone()[0].lower()
            if "'reminder'" not in outbox_sql or "unique (loan_id, event_type)" in outbox_sql:
                connection.execute("ALTER TABLE email_outbox RENAME TO email_outbox_legacy")
                connection.execute(outbox_schema)
                connection.execute(
                    """INSERT INTO email_outbox (
                        id, loan_id, user_id, event_type, recipient, subject, body,
                        status, attempt_count, last_error, created_at, sent_at
                    ) SELECT id, loan_id, user_id, event_type, recipient, subject, body,
                        status, attempt_count, last_error, created_at, sent_at
                    FROM email_outbox_legacy"""
                )
                connection.execute("DROP TABLE email_outbox_legacy")
            connection.execute(
                "CREATE UNIQUE INDEX IF NOT EXISTS idx_email_outbox_single_action "
                "ON email_outbox(loan_id, event_type) WHERE event_type IN ('issue', 'return')"
            )
            outbox_columns = {
                row[1] for row in connection.execute("PRAGMA table_info(email_outbox)")
            }
            if "body_html" not in outbox_columns:
                connection.execute(
                    "ALTER TABLE email_outbox ADD COLUMN body_html TEXT NOT NULL DEFAULT ''"
                )
            connection.execute(
                """CREATE TABLE IF NOT EXISTS email_outbox_loans (
                    outbox_id INTEGER NOT NULL REFERENCES email_outbox(id) ON DELETE CASCADE,
                    loan_id INTEGER NOT NULL REFERENCES loans(id) ON DELETE CASCADE,
                    PRIMARY KEY (outbox_id, loan_id)
                )"""
            )
            connection.execute(
                """INSERT OR IGNORE INTO email_outbox_loans (outbox_id, loan_id)
                SELECT id, loan_id FROM email_outbox"""
            )
            connection.execute(
                "UPDATE email_outbox SET status = 'pending' WHERE status = 'sending'"
            )

    def create_backup(self, destination_path: Path) -> Path:
        destination = destination_path.expanduser().resolve()
        if destination == self.database_path:
            raise ValueError("O arquivo de backup não pode substituir o banco em uso.")
        destination.parent.mkdir(parents=True, exist_ok=True)
        _copy_sqlite_database(self.database_path, destination)
        return destination

    def restore_backup(self, backup_path: Path) -> None:
        source = backup_path.expanduser().resolve()
        if source == self.database_path:
            raise ValueError("Selecione um arquivo de backup diferente do banco em uso.")
        if not source.is_file():
            raise FileNotFoundError(f"O arquivo selecionado não existe: {source}")
        _validate_import_database(source)

        staging = self.database_path.with_name(
            f".{self.database_path.stem}.{uuid.uuid4().hex}.restore"
        )
        try:
            _copy_sqlite_database(source, staging)
            CategoryRepository(staging)
            _validate_import_database(staging)
            os.replace(staging, self.database_path)
        finally:
            for temporary_path in (
                staging, Path(f"{staging}-journal"), Path(f"{staging}-wal"),
                Path(f"{staging}-shm"),
            ):
                temporary_path.unlink(missing_ok=True)

    @contextmanager
    def _connect(self):
        connection = sqlite3.connect(self.database_path, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA busy_timeout = 30000")
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def list_categories(self, search: str = "") -> list[sqlite3.Row]:
        with self._connect() as connection:
            return connection.execute(
                "SELECT id, name, created_at FROM categories "
                "WHERE name LIKE ? ORDER BY name COLLATE NOCASE",
                (f"%{search.strip()}%",),
            ).fetchall()

    def list_category_item_counts(self) -> list[sqlite3.Row]:
        with self._connect() as connection:
            return connection.execute(
                "SELECT categories.id, categories.name, COUNT(items.id) AS item_count "
                "FROM categories LEFT JOIN items ON items.category_id = categories.id "
                "GROUP BY categories.id ORDER BY categories.name COLLATE NOCASE"
            ).fetchall()

    def list_recent_items(
        self,
        limit: int = 30,
        category_id: int | None = None,
        search: str = "",
    ) -> list[sqlite3.Row]:
        filters = []
        parameters: list[object] = []
        if category_id is not None:
            filters.append("items.category_id = ?")
            parameters.append(category_id)
        if search.strip():
            filters.append(
                "(items.name LIKE ? OR CAST(items.number AS TEXT) LIKE ?)"
            )
            search_pattern = f"%{search.strip()}%"
            parameters.extend((search_pattern, search_pattern))
        where_clause = f"WHERE {' AND '.join(filters)}" if filters else ""
        parameters.append(limit)
        with self._connect() as connection:
            return connection.execute(
                "SELECT items.id, items.number, items.name, items.category_id, "
                "categories.name AS category_name, items.status, items.created_at "
                "FROM items JOIN categories ON categories.id = items.category_id "
                f"{where_clause} ORDER BY items.created_at DESC, items.id DESC LIMIT ?",
                parameters,
            ).fetchall()

    def count_categories(self) -> int:
        with self._connect() as connection:
            return int(connection.execute("SELECT COUNT(*) FROM categories").fetchone()[0])

    def create_category(self, name: str) -> None:
        normalized = self._validate_name(name)
        with self._connect() as connection:
            connection.execute(
                "INSERT INTO categories (name, created_at) VALUES (?, ?)",
                (normalized, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            )

    def update_category(self, category_id: int, name: str) -> None:
        normalized = self._validate_name(name)
        with self._connect() as connection:
            cursor = connection.execute(
                "UPDATE categories SET name = ? WHERE id = ?",
                (normalized, category_id),
            )
            if cursor.rowcount == 0:
                raise ValueError("A categoria selecionada não existe mais.")

    def delete_category(self, category_id: int) -> None:
        with self._connect() as connection:
            connection.execute("DELETE FROM categories WHERE id = ?", (category_id,))

    def list_items(
        self, search: str = "", category_id: int | None = None
    ) -> list[sqlite3.Row]:
        filters = []
        parameters: list[object] = []
        if search.strip():
            pattern = f"%{search.strip()}%"
            filters.append(
                "(CAST(items.number AS TEXT) LIKE ? OR items.name LIKE ? "
                "OR categories.name LIKE ?)"
            )
            parameters.extend((pattern, pattern, pattern))
        if category_id is not None:
            filters.append("items.category_id = ?")
            parameters.append(category_id)
        where_clause = f"WHERE {' AND '.join(filters)}" if filters else ""
        with self._connect() as connection:
            return connection.execute(
                "SELECT items.id, items.number, items.name, items.category_id, "
                "categories.name AS category_name, "
                "items.status, items.created_at "
                "FROM items JOIN categories ON categories.id = items.category_id "
                f"{where_clause} ORDER BY categories.name COLLATE NOCASE, items.number, items.name COLLATE NOCASE",
                parameters,
            ).fetchall()

    def count_items(self, status: str | None = None) -> int:
        with self._connect() as connection:
            if status is None:
                query = "SELECT COUNT(*) FROM items"
                parameters = ()
            else:
                query = "SELECT COUNT(*) FROM items WHERE status = ?"
                parameters = (status,)
            return int(connection.execute(query, parameters).fetchone()[0])

    def count_active_loans(self) -> int:
        with self._connect() as connection:
            return int(connection.execute(
                "SELECT COUNT(*) FROM loans WHERE returned_at IS NULL"
            ).fetchone()[0])

    def count_outbox_messages(self, status: str) -> int:
        with self._connect() as connection:
            return int(connection.execute(
                "SELECT COUNT(*) FROM email_outbox WHERE status = ?", (status,)
            ).fetchone()[0])

    def get_item(self, item_id: int) -> sqlite3.Row | None:
        with self._connect() as connection:
            return connection.execute(
                "SELECT items.id, items.number, items.name, items.category_id, "
                "categories.name AS category_name, "
                "items.status, items.created_at "
                "FROM items JOIN categories ON categories.id = items.category_id "
                "WHERE items.id = ?",
                (item_id,),
            ).fetchone()

    def create_item(
        self, number: str | int, name: str, category_id: int,
        status: str = DEFAULT_ITEM_STATUS,
    ) -> None:
        normalized_number = self._validate_item_number(number)
        normalized_name = self._validate_item_name(name)
        normalized_status = self._validate_item_status(status)
        with self._connect() as connection:
            connection.execute(
                "INSERT INTO items (number, name, category_id, status, created_at) "
                "VALUES (?, ?, ?, ?, ?)",
                (
                    normalized_number,
                    normalized_name,
                    category_id,
                    normalized_status,
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                ),
            )

    def update_item(
        self, item_id: int, number: str | int, name: str, category_id: int,
        status: str | None = None,
    ) -> None:
        normalized_number = self._validate_item_number(number)
        normalized_name = self._validate_item_name(name)
        with self._connect() as connection:
            if status is None:
                cursor = connection.execute(
                    "UPDATE items SET number = ?, name = ?, category_id = ? WHERE id = ?",
                    (normalized_number, normalized_name, category_id, item_id),
                )
            else:
                normalized_status = self._validate_item_status(status)
                cursor = connection.execute(
                    "UPDATE items SET number = ?, name = ?, category_id = ?, status = ? WHERE id = ?",
                    (normalized_number, normalized_name, category_id, normalized_status, item_id),
                )
            if cursor.rowcount == 0:
                raise ValueError("O item selecionado não existe mais.")

    def delete_item(self, item_id: int) -> None:
        with self._connect() as connection:
            connection.execute("DELETE FROM items WHERE id = ?", (item_id,))

    def list_available_items(self, category_id: int | None = None) -> list[sqlite3.Row]:
        query = (
            "SELECT items.id, items.number, items.name, items.category_id, "
            "categories.name AS category_name "
            "FROM items JOIN categories ON categories.id = items.category_id "
            "WHERE items.status = ?"
        )
        parameters: list[object] = [DEFAULT_ITEM_STATUS]
        if category_id is not None:
            query += " AND items.category_id = ?"
            parameters.append(category_id)
        query += " ORDER BY categories.name COLLATE NOCASE, items.number, items.name COLLATE NOCASE"
        with self._connect() as connection:
            return connection.execute(query, parameters).fetchall()

    def list_loans(
        self, active_only: bool = False, user_id: int | None = None,
    ) -> list[sqlite3.Row]:
        filters = []
        parameters: list[int] = []
        if active_only:
            filters.append("loans.returned_at IS NULL")
        if user_id is not None:
            filters.append("loans.user_id = ?")
            parameters.append(user_id)
        where_clause = f"WHERE {' AND '.join(filters)}" if filters else ""
        with self._connect() as connection:
            return connection.execute(
                "SELECT loans.id, loans.item_id, loans.user_id, loans.issued_at, loans.returned_at, "
                "items.number, items.name AS item_name, categories.name AS category_name, "
                "users.name AS user_name, users.matricula, users.email, "
                "COALESCE(issue_mail.status, 'pending') AS issue_email_status, "
                "COALESCE(return_mail.status, '') AS return_email_status, "
                "COALESCE(reminder_mail.status, '') AS reminder_email_status "
                "FROM loans JOIN items ON items.id = loans.item_id "
                "JOIN categories ON categories.id = items.category_id "
                "JOIN users ON users.id = loans.user_id "
                "LEFT JOIN email_outbox AS issue_mail "
                "ON issue_mail.id = (SELECT MAX(mail.id) FROM email_outbox AS mail "
                "JOIN email_outbox_loans AS link ON link.outbox_id = mail.id "
                "WHERE link.loan_id = loans.id AND mail.event_type = 'issue') "
                "LEFT JOIN email_outbox AS return_mail "
                "ON return_mail.id = (SELECT MAX(mail.id) FROM email_outbox AS mail "
                "JOIN email_outbox_loans AS link ON link.outbox_id = mail.id "
                "WHERE link.loan_id = loans.id AND mail.event_type = 'return') "
                "LEFT JOIN email_outbox AS reminder_mail "
                "ON reminder_mail.id = (SELECT MAX(reminder.id) FROM email_outbox AS reminder "
                "JOIN email_outbox_loans AS link ON link.outbox_id = reminder.id "
                "WHERE link.loan_id = loans.id AND reminder.event_type = 'reminder') "
                f"{where_clause} ORDER BY loans.issued_at DESC, loans.id DESC",
                parameters,
            ).fetchall()

    def get_active_loan_for_item(self, item_id: int) -> sqlite3.Row | None:
        with self._connect() as connection:
            return connection.execute(
                "SELECT loans.id, loans.issued_at, items.number, items.name AS item_name, "
                "categories.name AS category_name, users.name AS user_name, "
                "users.matricula, users.email, users.user_type "
                "FROM loans JOIN items ON items.id = loans.item_id "
                "JOIN categories ON categories.id = items.category_id "
                "JOIN users ON users.id = loans.user_id "
                "WHERE loans.item_id = ? AND loans.returned_at IS NULL",
                (item_id,),
            ).fetchone()

    def create_loan(self, item_id: int, user_id: int) -> tuple[int, int]:
        loan_ids, outbox_id = self.create_loans([item_id], user_id)
        return loan_ids[0], outbox_id

    def create_loans(self, item_ids: list[int], user_id: int) -> tuple[list[int], int]:
        unique_item_ids = list(dict.fromkeys(item_ids))
        if not unique_item_ids:
            raise ValueError("Selecione pelo menos um item disponível.")
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as connection:
            user = connection.execute(
                "SELECT id, name, matricula, email FROM users WHERE id = ?", (user_id,),
            ).fetchone()
            if user is None:
                raise ValueError("O usuário selecionado não existe mais.")
            loan_ids: list[int] = []
            notification_items: list[dict[str, str]] = []
            for item_id in unique_item_ids:
                item = connection.execute(
                    "SELECT items.id, items.number, items.name, items.status, "
                    "categories.name AS category_name FROM items "
                    "JOIN categories ON categories.id = items.category_id WHERE items.id = ?",
                    (item_id,),
                ).fetchone()
                if item is None:
                    raise ValueError("Um dos itens selecionados não existe mais.")
                if item["status"] != DEFAULT_ITEM_STATUS:
                    raise ValueError(f"O item Nº {item['number']} não está mais disponível.")
                changed = connection.execute(
                    "UPDATE items SET status = 'Emprestado' WHERE id = ? AND status = ?",
                    (item_id, DEFAULT_ITEM_STATUS),
                ).rowcount
                if changed != 1:
                    raise ValueError(f"O item Nº {item['number']} acabou de deixar de estar disponível.")
                loan_ids.append(int(connection.execute(
                    "INSERT INTO loans (item_id, user_id, issued_at) VALUES (?, ?, ?)",
                    (item_id, user_id, timestamp),
                ).lastrowid))
                notification_items.append({
                    "number": str(item["number"]), "name": str(item["name"]),
                    "category": str(item["category_name"]), "date": timestamp,
                })
            email_settings = connection.execute(
                "SELECT institution, signature FROM email_settings WHERE id = 1"
            ).fetchone()
            institution = email_settings["institution"].strip() or "Sistema de Empréstimos"
            signature = email_settings["signature"].strip() or DEFAULT_EMAIL_SIGNATURE
            subject = f"Retirada de {len(loan_ids)} item(ns) - {institution}"
            body = (
                f"Olá, {user['name']}.\n\n"
                f"A retirada dos itens abaixo foi registrada em {institution}.\n\n"
                + "\n".join(
                    f"- Nº {item['number']} · {item['name']} · {item['category']} · {timestamp}"
                    for item in notification_items
                )
                + "\n\n"
                f"Matrícula: {user['matricula']}\n"
                f"Guarde esta mensagem para consulta.\n\n{signature}"
            )
            outbox_id = int(connection.execute(
                """INSERT INTO email_outbox (
                    loan_id, user_id, event_type, recipient, subject, body, body_html, created_at
                ) VALUES (?, ?, 'issue', ?, ?, ?, ?, ?)""",
                (
                    loan_ids[0], user_id, user["email"], subject, body,
                    _notification_html(
                        str(user["name"]), f"A retirada foi registrada em {institution}.",
                        notification_items, "Guarde esta mensagem para consulta.", signature,
                    ),
                    timestamp,
                ),
            ).lastrowid)
            connection.executemany(
                "INSERT INTO email_outbox_loans (outbox_id, loan_id) VALUES (?, ?)",
                [(outbox_id, loan_id) for loan_id in loan_ids],
            )
        return loan_ids, outbox_id

    def return_loan(self, loan_id: int) -> int:
        return self.return_loans([loan_id])[1]

    def return_loans(self, loan_ids: list[int]) -> tuple[list[int], int]:
        unique_loan_ids = list(dict.fromkeys(loan_ids))
        if not unique_loan_ids:
            raise ValueError("Selecione pelo menos um empréstimo em aberto.")
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as connection:
            placeholders = ", ".join("?" for _ in unique_loan_ids)
            loans = connection.execute(
                "SELECT loans.id, loans.item_id, loans.user_id, loans.returned_at, "
                "items.number, items.name AS item_name, categories.name AS category_name, "
                "users.name AS user_name, users.matricula, users.email "
                "FROM loans JOIN items ON items.id = loans.item_id "
                "JOIN categories ON categories.id = items.category_id "
                "JOIN users ON users.id = loans.user_id "
                f"WHERE loans.id IN ({placeholders}) ORDER BY loans.id",
                unique_loan_ids,
            ).fetchall()
            if len(loans) != len(unique_loan_ids):
                raise ValueError("Um dos empréstimos selecionados não existe mais.")
            if any(loan["returned_at"] is not None for loan in loans):
                raise ValueError("Um dos itens selecionados já foi devolvido.")
            user_ids = {loan["user_id"] for loan in loans}
            if len(user_ids) != 1:
                raise ValueError("Devolva itens de um único usuário por vez.")
            user_id = next(iter(user_ids))
            for loan in loans:
                changed = connection.execute(
                    "UPDATE loans SET returned_at = ? WHERE id = ? AND returned_at IS NULL",
                    (timestamp, loan["id"]),
                ).rowcount
                if changed != 1:
                    raise ValueError(f"O empréstimo do item Nº {loan['number']} já foi alterado.")
                connection.execute(
                    "UPDATE items SET status = ? WHERE id = ?",
                    (DEFAULT_ITEM_STATUS, loan["item_id"]),
                )
            email_settings = connection.execute(
                "SELECT institution, signature FROM email_settings WHERE id = 1"
            ).fetchone()
            institution = email_settings["institution"].strip() or "Sistema de Empréstimos"
            signature = email_settings["signature"].strip() or DEFAULT_EMAIL_SIGNATURE
            user_name = str(loans[0]["user_name"])
            subject = f"Devolução de {len(loans)} item(ns) - {institution}"
            notification_items = [
                {
                    "number": str(loan["number"]), "name": str(loan["item_name"]),
                    "category": str(loan["category_name"]), "date": timestamp,
                }
                for loan in loans
            ]
            body = (
                f"Olá, {user_name}.\n\n"
                f"A devolução dos itens abaixo foi registrada em {institution}.\n\n"
                + "\n".join(
                    f"- Nº {item['number']} · {item['name']} · {item['category']} · {timestamp}"
                    for item in notification_items
                )
                + f"\n\nMatrícula: {loans[0]['matricula']}\n\nObrigado.\n\n{signature}"
            )
            outbox_id = int(connection.execute(
                """INSERT INTO email_outbox (
                    loan_id, user_id, event_type, recipient, subject, body, body_html, created_at
                ) VALUES (?, ?, 'return', ?, ?, ?, ?, ?)""",
                (
                    unique_loan_ids[0], user_id, loans[0]["email"], subject, body,
                    _notification_html(
                        user_name, f"A devolução foi registrada em {institution}.",
                        notification_items, "Obrigado por devolver os itens.", signature,
                    ),
                    timestamp,
                ),
            ).lastrowid)
            connection.executemany(
                "INSERT INTO email_outbox_loans (outbox_id, loan_id) VALUES (?, ?)",
                [(outbox_id, loan_id) for loan_id in unique_loan_ids],
            )
        return unique_loan_ids, outbox_id

    def create_loan_reminder(self, loan_id: int) -> int:
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as connection:
            selected_loan = connection.execute(
                "SELECT loans.id, loans.user_id, loans.issued_at, loans.returned_at, "
                "users.name AS user_name, users.matricula, users.email "
                "FROM loans JOIN items ON items.id = loans.item_id "
                "JOIN categories ON categories.id = items.category_id "
                "JOIN users ON users.id = loans.user_id WHERE loans.id = ?",
                (loan_id,),
            ).fetchone()
            if selected_loan is None:
                raise ValueError("O empréstimo selecionado não existe.")
            if selected_loan["returned_at"] is not None:
                raise ValueError("Só é possível cobrar itens que ainda estão emprestados.")
            loans = connection.execute(
                "SELECT loans.id, loans.issued_at, items.number, items.name AS item_name, "
                "categories.name AS category_name FROM loans "
                "JOIN items ON items.id = loans.item_id "
                "JOIN categories ON categories.id = items.category_id "
                "WHERE loans.user_id = ? AND loans.returned_at IS NULL "
                "ORDER BY loans.issued_at, loans.id",
                (selected_loan["user_id"],),
            ).fetchall()
            email_settings = connection.execute(
                "SELECT institution, signature FROM email_settings WHERE id = 1"
            ).fetchone()
            institution = email_settings["institution"].strip() or "Sistema de Empréstimos"
            signature = email_settings["signature"].strip() or DEFAULT_EMAIL_SIGNATURE
            notification_items = [
                {
                    "number": str(loan["number"]), "name": str(loan["item_name"]),
                    "category": str(loan["category_name"]), "date": str(loan["issued_at"]),
                }
                for loan in loans
            ]
            subject = f"Lembrete: {len(loans)} item(ns) emprestado(s) - {institution}"
            body = (
                f"Olá, {selected_loan['user_name']}.\n\n"
                f"Os itens abaixo ainda constam emprestados em {institution}:\n\n"
                + "\n".join(
                    f"- Nº {item['number']} · {item['name']} · {item['category']} · retirado em {item['date']}"
                    for item in notification_items
                )
                + f"\n\nMatrícula: {selected_loan['matricula']}\n\n"
                f"Por favor, providencie a devolução dos itens.\n\n{signature}"
            )
            outbox_id = int(connection.execute(
                """INSERT INTO email_outbox (
                    loan_id, user_id, event_type, recipient, subject, body, body_html, created_at
                ) VALUES (?, ?, 'reminder', ?, ?, ?, ?, ?)""",
                (
                    loans[0]["id"], selected_loan["user_id"], selected_loan["email"],
                    subject, body,
                    _notification_html(
                        str(selected_loan["user_name"]),
                        f"Segue a lista de itens ainda emprestados em {institution}.",
                        notification_items, "Por favor, providencie a devolução.", signature,
                    ),
                    timestamp,
                ),
            ).lastrowid)
            connection.executemany(
                "INSERT INTO email_outbox_loans (outbox_id, loan_id) VALUES (?, ?)",
                [(outbox_id, loan["id"]) for loan in loans],
            )
        return outbox_id

    def get_outbox_message(self, outbox_id: int) -> sqlite3.Row | None:
        with self._connect() as connection:
            return connection.execute(
                "SELECT email_outbox.*, "
                "COALESCE(NULLIF(email_outbox.recipient, ''), users.email) AS resolved_recipient "
                "FROM email_outbox JOIN users ON users.id = email_outbox.user_id "
                "WHERE email_outbox.id = ?",
                (outbox_id,),
            ).fetchone()

    def list_retryable_outbox_ids(self, loan_id: int | None = None) -> list[int]:
        with self._connect() as connection:
            if loan_id is None:
                query = (
                    "SELECT id FROM email_outbox WHERE status IN ('pending', 'failed') "
                    "ORDER BY id"
                )
                parameters: tuple[object, ...] = ()
            else:
                query = (
                    "SELECT DISTINCT email_outbox.id FROM email_outbox "
                    "JOIN email_outbox_loans ON email_outbox_loans.outbox_id = email_outbox.id "
                    "WHERE email_outbox_loans.loan_id = ? "
                    "AND email_outbox.status IN ('pending', 'failed') ORDER BY email_outbox.id"
                )
                parameters = (loan_id,)
            return [row[0] for row in connection.execute(
                query, parameters,
            ).fetchall()]

    def list_outbox_messages(self) -> list[sqlite3.Row]:
        with self._connect() as connection:
            return connection.execute(
                "SELECT email_outbox.id, email_outbox.event_type, email_outbox.subject, "
                "email_outbox.recipient, email_outbox.status, email_outbox.created_at, "
                "email_outbox.sent_at, email_outbox.last_error, users.name AS user_name, "
                "email_outbox.attempt_count, "
                "(SELECT GROUP_CONCAT(items.name, ', ') FROM email_outbox_loans AS links "
                "JOIN loans ON loans.id = links.loan_id "
                "JOIN items ON items.id = loans.item_id "
                "WHERE links.outbox_id = email_outbox.id) AS item_names "
                "FROM email_outbox JOIN users ON users.id = email_outbox.user_id "
                "ORDER BY email_outbox.id DESC"
            ).fetchall()

    def claim_outbox_message(self, outbox_id: int) -> bool:
        with self._connect() as connection:
            changed = connection.execute(
                "UPDATE email_outbox SET status = 'sending', attempt_count = attempt_count + 1 "
                "WHERE id = ? AND status IN ('pending', 'failed')",
                (outbox_id,),
            ).rowcount
            return changed == 1

    def mark_outbox_sent(self, outbox_id: int, recipient: str) -> None:
        with self._connect() as connection:
            connection.execute(
                "UPDATE email_outbox SET status = 'sent', recipient = ?, "
                "last_error = '', sent_at = ? WHERE id = ?",
                (recipient, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), outbox_id),
            )

    def mark_outbox_failed(self, outbox_id: int, error: str) -> None:
        with self._connect() as connection:
            connection.execute(
                "UPDATE email_outbox SET status = 'failed', last_error = ? WHERE id = ?",
                (error[:1000], outbox_id),
            )

    @staticmethod
    def _validate_item_status(status: str) -> str:
        normalized = status.strip()
        if normalized not in ITEM_STATUSES:
            raise ValueError("Selecione um status válido para o item.")
        return normalized

    def get_email_settings(self) -> dict[str, object]:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM email_settings WHERE id = 1").fetchone()
        settings = dict(row)
        stored_password = bytes(settings["app_password"] or b"")
        if stored_password.startswith(PLAIN_SECRET_PREFIX):
            settings["app_password"] = stored_password[len(PLAIN_SECRET_PREFIX):].decode("utf-8")
        elif stored_password:
            raise RuntimeError(
                "A senha salva em formato antigo precisa ser informada novamente nas configurações."
            )
        else:
            settings["app_password"] = ""
        settings["use_tls"] = bool(settings["use_tls"])
        settings["smtp_host"] = settings["smtp_host"] or DEFAULT_SMTP_HOST
        settings["signature"] = settings.get("signature") or DEFAULT_EMAIL_SIGNATURE
        return settings

    def save_email_settings(
        self, email: str, app_password: str, institution: str, sender_name: str,
        smtp_host: str, smtp_port: int, use_tls: bool, signature: str | None = None,
    ) -> None:
        normalized_email = email.strip()
        if normalized_email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", normalized_email):
            raise ValueError("Informe um endereço de e-mail válido.")
        if not 1 <= smtp_port <= 65535:
            raise ValueError("A porta SMTP deve estar entre 1 e 65535.")
        shared_password = PLAIN_SECRET_PREFIX + app_password.strip().encode("utf-8")
        normalized_signature = (
            DEFAULT_EMAIL_SIGNATURE if signature is None else signature.strip()
        )
        with self._connect() as connection:
            connection.execute(
                """INSERT INTO email_settings (
                    id, email, app_password, institution, sender_name, smtp_host,
                    smtp_port, use_tls, signature, updated_at
                ) VALUES (1, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(id) DO UPDATE SET
                    email = excluded.email,
                    app_password = excluded.app_password,
                    institution = excluded.institution,
                    sender_name = excluded.sender_name,
                    smtp_host = excluded.smtp_host,
                    smtp_port = excluded.smtp_port,
                    use_tls = excluded.use_tls,
                    signature = excluded.signature,
                    updated_at = excluded.updated_at""",
                (
                    normalized_email, shared_password, institution.strip(), sender_name.strip(),
                    smtp_host.strip() or DEFAULT_SMTP_HOST, smtp_port, int(use_tls),
                    normalized_signature,
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                ),
            )

    def list_users(
        self, search: str = "", user_type: str | None = None,
    ) -> list[sqlite3.Row]:
        filters = []
        parameters: list[str] = []
        if search.strip():
            pattern = f"%{search.strip()}%"
            filters.append("(name LIKE ? OR matricula LIKE ? OR email LIKE ? OR user_type LIKE ?)")
            parameters.extend((pattern, pattern, pattern, pattern))
        if user_type is not None:
            if user_type not in USER_TYPES:
                raise ValueError("Selecione um tipo de usuário válido.")
            filters.append("user_type = ?")
            parameters.append(user_type)
        where_clause = f"WHERE {' AND '.join(filters)}" if filters else ""
        with self._connect() as connection:
            return connection.execute(
                "SELECT id, name, matricula, email, user_type, created_at FROM users "
                f"{where_clause} ORDER BY name COLLATE NOCASE",
                parameters,
            ).fetchall()

    def count_users(self, user_type: str | None = None) -> int:
        with self._connect() as connection:
            if user_type is None:
                row = connection.execute("SELECT COUNT(*) FROM users").fetchone()
            else:
                row = connection.execute(
                    "SELECT COUNT(*) FROM users WHERE user_type = ?", (user_type,),
                ).fetchone()
            return int(row[0])

    def create_user(self, name: str, matricula: object, email: str, user_type: str) -> None:
        record = self._validate_user(name, matricula, email, user_type)
        with self._connect() as connection:
            connection.execute(
                "INSERT INTO users (name, matricula, email, user_type, created_at) "
                "VALUES (?, ?, ?, ?, ?)",
                (*record, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            )

    def get_user(self, user_id: int) -> sqlite3.Row | None:
        with self._connect() as connection:
            return connection.execute(
                "SELECT id, name, matricula, email, user_type, created_at "
                "FROM users WHERE id = ?",
                (user_id,),
            ).fetchone()

    def update_user(
        self, user_id: int, name: str, matricula: object, email: str, user_type: str,
    ) -> None:
        record = self._validate_user(name, matricula, email, user_type)
        with self._connect() as connection:
            cursor = connection.execute(
                "UPDATE users SET name = ?, matricula = ?, email = ?, user_type = ? "
                "WHERE id = ?",
                (*record, user_id),
            )
            if cursor.rowcount == 0:
                raise ValueError("O usuário selecionado não existe mais.")

    def sync_servers(
        self, records: list[tuple[object, object, object]],
    ) -> dict[str, int]:
        normalized = [
            self._validate_user(name, matricula, email, "Servidor")
            for matricula, name, email in records
        ]
        if not normalized:
            raise ValueError("A planilha não contém servidores válidos; nenhum dado foi alterado.")
        registrations = [record[1].casefold() for record in normalized]
        if len(registrations) != len(set(registrations)):
            raise ValueError("A planilha contém matrículas duplicadas; nenhum dado foi alterado.")
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as connection:
            conflicts = connection.execute(
                "SELECT matricula FROM users WHERE user_type != 'Servidor'"
            ).fetchall()
            conflicting_registrations = {
                str(row["matricula"]).casefold() for row in conflicts
            }.intersection(registrations)
            if conflicting_registrations:
                raise ValueError(
                    "Uma matrícula da planilha já pertence a um usuário Externo ou "
                    "Colaborador. Nenhum dado foi alterado."
                )
            existing = connection.execute(
                "SELECT id, name, matricula, email FROM users "
                "WHERE user_type = 'Servidor'"
            ).fetchall()
            unmatched_ids = {int(user["id"]) for user in existing}
            existing_by_id = {int(user["id"]): user for user in existing}
            import_name_counts: dict[str, int] = {}
            import_email_counts: dict[str, int] = {}
            for record in normalized:
                key = record[0].casefold()
                import_name_counts[key] = import_name_counts.get(key, 0) + 1
            for record in normalized:
                key = record[2].casefold()
                if key:
                    import_email_counts[key] = import_email_counts.get(key, 0) + 1
            assignments: dict[int, tuple[str, str, str, str]] = {}

            for record in normalized:
                registration_key = record[1].casefold()
                candidates = [
                    user_id for user_id in unmatched_ids
                    if str(existing_by_id[user_id]["matricula"]).casefold() == registration_key
                ]
                if len(candidates) == 1:
                    assignments[candidates[0]] = record
                    unmatched_ids.remove(candidates[0])
            for record in normalized:
                if record in assignments.values():
                    continue
                email_key = record[2].casefold()
                candidates = [
                    user_id for user_id in unmatched_ids
                    if email_key and str(existing_by_id[user_id]["email"]).casefold() == email_key
                ]
                if import_email_counts.get(email_key) == 1 and len(candidates) == 1:
                    assignments[candidates[0]] = record
                    unmatched_ids.remove(candidates[0])
            for record in normalized:
                if record in assignments.values():
                    continue
                name_key = record[0].casefold()
                candidates = [
                    user_id for user_id in unmatched_ids
                    if str(existing_by_id[user_id]["name"]).casefold() == name_key
                ]
                if import_name_counts[name_key] == 1 and len(candidates) == 1:
                    assignments[candidates[0]] = record
                    unmatched_ids.remove(candidates[0])

            for user_id, record in assignments.items():
                if str(existing_by_id[user_id]["matricula"]).casefold() != record[1].casefold():
                    temporary_registration = f"__sync_user_{uuid.uuid4().hex}"
                    connection.execute(
                        "UPDATE users SET matricula = ? WHERE id = ?",
                        (temporary_registration, user_id),
                    )
            for user_id, record in assignments.items():
                connection.execute(
                    "UPDATE users SET name = ?, matricula = ?, email = ?, "
                    "user_type = 'Servidor' WHERE id = ?",
                    (*record[:3], user_id),
                )

            assigned_records = list(assignments.values())
            new_records = [record for record in normalized if record not in assigned_records]
            connection.executemany(
                "INSERT INTO users (name, matricula, email, user_type, created_at) "
                "VALUES (?, ?, ?, ?, ?)",
                [(*record, timestamp) for record in new_records],
            )

            retained = 0
            removed = 0
            for user_id in unmatched_ids:
                has_history = connection.execute(
                    "SELECT EXISTS(SELECT 1 FROM loans WHERE user_id = ?) "
                    "OR EXISTS(SELECT 1 FROM email_outbox WHERE user_id = ?)",
                    (user_id, user_id),
                ).fetchone()[0]
                if has_history:
                    retained += 1
                else:
                    connection.execute("DELETE FROM users WHERE id = ?", (user_id,))
                    removed += 1
        return {
            "processed": len(normalized),
            "updated": len(assignments),
            "added": len(new_records),
            "retained_with_history": retained,
            "removed": removed,
        }

    def replace_servers(self, records: list[tuple[object, object, object]]) -> int:
        return self.sync_servers(records)["processed"]

    def delete_user(self, user_id: int) -> None:
        with self._connect() as connection:
            connection.execute("DELETE FROM users WHERE id = ?", (user_id,))

    @staticmethod
    def _validate_user(
        name: object, matricula: object, email: object, user_type: str,
    ) -> tuple[str, str, str, str]:
        normalized_name = str(name or "").strip()
        if isinstance(matricula, float) and matricula.is_integer():
            normalized_matricula = str(int(matricula))
        else:
            normalized_matricula = str(matricula or "").strip()
        normalized_email = str(email or "").strip()
        if not normalized_name or not normalized_matricula:
            raise ValueError("Nome e matrícula são obrigatórios.")
        if normalized_email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", normalized_email):
            raise ValueError(f"E-mail inválido para a matrícula {normalized_matricula}.")
        if user_type not in USER_TYPES:
            raise ValueError("Selecione um tipo de usuário válido.")
        return normalized_name, normalized_matricula, normalized_email, user_type

    @staticmethod
    def _validate_name(name: str) -> str:
        normalized = name.strip()
        if not normalized:
            raise ValueError("Informe o nome da categoria.")
        if len(normalized) > MAX_CATEGORY_LENGTH:
            raise ValueError(f"O nome deve ter até {MAX_CATEGORY_LENGTH} caracteres.")
        return normalized

    @staticmethod
    def _validate_item_number(number: str | int) -> int:
        try:
            normalized = int(str(number).strip())
        except ValueError as error:
            raise ValueError("O número do item deve ser um inteiro.") from error
        if normalized < 1:
            raise ValueError("O número do item deve ser maior que zero.")
        return normalized

    @staticmethod
    def _validate_item_name(name: str) -> str:
        normalized = name.strip()
        if not normalized:
            raise ValueError("Informe o nome do item.")
        if len(normalized) > MAX_ITEM_NAME_LENGTH:
            raise ValueError(f"O nome deve ter até {MAX_ITEM_NAME_LENGTH} caracteres.")
        return normalized


def read_server_spreadsheet(path: Path) -> list[tuple[str, str, str]]:
    path = Path(path)
    if path.suffix.lower() == ".xls":
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        sheet = workbook.sheet_by_index(0)
        rows = [
            [sheet.cell_value(row_index, column_index) if sheet.ncols > column_index else ""
             for column_index in (1, 2, 3)]
            for row_index in range(sheet.nrows)
        ]
        workbook.release_resources()
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(min_col=2, max_col=4, values_only=True)]
        workbook.close()
    else:
        raise ValueError("Selecione uma planilha .xls ou .xlsx.")

    def header_text(value: object) -> str:
        normalized = unicodedata.normalize("NFKD", str(value or "").lower())
        return re.sub(r"[^a-z0-9]", "", normalized.encode("ascii", "ignore").decode("ascii"))

    header_index = next((
        index for index, row in enumerate(rows[:20])
        if len(row) >= 3 and "matric" in header_text(row[0])
        and "nome" in header_text(row[1]) and "email" in header_text(row[2])
    ), None)
    records: list[tuple[str, str, str]] = []
    for index, row in enumerate(rows):
        if index == header_index:
            continue
        matricula, name, email = (list(row) + [""] * 3)[:3]
        values = (matricula, name, email)
        if not any(str(value or "").strip() for value in values):
            continue
        if not str(matricula or "").strip() or not str(name or "").strip():
            raise ValueError(f"Linha {index + 1}: matrícula e nome são obrigatórios.")
        records.append((matricula, name, email))
    if not records:
        raise ValueError("A planilha não contém usuários para importar.")
    return records

class PlaceholderEntry(ttk.Entry):
    def __init__(self, *args, placeholder: str = "", **kwargs):
        super().__init__(*args, **kwargs)
        self.placeholder = placeholder
        self.placeholder_active = False
        self.bind("<FocusIn>", self._clear_placeholder, add="+")
        self.bind("<FocusOut>", self._restore_placeholder, add="+")
        if not super().get():
            self._show_placeholder()

    def get(self) -> str:
        return "" if self.placeholder_active else super().get()

    def delete(self, first, last=None) -> None:
        self.placeholder_active = False
        self.configure(style="TEntry")
        super().delete(first, last)

    def insert(self, index, string) -> None:
        self.placeholder_active = False
        self.configure(style="TEntry")
        super().insert(index, string)

    def _show_placeholder(self) -> None:
        if self.placeholder and not super().get():
            self.placeholder_active = True
            self.configure(style="Placeholder.TEntry")
            super().insert(0, self.placeholder)

    def _clear_placeholder(self, _event=None) -> None:
        if self.placeholder_active:
            self.placeholder_active = False
            self.configure(style="TEntry")
            super().delete(0, "end")

    def _restore_placeholder(self, _event=None) -> None:
        if not super().get():
            self._show_placeholder()


class LoanAdminApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Empréstimo | Gestão do acervo")
        self.geometry("1320x820")
        self.minsize(1040, 680)
        self.configure(background=COLORS["background"])
        self._load_app_logo()
        self.selected_category_id: int | None = None
        self.selected_item_id: int | None = None
        self.selected_user_id: int | None = None
        self.selected_loan_id: int | None = None
        self.loan_selected_item_ids: set[int] = set()
        self.history_email_in_progress = False
        self.history_progress_dialog: tk.Toplevel | None = None
        self.history_progress_bar: ttk.Progressbar | None = None
        self.startup_loans_popup_shown = False
        self.backup_in_progress = False
        self.app_config = self._load_app_config()
        COLORS.update(THEMES.get(str(self.app_config.get("theme", "Verde")), THEMES["Verde"]))
        self.database_path = Path(str(self.app_config.get("database_path", DEFAULT_DATABASE)))
        self.repository = CategoryRepository(self.database_path)
        if IS_FROZEN and not CONFIG_FILE.exists():
            self.app_config.setdefault("database_path", str(self.database_path))
            self.app_config.setdefault(
                "backup_directory", str(self.database_path.parent),
            )
            try:
                self._save_app_config()
            except OSError as error:
                messagebox.showwarning(
                    "Configuração portátil",
                    "Não foi possível salvar as configurações junto ao executável. "
                    f"Verifique as permissões da pasta.\n\n{error}",
                    parent=self,
                )
        self.withdraw()
        self._configure_styles()
        self._build_layout()
        self.refresh_categories()
        self.refresh_users()
        self._show_splash()

    def _load_app_logo(self) -> None:
        try:
            from PIL import _imagingtk

            _imagingtk.tkinit(self.tk.interpaddr())
            logo = Image.open(resource_path("logo.png")).convert("RGBA")
            icon_image = logo.copy()
            icon_image.thumbnail((64, 64), Image.Resampling.LANCZOS)
            self.app_icon_photo = ImageTk.PhotoImage(icon_image)
            sidebar_image = logo.copy()
            sidebar_image.thumbnail((38, 38), Image.Resampling.LANCZOS)
            self.sidebar_logo_photo = ImageTk.PhotoImage(sidebar_image)
            self.iconphoto(True, self.app_icon_photo)
        except (ImportError, OSError, tk.TclError) as error:
            messagebox.showwarning(
                "Logo do aplicativo",
                f"Não foi possível carregar logo.png:\n{error}",
                parent=self,
            )

    def _load_app_config(self) -> dict[str, object]:
        try:
            legacy_config_paths = (
                (SHARED_CONFIG_FILE, LEGACY_CONFIG_FILE)
                if IS_FROZEN
                else (LEGACY_CONFIG_FILE,)
            )
            config_path = next(
                (
                    path for path in (CONFIG_FILE, *legacy_config_paths)
                    if path.exists()
                ),
                None,
            )
            if config_path is None:
                return {}
            saved = json.loads(config_path.read_text(encoding="utf-8"))
            if not isinstance(saved, dict):
                return {}
            if config_path != CONFIG_FILE:
                try:
                    if IS_FROZEN:
                        saved["backup_directory"] = str(APPLICATION_DIRECTORY)
                        old_database_path = Path(
                            str(saved.get("database_path", DEFAULT_DATABASE))
                        ).expanduser()
                        if (
                            not DEFAULT_DATABASE.exists()
                            and old_database_path.is_file()
                            and old_database_path.resolve() != DEFAULT_DATABASE.resolve()
                        ):
                            _copy_sqlite_database(old_database_path, DEFAULT_DATABASE)
                        if DEFAULT_DATABASE.exists() or not old_database_path.exists():
                            saved["database_path"] = str(DEFAULT_DATABASE)
                        else:
                            saved["database_path"] = str(old_database_path)
                    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
                    CONFIG_FILE.write_text(
                        json.dumps(saved, ensure_ascii=False, indent=2),
                        encoding="utf-8",
                    )
                except OSError as error:
                    messagebox.showwarning(
                        "Migração das configurações",
                        "Não foi possível salvar as configurações junto ao aplicativo. "
                        f"Verifique se a pasta do executável permite gravação.\n\n{error}",
                        parent=self,
                    )
                except (sqlite3.Error, ValueError) as error:
                    messagebox.showwarning(
                        "Migração do banco de dados",
                        "Não foi possível copiar o banco antigo para a pasta do executável. "
                        "O sistema continuará usando o banco antigo, sem apagar seus dados.\n\n"
                        f"{error}",
                        parent=self,
                    )
            return saved
        except (OSError, json.JSONDecodeError):
            return {}

    def _save_database_path(self, path: Path) -> None:
        CONFIG_DIR.mkdir(parents=True, exist_ok=True)
        self.app_config["database_path"] = str(path)
        CONFIG_FILE.write_text(
            json.dumps(self.app_config, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def _save_theme(self, theme_name: str) -> None:
        CONFIG_DIR.mkdir(parents=True, exist_ok=True)
        self.app_config["theme"] = theme_name
        CONFIG_FILE.write_text(
            json.dumps(self.app_config, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def _save_app_config(self) -> None:
        CONFIG_DIR.mkdir(parents=True, exist_ok=True)
        temporary_path = CONFIG_FILE.with_suffix(".tmp")
        temporary_path.write_text(
            json.dumps(self.app_config, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        temporary_path.replace(CONFIG_FILE)

    def _backup_directory(self) -> Path:
        configured_root = self.app_config.get("backup_directory")
        base_directory = (
            Path(str(configured_root)).expanduser().resolve()
            if configured_root is not None
            else self.database_path.parent.resolve()
        )
        if base_directory.name.casefold() == "backup":
            return base_directory
        return base_directory / "backup"

    def _ensure_backup_directory(self) -> Path:
        directory = self._backup_directory()
        directory.mkdir(parents=True, exist_ok=True)
        return directory

    def _backup_is_due(self) -> bool:
        frequency = str(self.app_config.get("backup_frequency", "Desativado"))
        intervals = {"Diário": 1, "Semanal": 7}
        days = intervals.get(frequency)
        if days is None:
            return False
        last_attempt = str(self.app_config.get("backup_last_attempt", ""))
        if not last_attempt:
            return True
        try:
            last_attempt_at = datetime.strptime(last_attempt, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return True
        return datetime.now() - last_attempt_at >= timedelta(days=days)

    def _check_scheduled_backup(self) -> None:
        if not self.winfo_exists():
            return
        self.after(60_000, self._check_scheduled_backup)
        if self.backup_in_progress or not self._backup_is_due():
            return
        self._run_database_backup(automatic=True)

    def _run_database_backup(self, automatic: bool = False) -> None:
        if self.backup_in_progress:
            if not automatic:
                messagebox.showinfo(
                    "Backup em andamento", "Aguarde a conclusão do backup atual.", parent=self,
                )
            return
        self.backup_in_progress = True
        now = datetime.now()
        timestamp = now.strftime("%Y%m%d_%H%M%S_%f")
        destination = self._backup_directory() / f"Emprestimo_backup_{timestamp}.sqlite3"
        try:
            destination = self.repository.create_backup(destination)
            self.app_config["backup_last_attempt"] = now.strftime("%Y-%m-%d %H:%M:%S")
            self.app_config["backup_last_success"] = self.app_config["backup_last_attempt"]
            self._save_app_config()
        except (OSError, sqlite3.Error, ValueError) as error:
            self.app_config["backup_last_attempt"] = now.strftime("%Y-%m-%d %H:%M:%S")
            try:
                self._save_app_config()
            except OSError as config_error:
                error = RuntimeError(f"{error}\nNão foi possível salvar as preferências: {config_error}")
            if automatic:
                messagebox.showerror(
                    "Falha no backup automático",
                    f"O backup programado não foi concluído.\n\n{error}",
                    parent=self,
                )
            else:
                messagebox.showerror("Falha no backup", str(error), parent=self)
            self.backup_in_progress = False
            return
        self.backup_in_progress = False
        self.status_message(f"Backup completo salvo em {destination}")
        if not automatic:
            messagebox.showinfo(
                "Backup concluído", f"O banco de dados foi copiado para:\n{destination}",
                parent=self,
            )

    def _import_database_backup(self) -> None:
        selected_path = filedialog.askopenfilename(
            parent=self, title="Selecionar backup para importar",
            filetypes=[
                ("Bancos SQLite", "*.db *.sqlite *.sqlite3"),
                ("Todos os arquivos", "*.*"),
            ],
        )
        if not selected_path:
            return
        backup_path = Path(selected_path)
        if backup_path.expanduser().resolve() == self.database_path:
            messagebox.showwarning(
                "Arquivo inválido", "Selecione um backup, não o banco que está em uso.",
                parent=self,
            )
            return
        if not messagebox.askyesno(
            "Confirmar importação do banco",
            "A importação substituirá os dados atuais do sistema e não criará um backup "
            "automático. Se quiser guardar os dados atuais, cancele e faça um backup manual "
            "primeiro. Feche o sistema nas outras sessões conectadas a este banco antes de "
            "continuar. Deseja importar mesmo assim?",
            parent=self,
        ):
            return
        if self.backup_in_progress:
            messagebox.showwarning(
                "Operação em andamento",
                "Aguarde a conclusão do backup antes de importar um banco.",
                parent=self,
            )
            return
        self.backup_in_progress = True
        try:
            self.repository.restore_backup(backup_path)
            self.repository = CategoryRepository(self.database_path)
            self.refresh_categories()
            self.refresh_users()
            self.refresh_loans()
            self.status_message("Banco de dados importado e listas atualizadas.")
        except (OSError, sqlite3.Error, ValueError, RuntimeError) as error:
            messagebox.showerror(
                "Falha ao importar o banco",
                f"O banco não foi importado com sucesso.\n\n{error}",
                parent=self,
            )
            return
        finally:
            self.backup_in_progress = False
        messagebox.showinfo(
            "Importação concluída",
            "O banco selecionado foi importado e as listas foram atualizadas. "
            "O conteúdo anterior foi substituído.",
            parent=self,
        )

    def _show_splash(self) -> None:
        splash = tk.Toplevel(self)
        splash.overrideredirect(True)
        splash.configure(background=COLORS["sidebar"])
        splash_width, splash_height = 680, 400
        x = (splash.winfo_screenwidth() - splash_width) // 2
        y = (splash.winfo_screenheight() - splash_height) // 2
        splash.geometry(f"{splash_width}x{splash_height}+{x}+{y}")
        panel = tk.Frame(splash, bg=COLORS["sidebar"], padx=30, pady=26)
        panel.pack(fill="both", expand=True)
        try:
            portrait = Image.open(resource_path("foto.webp"))
            portrait.thumbnail((230, 320), Image.Resampling.LANCZOS)
            self.splash_photo = ImageTk.PhotoImage(portrait)
            tk.Label(panel, image=self.splash_photo, bg=COLORS["sidebar"]).pack(side="left", padx=(0, 28))
        except (ImportError, OSError, tk.TclError) as error:
            splash.destroy()
            self._show_main_window()
            messagebox.showwarning("Imagem de abertura", f"Não foi possível carregar foto.webp:\n{error}", parent=self)
            return
        copy = tk.Frame(panel, bg=COLORS["sidebar"])
        copy.pack(side="left", fill="both", expand=True)
        tk.Label(
            copy, text="EMPRÉSTIMO", bg=COLORS["sidebar"], fg="#FFFFFF",
            font=("Segoe UI Semibold", 24), anchor="w",
        ).pack(anchor="w", pady=(70, 7))
        tk.Label(
            copy, text="Gestão simples do seu acervo", bg=COLORS["sidebar"],
            fg="#D5E3DC", font=("Segoe UI", 12), anchor="w",
        ).pack(anchor="w")
        tk.Label(
            copy, text="por Ozeias Meira Santos de Souza", bg=COLORS["sidebar"],
            fg=COLORS["accent"], font=("Segoe UI Semibold", 9), anchor="w",
        ).pack(anchor="w", pady=(22, 0))
        self.update_idletasks()
        splash.lift()
        self.after(1600, lambda: (splash.destroy(), self._show_main_window()))

    def _show_main_window(self) -> None:
        self.deiconify()
        self.state("zoomed")
        try:
            self._ensure_backup_directory()
        except OSError as error:
            messagebox.showerror(
                "Pasta de backup indisponível",
                f"Não foi possível criar a pasta de backup:\n"
                f"{self._backup_directory()}\n\n{error}",
                parent=self,
            )
        self.after(2500, self._check_scheduled_backup)
        if not self.startup_loans_popup_shown:
            self.startup_loans_popup_shown = True
            self.after(350, self._show_open_loans_popup)

    def _show_open_loans_popup(self) -> None:
        try:
            loans = self.repository.list_loans(active_only=True)
        except sqlite3.Error as error:
            messagebox.showerror(
                "Empréstimos em aberto", f"Não foi possível consultar as retiradas abertas.\n\n{error}",
                parent=self,
            )
            return
        if not loans or not self.winfo_exists():
            return
        dialog = tk.Toplevel(self)
        dialog.title("Empréstimos em aberto")
        dialog.transient(self)
        dialog.geometry("940x500")
        dialog.minsize(720, 380)
        dialog.columnconfigure(0, weight=1)
        dialog.rowconfigure(1, weight=1)
        ttk.Label(
            dialog,
            text=f"Há {len(loans)} item(ns) emprestado(s). Selecione um registro para cobrar "
            "todos os itens em aberto daquela pessoa.",
            style="Subtitle.TLabel", wraplength=880, justify="left",
        ).grid(row=0, column=0, sticky="w", padx=18, pady=(16, 10))
        table_frame = ttk.Frame(dialog)
        table_frame.grid(row=1, column=0, sticky="nsew", padx=18, pady=(0, 12))
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        table = ttk.Treeview(
            table_frame,
            columns=("user", "item", "category", "issued", "email"),
            show="headings", selectmode="browse",
        )
        for column, heading in (
            ("user", "USUÁRIO"), ("item", "ITEM"), ("category", "CATEGORIA"),
            ("issued", "RETIRADO EM"), ("email", "E-MAIL"),
        ):
            table.heading(column, text=heading)
        for column, width in {
            "user": 220, "item": 250, "category": 150, "issued": 155, "email": 180,
        }.items():
            table.column(
                column, width=width,
                anchor="center" if column == "issued" else "w",
            )
        for loan in loans:
            table.insert(
                "", "end", iid=str(loan["id"]),
                values=(
                    loan["user_name"], loan["item_name"], loan["category_name"],
                    loan["issued_at"], loan["email"] or "Sem e-mail",
                ),
            )
        table.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=table.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        table.configure(yscrollcommand=scrollbar.set)
        buttons = ttk.Frame(dialog)
        buttons.grid(row=2, column=0, sticky="e", padx=18, pady=(0, 16))

        def send_reminder() -> None:
            selection = table.selection()
            if not selection:
                messagebox.showinfo(
                    "Selecionar empréstimo",
                    "Selecione um item para identificar a pessoa que receberá a cobrança.",
                    parent=dialog,
                )
                return
            loan_id = int(selection[0])
            selected_loan = next(
                (loan for loan in loans if int(loan["id"]) == loan_id), None,
            )
            if selected_loan is None:
                return
            user_loans = [
                loan for loan in loans
                if int(loan["user_id"]) == int(selected_loan["user_id"])
            ]
            if not messagebox.askyesno(
                "Confirmar cobrança",
                f"Enviar um único e-mail de cobrança para {selected_loan['user_name']} "
                f"com {len(user_loans)} item(ns) em aberto?",
                parent=dialog,
            ):
                return
            try:
                outbox_id = self.repository.create_loan_reminder(loan_id)
            except ValueError as error:
                messagebox.showwarning("Cobrança não registrada", str(error), parent=dialog)
                return
            except sqlite3.Error as error:
                messagebox.showerror("Erro no banco de dados", str(error), parent=dialog)
                return
            dialog.destroy()
            self.status_message("Cobrança registrada; enviando e-mail.")
            self._dispatch_loan_emails(loan_id, [outbox_id])

        ttk.Button(buttons, text="Fechar", command=dialog.destroy).pack(
            side="right", padx=(8, 0),
        )
        ttk.Button(
            buttons, text="Cobrar pessoa selecionada",
            style="Accent.TButton", command=send_reminder,
        ).pack(side="right")
        dialog.update_idletasks()
        x = self.winfo_rootx() + max(0, (self.winfo_width() - dialog.winfo_width()) // 2)
        y = self.winfo_rooty() + max(0, (self.winfo_height() - dialog.winfo_height()) // 2)
        dialog.geometry(f"+{x}+{y}")
        dialog.grab_set()

    def _configure_styles(self) -> None:
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("TFrame", background=COLORS["background"])
        style.configure("Surface.TFrame", background=COLORS["surface"])
        style.configure(
            "TLabel", background=COLORS["background"], foreground=COLORS["text"],
            font=("Segoe UI", 10),
        )
        style.configure("Placeholder.TEntry", foreground=COLORS["muted"])
        style.configure(
            "Title.TLabel", font=("Segoe UI Semibold", 23), foreground=COLORS["text"]
        )
        style.configure(
            "Subtitle.TLabel", font=("Segoe UI", 10), foreground=COLORS["muted"]
        )
        style.configure(
            "CardTitle.TLabel", font=("Segoe UI Semibold", 10), foreground=COLORS["muted"]
        )
        style.configure(
            "CardValue.TLabel", font=("Segoe UI Semibold", 21), foreground=COLORS["text"]
        )
        style.configure(
            "TButton", font=("Segoe UI Semibold", 9), padding=(12, 9),
            background=COLORS["surface"], foreground=COLORS["text"], bordercolor=COLORS["border"],
        )
        style.map("TButton", background=[("active", "#E8EEE9")])
        style.configure(
            "Accent.TButton", background=COLORS["accent"], foreground="#192B28",
            bordercolor=COLORS["accent"],
        )
        style.map("Accent.TButton", background=[("active", "#E8AA4A")])
        style.configure(
            "Danger.TButton", foreground=COLORS["danger"], bordercolor=COLORS["border"]
        )
        style.configure(
            "Treeview", rowheight=38, font=("Segoe UI", 10), background=COLORS["surface"],
            fieldbackground=COLORS["surface"], foreground=COLORS["text"], borderwidth=0,
        )
        style.configure(
            "Treeview.Heading", font=("Segoe UI Semibold", 9), padding=(10, 11),
            background="#E9EFEB", foreground=COLORS["text"], relief="flat",
        )
        style.map(
            "Treeview",
            background=[("selected", COLORS["border"])],
            foreground=[("selected", COLORS["text"])],
        )

    def _build_layout(self) -> None:
        self.columnconfigure(1, weight=1)
        self.rowconfigure(0, weight=1)

        sidebar = tk.Frame(self, bg=COLORS["sidebar"], width=228)
        self.sidebar_frame = sidebar
        sidebar.grid(row=0, column=0, sticky="nsew")
        sidebar.grid_propagate(False)
        brand = tk.Frame(sidebar, bg=COLORS["sidebar"])
        self.sidebar_brand_frame = brand
        brand.pack(fill="x", padx=20, pady=(24, 6))
        if hasattr(self, "sidebar_logo_photo"):
            tk.Label(
                brand, image=self.sidebar_logo_photo, bg=COLORS["sidebar"],
            ).pack(side="left", padx=(0, 10))
        tk.Label(
            brand, text="EMPRÉSTIMO", bg=COLORS["sidebar"], fg="#F5F6EF",
            font=("Segoe UI Semibold", 15), anchor="w",
        ).pack(side="left", fill="x", expand=True)
        tk.Label(
            sidebar, text="PAINEL ADMINISTRATIVO", bg=COLORS["sidebar"], fg="#A9BCB4",
            font=("Segoe UI", 8), anchor="w",
        ).pack(fill="x", padx=24, pady=(0, 34))
        self._sidebar_label(sidebar, "MENU")
        self.overview_nav_button = self._sidebar_button(
            sidebar, "⌂   Visão geral", command=self._show_dashboard, active=True,
        )
        self._sidebar_label(sidebar, "CADASTROS")
        self.category_nav_button = self._sidebar_button(
            sidebar, "▦   Categorias", command=self._show_categories,
        )
        self.item_nav_button = self._sidebar_button(
            sidebar, "▣   Itens", command=self._show_items,
        )
        self.user_nav_button = self._sidebar_button(
            sidebar, "♙   Usuários", command=self._show_users,
        )
        self._sidebar_label(sidebar, "MOVIMENTAÇÃO")
        self.loan_nav_button = self._sidebar_button(
            sidebar, "⇄   Empréstimos", command=self._show_loans,
        )
        self.sidebar_separator = tk.Frame(sidebar, bg=COLORS["sidebar_hover"], height=1)
        self.sidebar_separator.pack(fill="x", padx=22, pady=(26, 18))
        self.settings_nav_button = self._sidebar_button(
            sidebar, "⚙   Configurações", command=self.open_settings,
        )
        self.about_nav_button = self._sidebar_button(
            sidebar, "ⓘ   Sobre", command=self._show_about,
        )
        tk.Label(
            sidebar, text="EMPRÉSTIMO DE ITENS\nGestão do acervo", bg=COLORS["sidebar"],
            fg="#A9BCB4", font=("Segoe UI", 8), justify="left", anchor="w",
        ).pack(side="bottom", fill="x", padx=24, pady=22)

        self.page_container = ttk.Frame(self)
        self.page_container.grid(row=0, column=1, sticky="nsew")
        self.page_container.columnconfigure(0, weight=1)
        self.page_container.rowconfigure(0, weight=1)
        self.content = ttk.Frame(self.page_container, padding=(34, 28, 34, 26))
        self.content.grid(row=0, column=0, sticky="nsew")
        self.content.columnconfigure(0, weight=1)
        self.content.rowconfigure(4, weight=1)

        topbar = ttk.Frame(self.content)
        topbar.grid(row=0, column=0, sticky="ew", pady=(0, 22))
        topbar.columnconfigure(0, weight=1)
        heading = ttk.Frame(topbar)
        heading.grid(row=0, column=0, sticky="w")
        ttk.Label(heading, text="Categorias", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            heading, text="Organize os tipos de itens disponíveis para empréstimo.",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(3, 0))
        ttk.Button(topbar, text="⚙  Banco de dados", command=self.open_settings).grid(
            row=0, column=1, sticky="e", padx=(14, 0)
        )

        cards = ttk.Frame(self.content)
        cards.grid(row=1, column=0, sticky="ew", pady=(0, 20))
        for column in range(3):
            cards.columnconfigure(column, weight=1, uniform="card")
        self.total_value = self._stat_card(cards, 0, "CATEGORIAS CADASTRADAS", "0")
        self.filtered_value = self._stat_card(cards, 1, "EXIBIDAS NA LISTA", "0")
        self.database_value = self._stat_card(cards, 2, "ARMAZENAMENTO", "SQLite")

        form = ttk.Frame(self.content, style="Surface.TFrame", padding=(18, 15))
        form.grid(row=2, column=0, sticky="ew", pady=(0, 14))
        form.columnconfigure(0, weight=1)
        ttk.Label(form, text="Nome da categoria", style="CardTitle.TLabel").grid(
            row=0, column=0, sticky="w", pady=(0, 8)
        )
        self.name_entry = PlaceholderEntry(
            form, font=("Segoe UI", 11), placeholder="Ex.: Equipamentos de informática",
        )
        self.name_entry.grid(row=1, column=0, sticky="ew", padx=(0, 12), ipady=5)
        self.name_entry.bind("<Return>", lambda _event: self.save_category())
        self.save_button = ttk.Button(
            form, text="＋  Adicionar categoria", style="Accent.TButton",
            command=self.save_category,
        )
        self.save_button.grid(row=1, column=1, sticky="e")
        self.cancel_edit_button = ttk.Button(
            form, text="Cancelar edição", command=self.clear_form
        )

        toolbar = ttk.Frame(self.content)
        toolbar.grid(row=3, column=0, sticky="ew", pady=(0, 10))
        toolbar.columnconfigure(0, weight=1)
        search_frame = ttk.Frame(toolbar, style="Surface.TFrame", padding=(10, 3))
        search_frame.grid(row=0, column=0, sticky="ew", padx=(0, 12))
        search_frame.columnconfigure(1, weight=1)
        ttk.Label(search_frame, text="⌕", style="Subtitle.TLabel").grid(
            row=0, column=0, padx=(2, 8)
        )
        self.search_var = tk.StringVar()
        self.search_entry = PlaceholderEntry(
            search_frame, textvariable=self.search_var, font=("Segoe UI", 10),
            placeholder="Pesquisar categoria...",
        )
        self.search_var.trace_add(
            "write",
            lambda *_args: self.refresh_categories()
            if not self.search_entry.placeholder_active else None,
        )
        search_entry = self.search_entry
        search_entry.grid(row=0, column=1, sticky="ew", ipady=6)
        ttk.Button(toolbar, text="⇩  Excel", command=self.export_excel).grid(row=0, column=1, padx=(0, 8))
        ttk.Button(toolbar, text="⇩  PDF", command=self.export_pdf).grid(row=0, column=2)

        table_frame = ttk.Frame(self.content, style="Surface.TFrame", padding=(0, 0, 0, 4))
        table_frame.grid(row=4, column=0, sticky="nsew")
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(1, weight=1)
        ttk.Label(table_frame, text="  LISTA DE CATEGORIAS", style="CardTitle.TLabel").grid(
            row=0, column=0, sticky="w", padx=16, pady=(15, 12)
        )
        columns = ("id", "name", "created_at")
        self.category_table = ttk.Treeview(
            table_frame, columns=columns, show="headings", selectmode="browse"
        )
        self.category_table.heading("id", text="ID")
        self.category_table.heading("name", text="CATEGORIA")
        self.category_table.heading("created_at", text="CADASTRADA EM")
        self.category_table.column("id", width=70, minwidth=60, anchor="center", stretch=False)
        self.category_table.column("name", width=420, minwidth=180, anchor="w")
        self.category_table.column("created_at", width=180, minwidth=150, anchor="center", stretch=False)
        self.category_table.grid(row=1, column=0, sticky="nsew", padx=(10, 0), pady=(0, 4))
        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=self.category_table.yview)
        scrollbar.grid(row=1, column=1, sticky="ns", padx=(0, 8), pady=(0, 4))
        self.category_table.configure(yscrollcommand=scrollbar.set)
        self.category_table.bind("<Double-1>", lambda _event: self.edit_selected())
        actions = ttk.Frame(self.content)
        actions.grid(row=5, column=0, sticky="e", pady=(12, 0))
        ttk.Button(actions, text="Editar selecionada", command=self.edit_selected).pack(side="left", padx=(0, 8))
        ttk.Button(
            actions, text="Excluir selecionada", style="Danger.TButton",
            command=self.delete_selected,
        ).pack(side="left")
        self._build_items_page()
        self._build_users_page()
        self._build_loans_page()
        self._build_dashboard_page()
        self._build_about_page()
        self.dashboard_page.tkraise()

    def _build_dashboard_page(self) -> None:
        self.dashboard_page = ttk.Frame(self.page_container, padding=(34, 28, 34, 26))
        self.dashboard_page.grid(row=0, column=0, sticky="nsew")
        self.dashboard_page.columnconfigure(0, weight=1)
        self.dashboard_page.columnconfigure(1, weight=4)
        self.dashboard_page.rowconfigure(2, weight=4)
        self.dashboard_page.rowconfigure(3, weight=0)

        topbar = ttk.Frame(self.dashboard_page)
        topbar.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 22))
        topbar.columnconfigure(0, weight=1)
        heading = ttk.Frame(topbar)
        heading.grid(row=0, column=0, sticky="w")
        ttk.Label(heading, text="Visão geral", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            heading, text="Visão rápida do acervo, movimentações e notificações.",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(3, 0))
        ttk.Button(
            topbar, text="＋  Novo item", style="Accent.TButton", command=self._show_items,
        ).grid(row=0, column=1, sticky="e", padx=(0, 8))
        ttk.Button(
            topbar, text="＋  Categoria", command=self._show_categories,
        ).grid(row=0, column=2, sticky="e", padx=(0, 8))
        ttk.Button(
            topbar, text="✉  Fila de e-mails", command=self._show_email_queue,
        ).grid(row=0, column=3, sticky="e", padx=(0, 8))
        ttk.Button(
            topbar, text="＋  Usuário", command=self._show_users,
        ).grid(row=0, column=4, sticky="e")

        cards = ttk.Frame(self.dashboard_page)
        cards.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(0, 20))
        for column in range(4):
            cards.columnconfigure(column, weight=1, uniform="overview-card")
        self.dashboard_items_value = self._stat_card(cards, 0, "ITENS NO ACERVO", "0")
        self.dashboard_available_value = self._stat_card(cards, 1, "DISPONÍVEIS", "0")
        self.dashboard_users_value = self._stat_card(cards, 2, "USUÁRIOS", "0")
        self.dashboard_loans_value = self._stat_card(cards, 3, "EMPRÉSTIMOS EM ABERTO", "0")

        category_panel = ttk.Frame(
            self.dashboard_page, style="Surface.TFrame", padding=(8, 10),
        )
        category_panel.grid(row=2, column=0, sticky="nsew", padx=(0, 8))
        category_panel.columnconfigure(0, weight=1)
        category_panel.columnconfigure(1, weight=0)
        category_panel.rowconfigure(1, weight=1)
        ttk.Label(
            category_panel, text="ACERVO POR CATEGORIA", style="CardTitle.TLabel",
        ).grid(row=0, column=0, sticky="w", pady=(0, 10))
        self.dashboard_category_table = ttk.Treeview(
            category_panel, columns=("category", "items"), show="headings",
        )
        self.dashboard_category_table.heading("category", text="CATEGORIA")
        self.dashboard_category_table.heading("items", text="ITENS")
        self.dashboard_category_table.column("category", width=100, minwidth=80, anchor="w")
        self.dashboard_category_table.column("items", width=42, minwidth=36, anchor="center", stretch=False)
        self.dashboard_category_table.grid(row=1, column=0, sticky="nsew")
        category_scrollbar = ttk.Scrollbar(
            category_panel, orient="vertical", command=self.dashboard_category_table.yview,
        )
        category_scrollbar.grid(row=1, column=1, sticky="ns")
        self.dashboard_category_table.configure(yscrollcommand=category_scrollbar.set)

        recent_panel = ttk.Frame(
            self.dashboard_page, style="Surface.TFrame", padding=(14, 14),
        )
        recent_panel.grid(row=2, column=1, sticky="nsew", padx=(8, 0))
        recent_panel.columnconfigure(0, weight=1)
        recent_panel.columnconfigure(1, weight=0)
        recent_panel.rowconfigure(2, weight=3)
        recent_panel.rowconfigure(4, weight=1)
        ttk.Label(
            recent_panel, text="ITENS ADICIONADOS RECENTEMENTE", style="CardTitle.TLabel",
        ).grid(row=0, column=0, sticky="w", pady=(0, 10))
        recent_filters = ttk.Frame(recent_panel)
        recent_filters.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        recent_filters.columnconfigure(1, weight=1)
        self.dashboard_recent_category_var = tk.StringVar(value="Todas as categorias")
        self.dashboard_recent_category_combo = ttk.Combobox(
            recent_filters, textvariable=self.dashboard_recent_category_var,
            state="readonly", width=24,
        )
        self.dashboard_recent_category_combo.grid(row=0, column=0, sticky="w", padx=(0, 8))
        self.dashboard_recent_category_combo.bind(
            "<<ComboboxSelected>>", lambda _event: self._refresh_dashboard_recent_items(),
        )
        self.dashboard_recent_search_var = tk.StringVar()
        self.dashboard_recent_search_entry = PlaceholderEntry(
            recent_filters, textvariable=self.dashboard_recent_search_var,
            placeholder="Filtrar por nome ou número do item...",
        )
        self.dashboard_recent_search_entry.grid(row=0, column=1, sticky="ew", ipady=4)
        self.dashboard_recent_search_var.trace_add(
            "write",
            lambda *_args: self._refresh_dashboard_recent_items()
            if not self.dashboard_recent_search_entry.placeholder_active else None,
        )
        self.dashboard_recent_table = ttk.Treeview(
            recent_panel, columns=("number", "name", "category", "status"), show="headings",
        )
        self.dashboard_recent_table.bind("<Double-1>", self._on_dashboard_item_double_click)
        self.dashboard_recent_table.heading("number", text="Nº")
        self.dashboard_recent_table.heading("name", text="ITEM")
        self.dashboard_recent_table.heading("category", text="CATEGORIA")
        self.dashboard_recent_table.heading("status", text="ESTADO")
        self.dashboard_recent_table.column("number", width=55, minwidth=45, anchor="center", stretch=False)
        self.dashboard_recent_table.column("name", width=280, minwidth=140, anchor="w")
        self.dashboard_recent_table.column("category", width=170, minwidth=100, anchor="w")
        self.dashboard_recent_table.column("status", width=85, minwidth=75, anchor="center", stretch=False)
        self.dashboard_recent_table.grid(row=2, column=0, sticky="nsew")
        recent_scrollbar = ttk.Scrollbar(
            recent_panel, orient="vertical", command=self.dashboard_recent_table.yview,
        )
        recent_scrollbar.grid(row=2, column=1, sticky="ns")
        self.dashboard_recent_table.configure(yscrollcommand=recent_scrollbar.set)
        ttk.Label(
            recent_panel,
            text="✓ Disponível   ↗ Emprestado   ⚒ Manutenção",
            style="Subtitle.TLabel",
        ).grid(row=5, column=0, sticky="w", pady=(6, 0))
        ttk.Label(
            recent_panel, text="USUÁRIOS POR TIPO", style="CardTitle.TLabel",
        ).grid(row=3, column=0, sticky="w", pady=(10, 6))
        self.dashboard_user_type_table = ttk.Treeview(
            recent_panel, columns=("type", "count"), show="headings", height=3,
        )
        self.dashboard_user_type_table.heading("type", text="TIPO")
        self.dashboard_user_type_table.heading("count", text="USUÁRIOS")
        self.dashboard_user_type_table.column("type", width=250, minwidth=150, anchor="w")
        self.dashboard_user_type_table.column("count", width=90, minwidth=70, anchor="center", stretch=False)
        self.dashboard_user_type_table.grid(row=4, column=0, sticky="nsew")

        self.dashboard_email_summary = ttk.Label(
            self.dashboard_page, style="Subtitle.TLabel", anchor="w",
        )
        self.dashboard_email_summary.grid(
            row=3, column=0, columnspan=2, sticky="ew", pady=(8, 0),
        )

    def _build_about_page(self) -> None:
        self.about_page = ttk.Frame(self.page_container, padding=(34, 28, 34, 26))
        self.about_page.grid(row=0, column=0, sticky="nsew")
        self.about_page.columnconfigure(0, weight=1)
        self.about_page.rowconfigure(2, weight=1)
        ttk.Label(self.about_page, text="Sobre", style="Title.TLabel").grid(
            row=0, column=0, sticky="w",
        )
        ttk.Label(
            self.about_page, text="Quem criou e o que você pode fazer com o sistema.",
            style="Subtitle.TLabel",
        ).grid(row=1, column=0, sticky="w", pady=(4, 20))
        card = ttk.Frame(self.about_page, style="Surface.TFrame", padding=24)
        card.grid(row=2, column=0, sticky="nsew")
        card.columnconfigure(1, weight=1)
        card.rowconfigure(5, weight=1)
        try:
            portrait = Image.open(resource_path("foto.webp"))
            portrait.thumbnail((300, 330), Image.Resampling.LANCZOS)
            self.about_photo = ImageTk.PhotoImage(portrait)
            ttk.Label(card, image=self.about_photo).grid(
                row=0, column=0, rowspan=6, sticky="nw", padx=(0, 28),
            )
        except (ImportError, OSError, tk.TclError) as error:
            messagebox.showwarning(
                "Foto indisponível",
                f"Não foi possível carregar foto.webp:\n{error}",
                parent=self,
            )
        ttk.Label(
            card, text="Ozeias Meira Santos de Souza", style="Title.TLabel",
        ).grid(row=0, column=1, sticky="w", pady=(10, 4))
        ttk.Label(
            card, text="Analista e Desenvolvedor de Sistemas",
            style="CardTitle.TLabel",
        ).grid(row=1, column=1, sticky="w")
        ttk.Label(
            card, text="✉  ozeeiiaass@gmail.com", style="Subtitle.TLabel",
        ).grid(row=2, column=1, sticky="w", pady=(12, 4))
        ttk.Label(
            card, text="☎  69 9 92726386  ·  Instagram: @ozeias_dev",
            style="Subtitle.TLabel",
        ).grid(row=3, column=1, sticky="w", pady=(0, 14))
        ttk.Label(
            card, text="Trajetória", style="CardTitle.TLabel",
        ).grid(row=4, column=1, sticky="nw")
        biography_frame = ttk.Frame(card)
        biography_frame.grid(row=5, column=1, sticky="nsew", pady=(8, 0))
        biography_frame.columnconfigure(0, weight=1)
        biography_frame.rowconfigure(0, weight=1)
        biography = tk.Text(
            biography_frame, wrap="word", height=12, width=70,
            relief="flat", borderwidth=0, padx=2, pady=2,
            background=COLORS["surface"], foreground=COLORS["text"],
            font=("Segoe UI", 10), spacing1=2, spacing3=3,
        )
        biography.insert(
            "1.0",
            "Graduado em Geografia pela Universidade Norte do Paraná (UNOPAR) e Análise e Desenvolvimento de Sistemas pela Faculdade São Francisco de Assis, com sólida "
            "formação interdisciplinar e trajetória acadêmica e profissional marcada pela "
            "integração entre educação, gestão e tecnologia.\n\n"
            "• Formação complementar:\n"
            "   – Pós-graduação Lato Sensu em Docência do Ensino Superior e Tutoria do Ensino a Distância;\n"
            "   – Pós-graduação Lato Sensu em Ensino de Geografia e História;\n"
            "   – Pós-graduação Lato Sensu em Educação de Jovens e Adultos (EJA);\n"
            "   – Pós-graduação Lato Sensu em Logística Empresarial;\n"
            "   – Pós-graduação Lato Sensu em Arquitetura de Sistemas e Cloud Computing;\n\n"
            "   – Pós-graduando Lato Sensu em Planejamento e Gestão Estratégica na Rede Federal de Educação.\n\n"
            "   – Futuro aluno do Mestrado Profissional em Propriedade Intelectual e Transferência de Tecnologia.\n"
            "\n• Propriedade intelectual:\n"
            "   – Detém 11 certificados de registro de software junto ao Instituto Nacional "
            "da Propriedade Industrial (INPI), evidenciando experiência prática em "
            "desenvolvimento e formalização de soluções tecnológicas.\n\n"
            "• Projetos de Estimada Relevância:\n"
            "   – Autor e Desenvolvedor do SatAgroBr (Projeto apresentado a Agência Nacional da Aeronáutica e Espaço - NASA);\n"
            "   – Autor e Desenvolvedor do SEINOTAS (Extensão do navegador com mais de 1.000 usuários).\n\n"
            "• Experiência profissional:\n"
            "   – Mais de quatro anos de experiência na área de gestão administrativa no "
            "Instituto Federal de Rondônia, atuando como Chefe de Gabinete, com desempenho em "
            "planejamento estratégico, assessoramento institucional, organização de "
            "processos administrativos e articulação entre setores.\n\n"
            "   – Cinco anos de experiência administrativa na Aeronáutica, com atuação "
            "voltada à organização, controle e gestão de rotinas administrativas.\n\n"
            "Sua trajetória profissional integra educação, gestão pública e tecnologia, "
            "com foco em inovação, organização institucional e desenvolvimento de "
            "soluções aplicadas.",
        )
        biography.configure(state="disabled")
        biography.grid(row=0, column=0, sticky="nsew")
        biography_scrollbar = ttk.Scrollbar(
            biography_frame, orient="vertical", command=biography.yview,
        )
        biography_scrollbar.grid(row=0, column=1, sticky="ns")
        biography.configure(yscrollcommand=biography_scrollbar.set)
        ttk.Label(
            self.about_page,
            text="Uma ferramenta para organizar o acervo, acompanhar retiradas e devoluções "
            "e manter os usuários informados por e-mail.",
            style="Subtitle.TLabel",
        ).grid(row=3, column=0, sticky="w", pady=(18, 0))

    def _show_about(self) -> None:
        self._set_active_navigation("about")
        self.about_page.tkraise()

    def _show_email_queue(self) -> None:
        self._show_loans()
        self.loans_notebook.select(2)
        self.refresh_email_queue()

    def refresh_dashboard(self) -> None:
        if not hasattr(self, "dashboard_category_table"):
            return
        categories = self.repository.list_category_item_counts()
        total_items = sum(row["item_count"] for row in categories)
        self.dashboard_items_value.configure(text=str(total_items))
        self.dashboard_available_value.configure(
            text=str(self.repository.count_items(DEFAULT_ITEM_STATUS)),
        )
        self.dashboard_users_value.configure(text=str(self.repository.count_users()))
        self.dashboard_loans_value.configure(text=str(self.repository.count_active_loans()))
        self.dashboard_email_summary.configure(
            text=(
                f"E-mails: {self.repository.count_outbox_messages('sent')} enviados  ·  "
                f"{self.repository.count_outbox_messages('pending')} na fila  ·  "
                f"{self.repository.count_outbox_messages('failed')} com falha"
            ),
        )

        for item_id in self.dashboard_category_table.get_children():
            self.dashboard_category_table.delete(item_id)
        for category in categories:
            self.dashboard_category_table.insert(
                "", "end", iid=str(category["id"]),
                values=(category["name"], category["item_count"]),
            )
        if not categories:
            self.dashboard_category_table.insert(
                "", "end", iid="empty-categories", values=("Nenhuma categoria cadastrada", "0"),
            )

        self.dashboard_recent_category_ids = {
            f"{category['name']} · #{category['id']}": int(category["id"])
            for category in categories
        }
        category_options = ["Todas as categorias", *self.dashboard_recent_category_ids]
        self.dashboard_recent_category_combo.configure(values=category_options)
        if self.dashboard_recent_category_var.get() not in category_options:
            self.dashboard_recent_category_var.set("Todas as categorias")
        self._refresh_dashboard_recent_items()

        for user_id in self.dashboard_user_type_table.get_children():
            self.dashboard_user_type_table.delete(user_id)
        for user_type in USER_TYPES:
            self.dashboard_user_type_table.insert(
                "", "end", iid=f"user-type-{user_type}",
                values=(user_type, self.repository.count_users(user_type)),
            )

    @staticmethod
    def _item_status_icon(status: str) -> str:
        return {
            "Disponível": "✓",
            "Emprestado": "↗",
            "Em manutenção": "⚒",
        }.get(status, "•")

    def _refresh_dashboard_recent_items(self) -> None:
        if not hasattr(self, "dashboard_recent_table"):
            return
        category_id = self.dashboard_recent_category_ids.get(
            self.dashboard_recent_category_var.get(),
        )
        search = (
            "" if self.dashboard_recent_search_entry.placeholder_active
            else self.dashboard_recent_search_var.get()
        )
        recent_items = self.repository.list_recent_items(
            limit=100, category_id=category_id, search=search,
        )
        for item_id in self.dashboard_recent_table.get_children():
            self.dashboard_recent_table.delete(item_id)
        for item in recent_items:
            self.dashboard_recent_table.insert(
                "", "end", iid=str(item["id"]),
                values=(
                    item["number"], item["name"], item["category_name"],
                    self._item_status_icon(str(item["status"])),
                ),
            )
        if not recent_items:
            self.dashboard_recent_table.insert(
                "", "end", iid="empty-items",
                values=("", "Nenhum item corresponde ao filtro", "", ""),
            )

    def _on_dashboard_item_double_click(self, event) -> None:
        item_id = self.dashboard_recent_table.identify_row(event.y)
        if not item_id or item_id.startswith("empty-"):
            return
        self.show_item_holder(int(item_id))

    def show_item_holder(self, item_id: int) -> None:
        item = self.repository.get_item(item_id)
        if item is None:
            messagebox.showwarning("Item não encontrado", "O item não existe mais.", parent=self)
            return
        loan = self.repository.get_active_loan_for_item(item_id)
        if loan is None:
            messagebox.showinfo(
                "Item disponível",
                f"{item['name']} (Nº {item['number']}) não possui empréstimo em aberto.",
                parent=self,
            )
            return
        messagebox.showinfo(
            "Responsável pelo item",
            f"Item: {loan['item_name']} (Nº {loan['number']})\n"
            f"Categoria: {loan['category_name']}\n"
            f"Usuário: {loan['user_name']}\n"
            f"Matrícula: {loan['matricula']}\n"
            f"E-mail: {loan['email'] or 'não informado'}\n"
            f"Retirado em: {loan['issued_at']}",
            parent=self,
        )

    def _build_items_page(self) -> None:
        self.items_page = ttk.Frame(self.page_container, padding=(34, 28, 34, 26))
        self.items_page.grid(row=0, column=0, sticky="nsew")
        self.items_page.columnconfigure(0, weight=1)
        self.items_page.rowconfigure(4, weight=1)

        topbar = ttk.Frame(self.items_page)
        topbar.grid(row=0, column=0, sticky="ew", pady=(0, 22))
        topbar.columnconfigure(0, weight=1)
        heading = ttk.Frame(topbar)
        heading.grid(row=0, column=0, sticky="w")
        ttk.Label(heading, text="Itens", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            heading,
            text="Cadastre e organize os itens por categoria.  ✓ disponível · ↗ emprestado · ⚒ manutenção",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(3, 0))
        ttk.Button(topbar, text="⚙  Banco de dados", command=self.open_settings).grid(
            row=0, column=1, sticky="e", padx=(14, 0)
        )

        cards = ttk.Frame(self.items_page)
        cards.grid(row=1, column=0, sticky="ew", pady=(0, 18))
        for column in range(3):
            cards.columnconfigure(column, weight=1, uniform="item-card")
        self.item_total_value = self._stat_card(cards, 0, "ITENS CADASTRADOS", "0")
        self.item_visible_value = self._stat_card(cards, 1, "EXIBIDOS NA LISTA", "0")
        self.item_category_count_value = self._stat_card(cards, 2, "CATEGORIAS DISPONÍVEIS", "0")

        form = ttk.Frame(self.items_page, style="Surface.TFrame", padding=(16, 13))
        form.grid(row=2, column=0, sticky="ew", pady=(0, 12))
        form.columnconfigure(1, weight=2)
        form.columnconfigure(2, weight=1)
        form.columnconfigure(3, weight=1)
        form.columnconfigure(4, weight=1)
        ttk.Label(form, text="Número", style="CardTitle.TLabel").grid(
            row=0, column=0, sticky="w", padx=(0, 10), pady=(0, 7)
        )
        ttk.Label(form, text="Nome do item", style="CardTitle.TLabel").grid(
            row=0, column=1, sticky="w", padx=(0, 10), pady=(0, 7)
        )
        ttk.Label(form, text="Categoria", style="CardTitle.TLabel").grid(
            row=0, column=2, sticky="w", pady=(0, 7), padx=(0, 10)
        )
        ttk.Label(form, text="Status", style="CardTitle.TLabel").grid(
            row=0, column=3, sticky="w", pady=(0, 7), padx=(0, 10)
        )
        self.item_number_entry = PlaceholderEntry(
            form, width=8, font=("Segoe UI", 10), placeholder="Nº",
        )
        self.item_number_entry.grid(row=1, column=0, sticky="ew", padx=(0, 10), ipady=5)
        self.item_name_entry = PlaceholderEntry(
            form, font=("Segoe UI", 10), placeholder="Ex.: Notebook Dell",
        )
        self.item_name_entry.grid(row=1, column=1, sticky="ew", padx=(0, 10), ipady=5)
        self.item_category_var = tk.StringVar()
        self.item_category_combo = ttk.Combobox(
            form, textvariable=self.item_category_var, state="readonly", width=18,
        )
        self.item_category_combo.grid(row=1, column=2, sticky="ew", padx=(0, 10), ipady=4)
        self.item_status_var = tk.StringVar(value=DEFAULT_ITEM_STATUS)
        self.item_status_combo = ttk.Combobox(
            form, textvariable=self.item_status_var, values=ITEM_STATUSES,
            state="readonly", width=16,
        )
        self.item_status_combo.grid(row=1, column=3, sticky="ew", padx=(0, 10), ipady=4)
        self.item_save_button = ttk.Button(
            form, text="＋  Adicionar item", style="Accent.TButton", command=self.save_item,
        )
        self.item_save_button.grid(row=1, column=4, sticky="ew")
        self.item_cancel_button = ttk.Button(
            form, text="Cancelar edição", command=self.clear_item_form,
        )
        self.item_number_entry.bind("<Return>", lambda _event: self.save_item())
        self.item_name_entry.bind("<Return>", lambda _event: self.save_item())

        toolbar = ttk.Frame(self.items_page)
        toolbar.grid(row=3, column=0, sticky="ew", pady=(0, 10))
        toolbar.columnconfigure(0, weight=1)
        search_frame = ttk.Frame(toolbar, style="Surface.TFrame", padding=(9, 3))
        search_frame.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        search_frame.columnconfigure(1, weight=1)
        ttk.Label(search_frame, text="⌕", style="Subtitle.TLabel").grid(
            row=0, column=0, padx=(2, 7)
        )
        self.item_search_var = tk.StringVar()
        self.item_search_entry = PlaceholderEntry(
            search_frame, textvariable=self.item_search_var, font=("Segoe UI", 9),
            placeholder="Pesquisar por número, item ou categoria...",
        )
        self.item_search_entry.grid(row=0, column=1, sticky="ew", ipady=5)
        self.item_search_var.trace_add(
            "write",
            lambda *_args: self.refresh_items()
            if not self.item_search_entry.placeholder_active else None,
        )

        self.item_filter_var = tk.StringVar(value="Todas as categorias")
        self.item_filter_combo = ttk.Combobox(
            toolbar, textvariable=self.item_filter_var, state="readonly", width=24,
        )
        self.item_filter_combo.grid(row=0, column=1, padx=(0, 8))
        self.item_filter_combo.bind("<<ComboboxSelected>>", lambda _event: self.refresh_items())
        self.item_view_var = tk.StringVar(value="list")
        view_switch = ttk.Frame(toolbar, style="Surface.TFrame", padding=2)
        view_switch.grid(row=0, column=2, padx=(0, 8))
        ttk.Radiobutton(
            view_switch, text="Lista", value="list", variable=self.item_view_var,
            style="Toolbutton", command=self.set_item_view,
        ).pack(side="left")
        ttk.Radiobutton(
            view_switch, text="Grade", value="grid", variable=self.item_view_var,
            style="Toolbutton", command=self.set_item_view,
        ).pack(side="left")
        ttk.Button(toolbar, text="Excel", command=self.export_items_excel).grid(
            row=0, column=3, padx=(0, 7)
        )
        ttk.Button(toolbar, text="PDF", command=self.export_items_pdf).grid(row=0, column=4)

        self.item_results = ttk.Frame(self.items_page, style="Surface.TFrame")
        self.item_results.grid(row=4, column=0, sticky="nsew")
        self.item_results.columnconfigure(0, weight=1)
        self.item_results.rowconfigure(0, weight=1)

        list_frame = ttk.Frame(self.item_results, style="Surface.TFrame", padding=(0, 4, 0, 4))
        list_frame.grid(row=0, column=0, sticky="nsew")
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)
        self.item_table = ttk.Treeview(
            list_frame, columns=("number", "name", "category", "status"), show="headings", selectmode="browse",
        )
        self.item_table.heading("number", text="NÚMERO")
        self.item_table.heading("name", text="ITEM")
        self.item_table.heading("category", text="CATEGORIA")
        self.item_table.heading("status", text="ESTADO")
        self.item_table.column("number", width=100, minwidth=80, anchor="center", stretch=False)
        self.item_table.column("name", width=320, minwidth=150, anchor="w")
        self.item_table.column("category", width=240, minwidth=150, anchor="w")
        self.item_table.column("status", width=72, minwidth=60, anchor="center", stretch=False)
        self.item_table.grid(row=0, column=0, sticky="nsew", padx=(10, 0))
        item_scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.item_table.yview)
        item_scrollbar.grid(row=0, column=1, sticky="ns", padx=(0, 8))
        self.item_table.configure(yscrollcommand=item_scrollbar.set)
        self.item_table.bind("<<TreeviewSelect>>", self._on_item_table_select)
        self.item_list_frame = list_frame

        grid_frame = ttk.Frame(self.item_results, style="Surface.TFrame")
        grid_frame.grid(row=0, column=0, sticky="nsew")
        grid_frame.columnconfigure(0, weight=1)
        grid_frame.rowconfigure(0, weight=1)
        self.item_grid_canvas = tk.Canvas(
            grid_frame, bg=COLORS["surface"], highlightthickness=0, borderwidth=0,
        )
        self.item_grid_canvas.grid(row=0, column=0, sticky="nsew", padx=(8, 0), pady=8)
        grid_scrollbar = ttk.Scrollbar(
            grid_frame, orient="vertical", command=self.item_grid_canvas.yview,
        )
        grid_scrollbar.grid(row=0, column=1, sticky="ns", padx=(0, 8), pady=8)
        self.item_grid_canvas.configure(yscrollcommand=grid_scrollbar.set)
        self.item_grid_inner = tk.Frame(self.item_grid_canvas, bg=COLORS["surface"])
        self.item_grid_window = self.item_grid_canvas.create_window(
            (0, 0), window=self.item_grid_inner, anchor="nw",
        )
        self.item_grid_inner.bind(
            "<Configure>",
            lambda _event: self.item_grid_canvas.configure(
                scrollregion=self.item_grid_canvas.bbox("all")
            ),
        )
        self.item_grid_canvas.bind("<Configure>", self._on_item_grid_resize)
        self.item_grid_cards: dict[int, tk.Frame] = {}
        self.item_grid_frame = grid_frame

        actions = ttk.Frame(self.items_page)
        actions.grid(row=5, column=0, sticky="e", pady=(12, 0))
        ttk.Button(actions, text="Editar selecionado", command=self.edit_selected_item).pack(
            side="left", padx=(0, 8)
        )
        ttk.Button(
            actions, text="Excluir selecionado", style="Danger.TButton",
            command=self.delete_selected_item,
        ).pack(side="left")
        self._refresh_item_category_choices()
        self.set_item_view()

    def _build_users_page(self) -> None:
        self.users_page = ttk.Frame(self.page_container, padding=(34, 28, 34, 26))
        self.users_page.grid(row=0, column=0, sticky="nsew")
        self.users_page.columnconfigure(0, weight=1)
        self.users_page.rowconfigure(5, weight=1)

        topbar = ttk.Frame(self.users_page)
        topbar.grid(row=0, column=0, sticky="ew", pady=(0, 20))
        topbar.columnconfigure(0, weight=1)
        heading = ttk.Frame(topbar)
        heading.grid(row=0, column=0, sticky="w")
        ttk.Label(heading, text="Usuários", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            heading, text="Cadastre usuários ou atualize a lista de servidores por planilha.",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(3, 0))

        cards = ttk.Frame(self.users_page)
        cards.grid(row=1, column=0, sticky="ew", pady=(0, 16))
        for column in range(3):
            cards.columnconfigure(column, weight=1, uniform="user-card")
        self.user_total_value = self._stat_card(cards, 0, "USUÁRIOS", "0")
        self.user_server_value = self._stat_card(cards, 1, "SERVIDORES", "0")
        self.user_other_value = self._stat_card(cards, 2, "EXTERNOS E COLABORADORES", "0")

        form = ttk.Frame(self.users_page, style="Surface.TFrame", padding=(16, 13))
        form.grid(row=2, column=0, sticky="ew", pady=(0, 12))
        for column, weight in enumerate((2, 1, 2, 1)):
            form.columnconfigure(column, weight=weight)
        for column, label in enumerate(("Nome", "Matrícula", "E-mail", "Tipo")):
            ttk.Label(form, text=label, style="CardTitle.TLabel").grid(
                row=0, column=column, sticky="w", padx=(0, 10), pady=(0, 7),
            )
        self.user_name_entry = PlaceholderEntry(
            form, font=("Segoe UI", 10), placeholder="Nome completo",
        )
        self.user_name_entry.grid(row=1, column=0, sticky="ew", padx=(0, 10), ipady=5)
        self.user_registration_entry = PlaceholderEntry(
            form, font=("Segoe UI", 10), placeholder="Ex.: 123456",
        )
        self.user_registration_entry.grid(row=1, column=1, sticky="ew", padx=(0, 10), ipady=5)
        self.user_email_entry = PlaceholderEntry(
            form, font=("Segoe UI", 10), placeholder="nome@instituicao.gov.br",
        )
        self.user_email_entry.grid(row=1, column=2, sticky="ew", padx=(0, 10), ipady=5)
        self.user_type_var = tk.StringVar(value="Servidor")
        self.user_type_combo = ttk.Combobox(
            form, textvariable=self.user_type_var, values=USER_TYPES, state="readonly",
        )
        self.user_type_combo.grid(row=1, column=3, sticky="ew", ipady=4)
        self.user_name_entry.bind("<Return>", lambda _event: self.save_user())
        self.user_registration_entry.bind("<Return>", lambda _event: self.save_user())
        self.user_email_entry.bind("<Return>", lambda _event: self.save_user())
        self.user_save_button = ttk.Button(
            form, text="＋  Adicionar usuário", style="Accent.TButton", command=self.save_user,
        )
        self.user_save_button.grid(row=2, column=3, sticky="e", pady=(10, 0))
        self.user_cancel_button = ttk.Button(
            form, text="Cancelar edição", command=self.clear_user_form,
        )

        suap_help = ttk.Frame(self.users_page, style="Surface.TFrame", padding=(14, 10))
        suap_help.grid(row=3, column=0, sticky="ew", pady=(0, 10))
        suap_help.columnconfigure(0, weight=1)
        ttk.Label(
            suap_help,
            text=(
                "Importar servidores do SUAP IFRO: acesse Gestão de Pessoas > Servidores, "
                "filtre pelo campus e situação Ativo, clique em Pesquisar e depois em Exportar XLS. "
                "Selecione o arquivo baixado no botão de importação abaixo."
            ),
            style="Subtitle.TLabel", wraplength=900, justify="left",
        ).grid(row=0, column=0, sticky="w")
        ttk.Button(
            suap_help, text="Abrir SUAP IFRO", command=self._open_suap,
        ).grid(row=0, column=1, sticky="e", padx=(12, 0))

        toolbar = ttk.Frame(self.users_page)
        toolbar.grid(row=4, column=0, sticky="ew", pady=(0, 10))
        toolbar.columnconfigure(0, weight=1)
        search_frame = ttk.Frame(toolbar, style="Surface.TFrame", padding=(9, 3))
        search_frame.grid(row=0, column=0, sticky="ew", padx=(0, 10))
        search_frame.columnconfigure(1, weight=1)
        ttk.Label(search_frame, text="⌕", style="Subtitle.TLabel").grid(
            row=0, column=0, padx=(2, 7),
        )
        self.user_search_var = tk.StringVar()
        self.user_search_entry = PlaceholderEntry(
            search_frame, textvariable=self.user_search_var, font=("Segoe UI", 9),
            placeholder="Pesquisar por nome, matrícula ou e-mail...",
        )
        self.user_search_entry.grid(row=0, column=1, sticky="ew", ipady=5)
        self.user_search_var.trace_add(
            "write",
            lambda *_args: self.refresh_users()
            if not self.user_search_entry.placeholder_active else None,
        )
        self.user_type_filter_var = tk.StringVar(value="Todos os tipos")
        self.user_type_filter_combo = ttk.Combobox(
            toolbar, textvariable=self.user_type_filter_var,
            values=("Todos os tipos", *USER_TYPES), state="readonly", width=18,
        )
        self.user_type_filter_combo.grid(row=0, column=1, padx=(0, 8))
        self.user_type_filter_combo.bind(
            "<<ComboboxSelected>>", lambda _event: self.refresh_users(),
        )
        ttk.Button(
            toolbar, text="Importar servidores (.xls/.xlsx)",
            command=self.import_servers_from_file,
        ).grid(row=0, column=2, padx=(0, 7))
        ttk.Button(toolbar, text="Excel", command=self.export_users_excel).grid(
            row=0, column=3, padx=(0, 7),
        )
        ttk.Button(toolbar, text="PDF", command=self.export_users_pdf).grid(row=0, column=4)

        table_frame = ttk.Frame(self.users_page, style="Surface.TFrame", padding=(10, 8))
        table_frame.grid(row=5, column=0, sticky="nsew")
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        self.user_table = ttk.Treeview(
            table_frame, columns=("matricula", "name", "email", "type"),
            show="headings", selectmode="browse",
        )
        self.user_table.heading("matricula", text="MATRÍCULA")
        self.user_table.heading("name", text="NOME")
        self.user_table.heading("email", text="E-MAIL")
        self.user_table.heading("type", text="TIPO")
        self.user_table.column("matricula", width=130, minwidth=110, anchor="center", stretch=False)
        self.user_table.column("name", width=250, minwidth=150, anchor="w")
        self.user_table.column("email", width=280, minwidth=160, anchor="w")
        self.user_table.column("type", width=140, minwidth=120, anchor="center", stretch=False)
        self.user_table.grid(row=0, column=0, sticky="nsew")
        user_scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=self.user_table.yview)
        user_scrollbar.grid(row=0, column=1, sticky="ns")
        self.user_table.configure(yscrollcommand=user_scrollbar.set)
        self.user_table.bind("<Delete>", lambda _event: self.delete_selected_user())
        self.user_table.bind("<Double-1>", lambda _event: self.edit_selected_user())
        actions = ttk.Frame(self.users_page)
        actions.grid(row=6, column=0, sticky="e", pady=(10, 0))
        self.user_history_email_button = ttk.Button(
            actions, text="Enviar histórico por e-mail", style="Accent.TButton",
            command=self.email_selected_user_loan_history,
        )
        self.user_history_email_button.pack(side="right", padx=(0, 8))
        ttk.Button(
            actions, text="Editar selecionado", command=self.edit_selected_user,
        ).pack(side="right", padx=(0, 8))
        ttk.Button(
            actions, text="Excluir usuário selecionado", style="Danger.TButton",
            command=self.delete_selected_user,
        ).pack(side="right")

    def _open_suap(self) -> None:
        if not webbrowser.open("https://suap.ifro.edu.br/"):
            messagebox.showerror(
                "Abrir SUAP IFRO",
                "Não foi possível abrir o endereço no navegador padrão.",
                parent=self,
            )

    def _build_loans_page(self) -> None:
        self.loans_page = ttk.Frame(self.page_container, padding=(34, 28, 34, 26))
        self.loans_page.grid(row=0, column=0, sticky="nsew")
        self.loans_page.columnconfigure(0, weight=1)
        self.loans_page.rowconfigure(3, weight=1)

        topbar = ttk.Frame(self.loans_page)
        topbar.grid(row=0, column=0, sticky="ew", pady=(0, 20))
        topbar.columnconfigure(0, weight=1)
        heading = ttk.Frame(topbar)
        heading.grid(row=0, column=0, sticky="w")
        ttk.Label(heading, text="Empréstimos", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            heading, text="Registre retiradas, devoluções e notificações por e-mail.",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(3, 0))
        report_buttons = ttk.Frame(topbar)
        report_buttons.grid(row=0, column=1, sticky="e")
        ttk.Button(report_buttons, text="Extrato Excel", command=self.export_loans_excel).pack(
            side="left", padx=(0, 7),
        )
        ttk.Button(report_buttons, text="Extrato PDF", command=self.export_loans_pdf).pack(side="left")

        form = ttk.Frame(self.loans_page, style="Surface.TFrame", padding=(16, 14))
        form.grid(row=1, column=0, sticky="ew", pady=(0, 12))
        for column, weight in enumerate((1, 2, 2, 2, 0)):
            form.columnconfigure(column, weight=weight)
        ttk.Label(form, text="Categoria", style="CardTitle.TLabel").grid(
            row=0, column=0, sticky="w", padx=(0, 10), pady=(0, 7),
        )
        ttk.Label(form, text="Itens disponíveis (clique para marcar)", style="CardTitle.TLabel").grid(
            row=0, column=1, sticky="w", padx=(0, 10), pady=(0, 7),
        )
        ttk.Label(form, text="Itens selecionados", style="CardTitle.TLabel").grid(
            row=0, column=2, sticky="w", padx=(0, 10), pady=(0, 7),
        )
        ttk.Label(form, text="Usuário", style="CardTitle.TLabel").grid(
            row=0, column=3, sticky="w", padx=(0, 10), pady=(0, 7),
        )
        self.loan_category_var = tk.StringVar(value="Todas as categorias")
        self.loan_category_combo = ttk.Combobox(
            form, textvariable=self.loan_category_var, state="readonly",
        )
        self.loan_category_combo.grid(row=1, column=0, sticky="ew", padx=(0, 10), ipady=4)
        self.loan_category_combo.bind(
            "<<ComboboxSelected>>", lambda _event: self._refresh_loan_item_choices(),
        )
        loan_items_frame = ttk.Frame(form)
        loan_items_frame.grid(row=1, column=1, sticky="ew", padx=(0, 10))
        loan_items_frame.columnconfigure(0, weight=1)
        self.loan_item_search_var = tk.StringVar()
        self.loan_item_search_entry = PlaceholderEntry(
            loan_items_frame,
            textvariable=self.loan_item_search_var,
            placeholder="Pesquisar por número, item ou categoria...",
        )
        self.loan_item_search_entry.grid(row=0, column=0, sticky="ew", pady=(0, 6), ipady=3)
        self.loan_item_search_var.trace_add(
            "write", lambda *_args: self._refresh_loan_item_choices(),
        )
        self.loan_items_table = ttk.Treeview(
            loan_items_frame, columns=("check", "number", "item", "category"), show="headings",
            selectmode="browse", height=4,
        )
        for column, heading in (
            ("check", "✓"), ("number", "Nº"), ("item", "ITEM"), ("category", "CATEGORIA"),
        ):
            self.loan_items_table.heading(column, text=heading)
        self.loan_items_table.column("check", width=34, anchor="center", stretch=False)
        self.loan_items_table.column("number", width=48, anchor="center", stretch=False)
        self.loan_items_table.column("item", width=145, anchor="w")
        self.loan_items_table.column("category", width=100, anchor="w")
        self.loan_items_table.grid(row=1, column=0, sticky="ew")
        loan_items_scrollbar = ttk.Scrollbar(
            loan_items_frame, orient="vertical", command=self.loan_items_table.yview,
        )
        loan_items_scrollbar.grid(row=1, column=1, sticky="ns")
        self.loan_items_table.configure(yscrollcommand=loan_items_scrollbar.set)
        self.loan_items_table.bind("<ButtonRelease-1>", self._toggle_loan_item_selection)

        selected_items_frame = ttk.Frame(form)
        selected_items_frame.grid(row=1, column=2, sticky="ew", padx=(0, 10))
        selected_items_frame.columnconfigure(0, weight=1)
        self.loan_selected_items_table = ttk.Treeview(
            selected_items_frame, columns=("number", "item", "category"),
            show="headings", selectmode="extended", height=4,
        )
        for column, heading in (
            ("number", "Nº"), ("item", "ITEM"), ("category", "CATEGORIA"),
        ):
            self.loan_selected_items_table.heading(column, text=heading)
        self.loan_selected_items_table.column("number", width=48, anchor="center", stretch=False)
        self.loan_selected_items_table.column("item", width=145, anchor="w")
        self.loan_selected_items_table.column("category", width=100, anchor="w")
        self.loan_selected_items_table.grid(row=0, column=0, sticky="ew")
        selected_items_scrollbar = ttk.Scrollbar(
            selected_items_frame, orient="vertical",
            command=self.loan_selected_items_table.yview,
        )
        selected_items_scrollbar.grid(row=0, column=1, sticky="ns")
        self.loan_selected_items_table.configure(
            yscrollcommand=selected_items_scrollbar.set,
        )
        ttk.Button(
            form, text="Remover marcados", command=self.remove_selected_loan_items,
        ).grid(row=2, column=2, sticky="e", pady=(5, 0))

        self.loan_user_var = tk.StringVar()
        self.loan_user_select_button = ttk.Button(
            form, text="Selecionar usuário...", command=self._open_loan_user_picker,
        )
        self.loan_user_select_button.grid(
            row=1, column=3, sticky="ew", padx=(0, 10), ipady=4,
        )
        ttk.Button(
            form, text="Registrar retirada", style="Accent.TButton",
            command=self.create_loan_from_form,
        ).grid(row=1, column=4, sticky="e")
        ttk.Label(
            form,
            text="Alterne as categorias sem perder a seleção; a retirada envia um único e-mail com todos os itens.",
            style="Subtitle.TLabel",
        ).grid(row=3, column=0, columnspan=5, sticky="w", pady=(9, 0))

        self.loans_notebook = ttk.Notebook(self.loans_page)
        self.loans_notebook.grid(row=3, column=0, sticky="nsew")
        active_tab = ttk.Frame(self.loans_notebook, padding=(10, 10))
        history_tab = ttk.Frame(self.loans_notebook, padding=(10, 10))
        email_tab = ttk.Frame(self.loans_notebook, padding=(10, 10))
        self.loans_notebook.add(active_tab, text="Em aberto")
        self.loans_notebook.add(history_tab, text="Histórico")
        self.loans_notebook.add(email_tab, text="Fila de e-mails")
        active_tab.columnconfigure(0, weight=1)
        active_tab.rowconfigure(0, weight=1)
        history_tab.columnconfigure(0, weight=1)
        history_tab.rowconfigure(0, weight=1)
        email_tab.columnconfigure(0, weight=1)
        email_tab.rowconfigure(1, weight=1)

        active_columns = ("item", "category", "user", "matricula", "issued", "email")
        self.active_loans_table = ttk.Treeview(
            active_tab, columns=active_columns, show="headings", selectmode="extended",
        )
        self.active_loans_table.heading("item", text="ITEM")
        self.active_loans_table.heading("category", text="CATEGORIA")
        self.active_loans_table.heading("user", text="USUÁRIO")
        self.active_loans_table.heading("matricula", text="MATRÍCULA")
        self.active_loans_table.heading("issued", text="RETIRADO EM")
        self.active_loans_table.heading("email", text="E-MAIL DE SAÍDA")
        for column, width in {
            "item": 210, "category": 135, "user": 180, "matricula": 105,
            "issued": 145, "email": 130,
        }.items():
            self.active_loans_table.column(column, width=width, minwidth=85, anchor="w")
        self.active_loans_table.column("matricula", anchor="center", stretch=False)
        self.active_loans_table.column("email", anchor="center", stretch=False)
        self.active_loans_table.grid(row=0, column=0, sticky="nsew")
        active_scrollbar = ttk.Scrollbar(
            active_tab, orient="vertical", command=self.active_loans_table.yview,
        )
        active_scrollbar.grid(row=0, column=1, sticky="ns")
        self.active_loans_table.configure(yscrollcommand=active_scrollbar.set)

        history_columns = (
            "item", "category", "user", "matricula", "issued", "returned",
            "issue_email", "return_email", "reminder_email",
        )
        self.loan_history_table = ttk.Treeview(
            history_tab, columns=history_columns, show="headings", selectmode="browse",
        )
        for column, label in {
            "item": "ITEM", "category": "CATEGORIA", "user": "USUÁRIO",
            "matricula": "MATRÍCULA", "issued": "RETIRADO EM", "returned": "DEVOLVIDO EM",
            "issue_email": "E-MAIL SAÍDA", "return_email": "E-MAIL RETORNO",
            "reminder_email": "ÚLTIMA COBRANÇA",
        }.items():
            self.loan_history_table.heading(column, text=label)
        for column, width in {
            "item": 170, "category": 115, "user": 150, "matricula": 100,
            "issued": 135, "returned": 135, "issue_email": 100, "return_email": 110,
            "reminder_email": 110,
        }.items():
            self.loan_history_table.column(column, width=width, minwidth=75, anchor="w")
        self.loan_history_table.grid(row=0, column=0, sticky="nsew")
        history_scrollbar = ttk.Scrollbar(
            history_tab, orient="vertical", command=self.loan_history_table.yview,
        )
        history_scrollbar.grid(row=0, column=1, sticky="ns")
        self.loan_history_table.configure(yscrollcommand=history_scrollbar.set)

        ttk.Label(
            email_tab, text="Status:  ● Enviado     ◷ Na fila     ✕ Falhou     ↻ Enviando",
            style="Subtitle.TLabel",
        ).grid(row=0, column=0, sticky="w", pady=(0, 8))
        self.email_queue_table = ttk.Treeview(
            email_tab,
            columns=("status", "type", "user", "items", "recipient", "created", "attempts"),
            show="headings", selectmode="browse",
        )
        for column, label in {
            "status": "STATUS", "type": "TIPO", "user": "USUÁRIO", "items": "ITENS",
            "recipient": "DESTINATÁRIO", "created": "CRIADO EM", "attempts": "TENTATIVAS",
        }.items():
            self.email_queue_table.heading(column, text=label)
        for column, width in {
            "status": 110, "type": 100, "user": 145, "items": 240,
            "recipient": 180, "created": 145, "attempts": 90,
        }.items():
            self.email_queue_table.column(column, width=width, minwidth=75, anchor="w")
        self.email_queue_table.tag_configure("sent", foreground="#2D6A4F")
        self.email_queue_table.tag_configure("pending", foreground="#9A6816")
        self.email_queue_table.tag_configure("failed", foreground=COLORS["danger"])
        self.email_queue_table.tag_configure("sending", foreground=COLORS["sidebar"])
        self.email_queue_table.grid(row=1, column=0, sticky="nsew")
        self.email_queue_table.bind("<Double-1>", self.show_email_queue_details)
        email_scrollbar = ttk.Scrollbar(
            email_tab, orient="vertical", command=self.email_queue_table.yview,
        )
        email_scrollbar.grid(row=1, column=1, sticky="ns")
        self.email_queue_table.configure(yscrollcommand=email_scrollbar.set)
        email_horizontal_scrollbar = ttk.Scrollbar(
            email_tab, orient="horizontal", command=self.email_queue_table.xview,
        )
        email_horizontal_scrollbar.grid(row=2, column=0, sticky="ew")
        self.email_queue_table.configure(xscrollcommand=email_horizontal_scrollbar.set)
        queue_actions = ttk.Frame(email_tab)
        queue_actions.grid(row=3, column=0, sticky="e", pady=(8, 0))
        ttk.Button(queue_actions, text="Atualizar", command=self.refresh_email_queue).pack(
            side="left", padx=(0, 8),
        )
        ttk.Button(
            queue_actions, text="Reenviar selecionado", style="Accent.TButton",
            command=self.retry_selected_email,
        ).pack(side="left")

        actions = ttk.Frame(self.loans_page)
        actions.grid(row=4, column=0, sticky="e", pady=(10, 0))
        ttk.Button(
            actions, text="Enviar cobrança", command=self.send_selected_loan_reminder,
        ).pack(side="right", padx=(0, 8))
        ttk.Button(
            actions, text="Reenviar e-mails", command=self.retry_selected_loan_emails,
        ).pack(side="right")
        ttk.Button(
            actions, text="Registrar devolução", style="Accent.TButton",
            command=self.return_selected_loan,
        ).pack(side="right", padx=(0, 8))
        self._refresh_loan_category_choices()
        self._refresh_loan_user_choices()

    def _refresh_loan_category_choices(self) -> None:
        if not hasattr(self, "loan_category_combo"):
            return
        categories = self.repository.list_categories()
        self.loan_category_ids = {
            f"{category['name']} · #{category['id']}": category["id"]
            for category in categories
        }
        values = ["Todas as categorias", *self.loan_category_ids]
        self.loan_category_combo.configure(values=values)
        if self.loan_category_var.get() not in values:
            self.loan_category_var.set(values[0])
        self._refresh_loan_item_choices()

    def _refresh_loan_item_choices(self) -> None:
        if not hasattr(self, "loan_items_table"):
            return
        label = self.loan_category_var.get()
        category_id = self.loan_category_ids.get(label)
        items = self.repository.list_available_items(category_id)
        search_term = self.loan_item_search_entry.get().strip().casefold()
        if search_term:
            items = [
                item for item in items
                if search_term in " ".join(
                    str(item[key] or "") for key in ("number", "name", "category_name")
                ).casefold()
            ]
        available_items = self.repository.list_available_items()
        available_item_ids = {int(item["id"]) for item in available_items}
        self.loan_selected_item_ids.intersection_update(available_item_ids)
        for row_id in self.loan_items_table.get_children():
            self.loan_items_table.delete(row_id)
        for item in items:
            row_id = str(item["id"])
            self.loan_items_table.insert(
                "", "end", iid=row_id,
                values=(
                    "☑" if int(item["id"]) in self.loan_selected_item_ids else "☐",
                    item["number"], item["name"], item["category_name"],
                ),
            )
        self._refresh_selected_loan_items(available_items)

    def _refresh_selected_loan_items(
        self, available_items: list[sqlite3.Row] | None = None,
    ) -> None:
        if not hasattr(self, "loan_selected_items_table"):
            return
        if available_items is None:
            available_items = self.repository.list_available_items()
        items_by_id = {int(item["id"]): item for item in available_items}
        self.loan_selected_item_ids.intersection_update(items_by_id)
        for row_id in self.loan_selected_items_table.get_children():
            self.loan_selected_items_table.delete(row_id)
        for item_id in sorted(self.loan_selected_item_ids):
            item = items_by_id[item_id]
            self.loan_selected_items_table.insert(
                "", "end", iid=str(item_id),
                values=(item["number"], item["name"], item["category_name"]),
            )

    def _toggle_loan_item_selection(self, event: tk.Event) -> str:
        row_id = self.loan_items_table.identify_row(event.y)
        if not row_id:
            return "break"
        item_id = int(row_id)
        if item_id in self.loan_selected_item_ids:
            self.loan_selected_item_ids.remove(item_id)
        else:
            self.loan_selected_item_ids.add(item_id)
        for item in self.loan_items_table.get_children():
            if item == row_id:
                values = list(self.loan_items_table.item(item, "values"))
                values[0] = "☑" if item_id in self.loan_selected_item_ids else "☐"
                self.loan_items_table.item(item, values=values)
                break
        self._refresh_selected_loan_items()
        return "break"

    def remove_selected_loan_items(self) -> None:
        selected = self.loan_selected_items_table.selection()
        if not selected:
            messagebox.showinfo(
                "Remover itens", "Selecione os itens que deseja remover da lista.", parent=self,
            )
            return
        self.loan_selected_item_ids.difference_update(int(item_id) for item_id in selected)
        self._refresh_loan_item_choices()

    def _refresh_loan_user_choices(self) -> None:
        if not hasattr(self, "loan_user_select_button"):
            return
        users = self.repository.list_users()
        self.loan_user_ids = {
            f"{user['name']} · {user['matricula']} · {user['user_type']}": user["id"]
            for user in users
        }
        if self.loan_user_var.get() not in self.loan_user_ids:
            self.loan_user_var.set(next(iter(self.loan_user_ids), ""))
        self._update_loan_user_select_button()

    def _update_loan_user_select_button(self) -> None:
        selected = self.loan_user_var.get()
        self.loan_user_select_button.configure(
            text=selected if selected else "Selecionar usuário...",
        )

    def _open_loan_user_picker(self) -> None:
        users = self.repository.list_users()
        if not users:
            messagebox.showinfo(
                "Selecionar usuário",
                "Cadastre ou importe usuários antes de registrar um empréstimo.",
                parent=self,
            )
            return
        dialog = tk.Toplevel(self)
        dialog.title("Selecionar servidor ou usuário")
        dialog.transient(self)
        dialog.geometry("780x560")
        dialog.minsize(650, 440)
        dialog.columnconfigure(0, weight=1)
        dialog.rowconfigure(1, weight=1)
        search_panel = ttk.Frame(dialog)
        search_panel.grid(row=0, column=0, sticky="ew", padx=18, pady=(16, 12))
        search_panel.columnconfigure(0, weight=1)
        ttk.Label(
            search_panel, text="Localize o usuário pelo nome, matrícula ou tipo.",
            style="Subtitle.TLabel",
        ).grid(row=0, column=0, sticky="w")
        search_var = tk.StringVar()
        search_entry = PlaceholderEntry(
            search_panel, textvariable=search_var,
            placeholder="Digite para filtrar usuários...",
        )
        search_entry.grid(row=1, column=0, sticky="ew", pady=(8, 0), ipady=5)

        list_frame = ttk.Frame(dialog)
        list_frame.grid(row=1, column=0, sticky="nsew", padx=18, pady=(0, 10))
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)
        table = ttk.Treeview(
            list_frame, columns=("name", "matricula", "type"), show="headings",
            selectmode="browse",
        )
        table.heading("name", text="NOME")
        table.heading("matricula", text="MATRÍCULA")
        table.heading("type", text="TIPO")
        table.column("name", width=430, minwidth=240, anchor="w")
        table.column("matricula", width=150, minwidth=120, anchor="center")
        table.column("type", width=150, minwidth=120, anchor="center")
        table.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=table.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        table.configure(yscrollcommand=scrollbar.set)
        buttons = ttk.Frame(dialog)
        buttons.grid(row=2, column=0, sticky="e", padx=18, pady=(0, 16))
        users_by_id = {str(user["id"]): user for user in users}

        def refresh_user_rows(*_args) -> None:
            term = "" if search_entry.placeholder_active else search_var.get().strip().casefold()
            previous_selection = table.selection()
            for row_id in table.get_children():
                table.delete(row_id)
            for user in users:
                searchable = " ".join(
                    str(user[key] or "") for key in ("name", "matricula", "email", "user_type")
                ).casefold()
                if term and term not in searchable:
                    continue
                table.insert(
                    "", "end", iid=str(user["id"]),
                    values=(user["name"], user["matricula"], user["user_type"]),
                )
            if previous_selection and table.exists(previous_selection[0]):
                table.selection_set(previous_selection[0])
            elif table.get_children():
                table.selection_set(table.get_children()[0])

        def confirm_selection(_event=None) -> None:
            selection = table.selection()
            if not selection:
                messagebox.showinfo(
                    "Selecionar usuário", "Selecione uma pessoa na lista.", parent=dialog,
                )
                return
            user = users_by_id.get(selection[0])
            if user is None:
                return
            self.loan_user_var.set(
                f"{user['name']} · {user['matricula']} · {user['user_type']}"
            )
            self._update_loan_user_select_button()
            dialog.destroy()

        search_var.trace_add("write", refresh_user_rows)
        table.bind("<Double-1>", confirm_selection)
        ttk.Button(buttons, text="Cancelar", command=dialog.destroy).pack(
            side="right", padx=(8, 0),
        )
        ttk.Button(
            buttons, text="Selecionar", style="Accent.TButton",
            command=confirm_selection,
        ).pack(side="right")
        refresh_user_rows()
        dialog.update_idletasks()
        x = self.winfo_rootx() + max(0, (self.winfo_width() - dialog.winfo_width()) // 2)
        y = self.winfo_rooty() + max(0, (self.winfo_height() - dialog.winfo_height()) // 2)
        dialog.geometry(f"+{x}+{y}")
        dialog.grab_set()
        search_entry.focus_set()
        self.wait_window(dialog)

    def _show_loans(self) -> None:
        self._set_active_navigation("loans")
        self.loans_page.tkraise()
        self._refresh_loan_category_choices()
        self._refresh_loan_user_choices()
        self.refresh_loans()

    @staticmethod
    def _email_status_label(status: str) -> str:
        return {
            "pending": "◷ Na fila", "sending": "↻ Enviando", "sent": "● Enviado",
            "failed": "✕ Falhou", "": "—",
        }.get(status, status or "—")

    @classmethod
    def _outbox_status_label(cls, status: str, recipient: str) -> str:
        if status == "failed" and not recipient.strip():
            return "✕ Sem e-mail"
        return cls._email_status_label(status)

    def refresh_loans(self) -> None:
        if not hasattr(self, "active_loans_table"):
            return
        active_loans = self.repository.list_loans(active_only=True)
        all_loans = self.repository.list_loans()
        for table in (self.active_loans_table, self.loan_history_table):
            for row_id in table.get_children():
                table.delete(row_id)
        for loan in active_loans:
            self.active_loans_table.insert(
                "", "end", iid=str(loan["id"]),
                values=(
                    f"Nº {loan['number']} · {loan['item_name']}", loan["category_name"],
                    loan["user_name"], loan["matricula"], loan["issued_at"],
                    self._email_status_label(loan["issue_email_status"]),
                ),
            )
        for loan in all_loans:
            self.loan_history_table.insert(
                "", "end", iid=str(loan["id"]),
                values=(
                    f"Nº {loan['number']} · {loan['item_name']}", loan["category_name"],
                    loan["user_name"], loan["matricula"], loan["issued_at"],
                    f"✓ {loan['returned_at']}" if loan["returned_at"] else "● Em aberto",
                    self._email_status_label(loan["issue_email_status"]),
                    self._email_status_label(loan["return_email_status"]),
                    self._email_status_label(loan["reminder_email_status"]),
                ),
            )
        self._refresh_loan_item_choices()
        self.refresh_email_queue()
        self.refresh_dashboard()

    def refresh_email_queue(self) -> None:
        if not hasattr(self, "email_queue_table"):
            return
        for row_id in self.email_queue_table.get_children():
            self.email_queue_table.delete(row_id)
        for message in self.repository.list_outbox_messages():
            status = str(message["status"])
            event_type = {
                "issue": "Retirada", "return": "Devolução", "reminder": "Cobrança",
            }.get(str(message["event_type"]), str(message["event_type"]))
            self.email_queue_table.insert(
                "", "end", iid=str(message["id"]),
                values=(
                    self._outbox_status_label(status, str(message["recipient"] or "")),
                    event_type, message["user_name"],
                    message["item_names"] or "—", message["recipient"] or "Sem e-mail",
                    message["sent_at"] or message["created_at"], message["attempt_count"],
                ),
                tags=(status,),
            )

    def show_email_queue_details(self, _event=None) -> None:
        selection = self.email_queue_table.selection()
        if not selection:
            return
        message = self.repository.get_outbox_message(int(selection[0]))
        if message is None:
            self.refresh_email_queue()
            return
        details = (
            f"Status: {self._email_status_label(str(message['status']))}\n"
            f"Destinatário: {message['resolved_recipient'] or 'não informado'}\n"
            f"Assunto: {message['subject']}\n"
            f"Tentativas: {message['attempt_count']}"
        )
        if message["last_error"]:
            details += f"\n\nMotivo da falha:\n{message['last_error']}"
        messagebox.showinfo("Detalhes do e-mail", details, parent=self)

    def create_loan_from_form(self) -> None:
        item_ids = sorted(self.loan_selected_item_ids)
        user_id = self.loan_user_ids.get(self.loan_user_var.get())
        if not item_ids:
            messagebox.showwarning("Itens necessários", "Selecione um ou mais itens disponíveis.", parent=self)
            return
        if user_id is None:
            messagebox.showwarning("Usuário necessário", "Cadastre ou selecione um usuário.", parent=self)
            return
        try:
            loan_ids, outbox_id = self.repository.create_loans(item_ids, user_id)
        except ValueError as error:
            messagebox.showwarning("Não foi possível registrar a retirada", str(error), parent=self)
            self.refresh_loans()
            return
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        self.loan_selected_item_ids.clear()
        self.refresh_loans()
        self.refresh_items()
        self.status_message(
            f"{len(loan_ids)} item(ns) retirado(s); enviando uma confirmação por e-mail."
        )
        self._dispatch_loan_emails(loan_ids[0], [outbox_id])

    def return_selected_loan(self) -> None:
        selection = self.active_loans_table.selection()
        if not selection:
            messagebox.showinfo(
                "Registrar devolução", "Selecione uma retirada na aba Em aberto.", parent=self,
            )
            return
        selected_id = int(selection[0])
        active_loans = self.repository.list_loans(active_only=True)
        selected_loan = next(
            (loan for loan in active_loans if int(loan["id"]) == selected_id),
            None,
        )
        if selected_loan is None:
            self.refresh_loans()
            messagebox.showwarning(
                "Empréstimo encerrado", "Este item já foi devolvido.", parent=self,
            )
            return
        loans_for_user = [
            loan for loan in active_loans
            if int(loan["user_id"]) == int(selected_loan["user_id"])
        ]
        loan_ids = self._choose_user_loan_returns(
            str(selected_loan["user_name"]), loans_for_user, selected_id,
        )
        if not loan_ids:
            return
        try:
            returned_ids, outbox_id = self.repository.return_loans(loan_ids)
        except ValueError as error:
            messagebox.showwarning("Não foi possível registrar a devolução", str(error), parent=self)
            self.refresh_loans()
            return
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        self.refresh_loans()
        self.refresh_items()
        self.status_message(
            f"{len(returned_ids)} item(ns) devolvido(s); enviando uma confirmação por e-mail."
        )
        self._dispatch_loan_emails(returned_ids[0], [outbox_id])

    def _choose_user_loan_returns(
        self,
        user_name: str,
        loans: list[sqlite3.Row],
        selected_loan_id: int,
    ) -> list[int]:
        dialog = tk.Toplevel(self)
        dialog.title(f"Devoluções em aberto — {user_name}")
        dialog.transient(self)
        dialog.geometry("820x480")
        dialog.minsize(680, 380)
        dialog.columnconfigure(0, weight=1)
        dialog.rowconfigure(1, weight=1)
        ttk.Label(
            dialog,
            text=f"{user_name} possui {len(loans)} item(ns) com empréstimo em aberto. "
            "Use Ctrl ou Shift para selecionar vários itens específicos.",
            style="Subtitle.TLabel",
        ).grid(row=0, column=0, sticky="w", padx=18, pady=(16, 10))
        list_frame = ttk.Frame(dialog)
        list_frame.grid(row=1, column=0, sticky="nsew", padx=18, pady=(0, 12))
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)
        table = ttk.Treeview(
            list_frame,
            columns=("number", "item", "category", "issued", "state"),
            show="headings", selectmode="extended",
        )
        for column, heading in (
            ("number", "Nº"), ("item", "ITEM"), ("category", "CATEGORIA"),
            ("issued", "RETIRADO EM"), ("state", "SITUAÇÃO"),
        ):
            table.heading(column, text=heading)
        for column, width in {
            "number": 65, "item": 300, "category": 150, "issued": 150, "state": 110,
        }.items():
            table.column(
                column, width=width,
                anchor="center" if column in ("number", "issued", "state") else "w",
            )
        for loan in loans:
            table.insert(
                "", "end", iid=str(loan["id"]),
                values=(
                    loan["number"], loan["item_name"], loan["category_name"],
                    loan["issued_at"], "● Em aberto",
                ),
            )
        table.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=table.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        table.configure(yscrollcommand=scrollbar.set)
        if table.exists(str(selected_loan_id)):
            table.selection_set(str(selected_loan_id))
            table.focus(str(selected_loan_id))

        buttons = ttk.Frame(dialog)
        buttons.grid(row=2, column=0, sticky="e", padx=18, pady=(0, 16))
        result: list[int] = []

        def return_selected() -> None:
            selection = table.selection()
            if not selection:
                messagebox.showinfo(
                    "Selecionar itens", "Marque um ou mais itens para devolver.", parent=dialog,
                )
                return
            result.extend(int(loan_id) for loan_id in selection)
            dialog.destroy()

        def return_all() -> None:
            if not messagebox.askyesno(
                "Confirmar devolução de todos",
                f"Dar baixa nos {len(loans)} itens em aberto de {user_name} "
                "e enviar uma única mensagem com a listagem?",
                parent=dialog,
            ):
                return
            result.extend(int(loan["id"]) for loan in loans)
            dialog.destroy()

        ttk.Button(buttons, text="Cancelar", command=dialog.destroy).pack(
            side="right", padx=(8, 0),
        )
        ttk.Button(
            buttons, text="Devolver todos", command=return_all,
        ).pack(side="right", padx=(8, 0))
        ttk.Button(
            buttons, text="Devolver selecionados", style="Accent.TButton",
            command=return_selected,
        ).pack(side="right")
        dialog.update_idletasks()
        x = self.winfo_rootx() + max(0, (self.winfo_width() - dialog.winfo_width()) // 2)
        y = self.winfo_rooty() + max(0, (self.winfo_height() - dialog.winfo_height()) // 2)
        dialog.geometry(f"+{x}+{y}")
        dialog.grab_set()
        self.wait_window(dialog)
        return result

    def send_selected_loan_reminder(self) -> None:
        selection = self.active_loans_table.selection()
        if not selection:
            messagebox.showinfo(
                "Enviar cobrança", "Selecione um empréstimo em aberto.", parent=self,
            )
            return
        loan_id = int(selection[0])
        selected_loan = next(
            (loan for loan in self.repository.list_loans(active_only=True) if loan["id"] == loan_id),
            None,
        )
        if selected_loan is None:
            self.refresh_loans()
            messagebox.showwarning("Empréstimo encerrado", "Este item já foi devolvido.", parent=self)
            return
        active_selected = {
            int(row["id"]): row for row in self.repository.list_loans(active_only=True)
        }
        selected_users = {
            int(active_selected[selected_id]["user_id"])
            for selected_id in map(int, selection)
            if selected_id in active_selected
        }
        if len(selected_users) != 1:
            messagebox.showwarning(
                "Seleção inválida", "Selecione empréstimos de um único usuário para cobrar.",
                parent=self,
            )
            return
        active_for_user = [
            loan for loan in active_selected.values()
            if int(loan["user_id"]) == int(selected_loan["user_id"])
        ]
        if not messagebox.askyesno(
            "Confirmar cobrança",
            f"Enviar um único lembrete com os {len(active_for_user)} itens em aberto de "
            f"{selected_loan['user_name']}?",
            parent=self,
        ):
            return
        try:
            outbox_id = self.repository.create_loan_reminder(loan_id)
        except ValueError as error:
            messagebox.showwarning("Não foi possível enviar a cobrança", str(error), parent=self)
            self.refresh_loans()
            return
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        self.refresh_loans()
        self.status_message("Cobrança registrada; enviando e-mail.")
        self._dispatch_loan_emails(loan_id, [outbox_id])

    def retry_selected_loan_emails(self) -> None:
        if self.loans_notebook.index(self.loans_notebook.select()) == 0:
            selection = self.active_loans_table.selection()
        else:
            selection = self.loan_history_table.selection()
        if not selection:
            messagebox.showinfo("Reenviar e-mails", "Selecione um empréstimo no histórico ou em aberto.", parent=self)
            return
        loan_id = int(selection[0])
        outbox_ids = self.repository.list_retryable_outbox_ids(loan_id)
        if not outbox_ids:
            messagebox.showinfo("Reenviar e-mails", "Todas as notificações deste empréstimo já foram enviadas.", parent=self)
            return
        self._dispatch_loan_emails(loan_id, outbox_ids)
        self.status_message("Notificação(ões) enviada(s) para reprocessamento.")

    def retry_selected_email(self) -> None:
        selection = self.email_queue_table.selection()
        if not selection:
            messagebox.showinfo(
                "Reenviar e-mail", "Selecione uma mensagem na fila.", parent=self,
            )
            return
        outbox_id = int(selection[0])
        message = self.repository.get_outbox_message(outbox_id)
        if message is None:
            self.refresh_email_queue()
            messagebox.showwarning("E-mail não encontrado", "A mensagem não existe mais.", parent=self)
            return
        if message["status"] not in ("pending", "failed"):
            messagebox.showinfo(
                "E-mail indisponível",
                "Somente mensagens na fila ou com falha podem ser reenviadas.",
                parent=self,
            )
            return
        self._dispatch_loan_emails(int(message["loan_id"]), [outbox_id])
        self.status_message("E-mail enviado para reprocessamento.")

    def _dispatch_loan_emails(self, loan_id: int, outbox_ids: list[int]) -> None:
        for outbox_id in outbox_ids:
            worker = threading.Thread(
                target=self._deliver_loan_email, args=(loan_id, outbox_id), daemon=True,
            )
            worker.start()

    def _deliver_loan_email(self, loan_id: int, outbox_id: int) -> None:
        error_message = ""
        try:
            if not self.repository.claim_outbox_message(outbox_id):
                return
            outbox = self.repository.get_outbox_message(outbox_id)
            if outbox is None:
                raise ValueError("A notificação não existe mais.")
            settings = self.repository.get_email_settings()
            recipient = str(outbox["resolved_recipient"] or "").strip()
            if not recipient:
                raise ValueError("O usuário não possui e-mail cadastrado.")
            sender_email = str(settings["email"] or "").strip()
            sender_password = str(settings["app_password"] or "")
            if not sender_email or not sender_password:
                raise ValueError("Configure o e-mail de envio e a senha de app nas Configurações.")
            message = EmailMessage()
            sender_name = str(settings["sender_name"] or settings["institution"] or "Empréstimos")
            message["From"] = formataddr((sender_name, sender_email))
            message["To"] = recipient
            message["Subject"] = str(outbox["subject"])
            message.set_content(str(outbox["body"]))
            if outbox["body_html"]:
                message.add_alternative(str(outbox["body_html"]), subtype="html")
            self._send_email_message(message, settings)
            self.repository.mark_outbox_sent(outbox_id, recipient)
        except Exception as error:
            error_message = str(error)
            try:
                self.repository.mark_outbox_failed(outbox_id, error_message)
            except sqlite3.Error as status_error:
                error_message = (
                    f"{error_message}\nNão foi possível registrar a falha na fila: {status_error}"
                )
        try:
            self.after(0, lambda: self._finish_loan_email_delivery(loan_id, error_message))
        except tk.TclError:
            pass

    def _finish_loan_email_delivery(self, loan_id: int, error_message: str) -> None:
        self.refresh_loans()
        self.refresh_email_queue()
        self.refresh_dashboard()
        if error_message:
            messagebox.showwarning(
                "E-mail não enviado",
                f"A movimentação foi registrada, mas o e-mail falhou e foi mantido na fila. "
                f"Abra a aba Fila de e-mails para consultar o motivo ou tentar novamente.\n\n{error_message}",
                parent=self,
            )
        else:
            self.status_message("Notificação enviada por e-mail.")

    @staticmethod
    def _send_email_message(
        message: EmailMessage, settings: dict[str, object],
    ) -> None:
        sender_email = str(settings["email"] or "").strip()
        sender_password = str(settings["app_password"] or "")
        host = str(settings["smtp_host"] or DEFAULT_SMTP_HOST)
        port = int(settings["smtp_port"] or 587)
        if settings["use_tls"] and port == 465:
            with smtplib.SMTP_SSL(host, port, timeout=20) as server:
                server.login(sender_email, sender_password)
                server.send_message(message)
        else:
            with smtplib.SMTP(host, port, timeout=20) as server:
                if settings["use_tls"]:
                    server.starttls(context=ssl.create_default_context())
                server.login(sender_email, sender_password)
                server.send_message(message)

    def email_selected_user_loan_history(self) -> None:
        if self.history_email_in_progress:
            messagebox.showinfo(
                "Histórico em envio",
                "Aguarde a conclusão do envio iniciado.",
                parent=self,
            )
            return
        selection = self.user_table.selection()
        if not selection:
            messagebox.showinfo(
                "Histórico de empréstimos",
                "Selecione uma pessoa na lista de usuários.",
                parent=self,
            )
            return
        user_id = int(selection[0])
        user = self.repository.get_user(user_id)
        if user is None:
            messagebox.showerror(
                "Histórico de empréstimos",
                "O usuário selecionado não foi encontrado. Atualize a lista e tente novamente.",
                parent=self,
            )
            self.refresh_users()
            return
        recipient = str(user["email"] or "").strip()
        if not recipient:
            messagebox.showwarning(
                "E-mail não cadastrado",
                f'{user["name"]} não possui um endereço de e-mail cadastrado.',
                parent=self,
            )
            return
        loans = self.repository.list_loans(user_id=user_id)
        if not loans:
            messagebox.showinfo(
                "Histórico de empréstimos",
                f'Não há empréstimos registrados para {user["name"]}.',
                parent=self,
            )
            return
        user_data = dict(user)
        worker = threading.Thread(
            target=self._deliver_user_loan_history_email,
            args=(user_data, recipient, len(loans), loans),
            daemon=True,
        )
        self.history_email_in_progress = True
        self.user_history_email_button.configure(state="disabled")
        self._show_history_email_progress(str(user["name"]))
        try:
            worker.start()
        except RuntimeError as error:
            self.history_email_in_progress = False
            self.user_history_email_button.configure(state="normal")
            self._close_history_email_progress()
            messagebox.showerror("Falha ao iniciar envio", str(error), parent=self)
            return
        self.status_message(f"Enviando histórico de {user['name']} por e-mail.")

    def _show_history_email_progress(self, user_name: str) -> None:
        dialog = tk.Toplevel(self)
        dialog.title("Processando histórico")
        dialog.transient(self)
        dialog.resizable(False, False)
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)
        panel = ttk.Frame(dialog, padding=24)
        panel.pack(fill="both", expand=True)
        ttk.Label(
            panel, text="Preparando e enviando o histórico",
            style="CardTitle.TLabel",
        ).pack(anchor="w")
        ttk.Label(
            panel,
            text=f"Gerando os documentos e enviando-os para {user_name}.\n"
            "Aguarde a confirmação antes de iniciar outra operação.",
            style="Subtitle.TLabel", wraplength=420, justify="left",
        ).pack(anchor="w", pady=(8, 16))
        progress = ttk.Progressbar(panel, mode="indeterminate", length=420)
        progress.pack(fill="x")
        progress.start(12)
        dialog.update_idletasks()
        x = self.winfo_rootx() + max(0, (self.winfo_width() - dialog.winfo_width()) // 2)
        y = self.winfo_rooty() + max(0, (self.winfo_height() - dialog.winfo_height()) // 2)
        dialog.geometry(f"+{x}+{y}")
        dialog.grab_set()
        self.history_progress_dialog = dialog
        self.history_progress_bar = progress

    def _close_history_email_progress(self) -> None:
        if self.history_progress_dialog is None:
            return
        if self.history_progress_bar is not None:
            self.history_progress_bar.stop()
            self.history_progress_bar = None
        self.history_progress_dialog.grab_release()
        self.history_progress_dialog.destroy()
        self.history_progress_dialog = None

    def _deliver_user_loan_history_email(
        self,
        user: dict[str, object],
        recipient: str,
        loan_count: int,
        loans: list[sqlite3.Row],
    ) -> None:
        error_message = ""
        try:
            excel_data = self._build_loan_excel_bytes(loans)
            pdf_data = self._build_loan_pdf_bytes(
                loans,
                title="Histórico de empréstimos",
                subtitle=f'{user["name"]} · Matrícula {user["matricula"]}',
            )
            settings = self.repository.get_email_settings()
            sender_email = str(settings["email"] or "").strip()
            sender_password = str(settings["app_password"] or "")
            if not sender_email or not sender_password:
                raise ValueError(
                    "Configure o e-mail de envio e a senha de app nas Configurações."
                )
            sender_name = str(
                settings["sender_name"] or settings["institution"] or "Empréstimos"
            )
            user_name = str(user["name"])
            safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", user_name).strip("_") or "usuario"
            message = EmailMessage()
            message["From"] = formataddr((sender_name, sender_email))
            message["To"] = recipient
            message["Subject"] = f"Histórico de empréstimos - {user_name}"
            message.set_content(
                f"Olá, {user_name}.\n\n"
                f"Segue seu histórico completo de empréstimos ({loan_count} registro(s)) "
                "em Excel e PDF.\n\n"
                f'{settings["signature"] or DEFAULT_EMAIL_SIGNATURE}'
            )
            message.add_attachment(
                excel_data,
                maintype="application",
                subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                filename=f"historico_emprestimos_{safe_name}.xlsx",
            )
            message.add_attachment(
                pdf_data,
                maintype="application",
                subtype="pdf",
                filename=f"historico_emprestimos_{safe_name}.pdf",
            )
            self._send_email_message(message, settings)
        except Exception as error:
            error_message = str(error)
        try:
            self.after(
                0,
                lambda: self._finish_user_loan_history_email(
                    str(user["name"]), recipient, error_message,
                ),
            )
        except tk.TclError:
            pass

    def _finish_user_loan_history_email(
        self, user_name: str, recipient: str, error_message: str,
    ) -> None:
        self._close_history_email_progress()
        self.history_email_in_progress = False
        self.user_history_email_button.configure(state="normal")
        if error_message:
            messagebox.showerror(
                "Falha ao enviar histórico",
                f"Não foi possível enviar os arquivos para {recipient}.\n\n{error_message}",
                parent=self,
            )
            return
        messagebox.showinfo(
            "Histórico enviado",
            f"Os relatórios Excel e PDF do histórico de {user_name} foram enviados para {recipient}.",
            parent=self,
        )

    def _build_loan_excel_bytes(self, loans: list[sqlite3.Row]) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Empréstimos"
        sheet.append([
            "Item", "Número", "Categoria", "Usuário", "Matrícula", "E-mail",
            "Retirada", "Devolução", "Situação", "E-mail de saída",
            "E-mail de retorno", "Última cobrança",
        ])
        for loan in loans:
            sheet.append([
                loan["item_name"], loan["number"], loan["category_name"],
                loan["user_name"], loan["matricula"], loan["email"],
                loan["issued_at"], loan["returned_at"] or "",
                "Devolvido" if loan["returned_at"] else "Em aberto",
                self._email_status_label(loan["issue_email_status"]),
                self._email_status_label(loan["return_email_status"]),
                self._email_status_label(loan["reminder_email_status"]),
            ])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183B3A")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for column, width in {
            "A": 36, "B": 12, "C": 26, "D": 34, "E": 16, "F": 32,
            "G": 22, "H": 22, "I": 16, "J": 18, "K": 18, "L": 18,
        }.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _build_loan_pdf_bytes(
        self,
        loans: list[sqlite3.Row],
        title: str = "Extrato de empréstimos",
        subtitle: str | None = None,
    ) -> bytes:
        styles = getSampleStyleSheet()
        styles["Title"].textColor = colors.HexColor(COLORS["sidebar"])
        styles["Title"].fontName = "Helvetica-Bold"
        body_style = styles["BodyText"]
        body_style.wordWrap = "CJK"
        date_style = copy(body_style)
        date_style.fontSize = 7
        date_style.leading = 8
        date_style.alignment = 1
        data = [[
            "Nº", "Item", "Categoria", "Usuário", "Matrícula",
            "Retirada", "Devolução", "Situação",
        ]]
        data.extend([
            [
                str(loan["number"]),
                Paragraph(escape(str(loan["item_name"])), body_style),
                Paragraph(escape(str(loan["category_name"])), body_style),
                Paragraph(escape(str(loan["user_name"])), body_style),
                Paragraph(escape(str(loan["matricula"])), body_style),
                self._pdf_datetime_cell(loan["issued_at"], date_style),
                self._pdf_datetime_cell(loan["returned_at"], date_style),
                "Devolvido" if loan["returned_at"] else "Em aberto",
            ]
            for loan in loans
        ])
        table = Table(
            data,
            colWidths=[
                12 * mm, 68 * mm, 38 * mm, 44 * mm,
                24 * mm, 30 * mm, 30 * mm, 22 * mm,
            ],
            repeatRows=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COLORS["sidebar"])),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F0")]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(COLORS["border"])),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        output = io.BytesIO()
        document = SimpleDocTemplate(
            output, pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm,
            topMargin=15 * mm, bottomMargin=15 * mm,
        )
        story = [
            Paragraph(escape(title), styles["Title"]),
            Paragraph(
                escape(subtitle) if subtitle else
                f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} · "
                f"{len(loans)} registro(s)",
                styles["Normal"],
            ),
            Spacer(1, 6 * mm),
            table,
        ]
        document.build(story)
        return output.getvalue()

    @staticmethod
    def _pdf_datetime_cell(value: object, style) -> Paragraph:
        if not value:
            return Paragraph("—", style)
        timestamp = str(value)
        try:
            timestamp = datetime.strptime(
                timestamp, "%Y-%m-%d %H:%M:%S",
            ).strftime("%d/%m/%Y<br/>%H:%M")
        except ValueError:
            timestamp = escape(timestamp).replace(" ", "<br/>", 1)
        return Paragraph(timestamp, style)

    def export_loans_excel(self) -> None:
        loans = self.repository.list_loans()
        if not loans:
            messagebox.showinfo("Extrato Excel", "Não há empréstimos para exportar.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Salvar extrato de empréstimos", defaultextension=".xlsx",
            filetypes=[("Planilha Excel", "*.xlsx")], initialfile="extrato_emprestimos.xlsx",
        )
        if not path:
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Empréstimos"
        sheet.append([
            "Item", "Número", "Categoria", "Usuário", "Matrícula", "E-mail",
            "Retirada", "Devolução", "Situação", "E-mail de saída",
            "E-mail de retorno", "Última cobrança",
        ])
        for loan in loans:
            sheet.append([
                loan["item_name"], loan["number"], loan["category_name"],
                loan["user_name"], loan["matricula"], loan["email"],
                loan["issued_at"], loan["returned_at"] or "",
                "Devolvido" if loan["returned_at"] else "Em aberto",
                self._email_status_label(loan["issue_email_status"]),
                self._email_status_label(loan["return_email_status"]),
                self._email_status_label(loan["reminder_email_status"]),
            ])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183B3A")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for column, width in {
            "A": 28, "B": 12, "C": 22, "D": 28, "E": 16, "F": 32,
            "G": 22, "H": 22, "I": 16, "J": 18, "K": 18, "L": 18,
        }.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        try:
            workbook.save(path)
        except OSError as error:
            messagebox.showerror("Falha ao exportar", str(error), parent=self)
            return
        self.status_message(f"Extrato salvo em {path}")

    def export_loans_pdf(self) -> None:
        loans = self.repository.list_loans()
        if not loans:
            messagebox.showinfo("Extrato PDF", "Não há empréstimos para exportar.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Salvar extrato de empréstimos", defaultextension=".pdf",
            filetypes=[("Documento PDF", "*.pdf")], initialfile="extrato_emprestimos.pdf",
        )
        if not path:
            return
        styles = getSampleStyleSheet()
        styles["Title"].textColor = colors.HexColor(COLORS["sidebar"])
        styles["Title"].fontName = "Helvetica-Bold"
        body_style = styles["BodyText"]
        body_style.wordWrap = "CJK"
        date_style = copy(body_style)
        date_style.fontSize = 7
        date_style.leading = 8
        date_style.alignment = 1
        data = [["Nº", "Item", "Categoria", "Usuário", "Matrícula", "Retirada", "Devolução", "Situação"]]
        data.extend([
            [
                str(loan["number"]),
                Paragraph(escape(loan["item_name"]), body_style),
                Paragraph(escape(loan["category_name"]), body_style),
                Paragraph(escape(loan["user_name"]), body_style),
                Paragraph(escape(loan["matricula"]), body_style),
                self._pdf_datetime_cell(loan["issued_at"], date_style),
                self._pdf_datetime_cell(loan["returned_at"], date_style),
                "Devolvido" if loan["returned_at"] else "Em aberto",
            ]
            for loan in loans
        ])
        table = Table(
            data,
            colWidths=[
                12 * mm, 68 * mm, 38 * mm, 44 * mm,
                24 * mm, 30 * mm, 30 * mm, 22 * mm,
            ],
            repeatRows=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COLORS["sidebar"])),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F0")]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(COLORS["border"])),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        document = SimpleDocTemplate(
            path, pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm,
            topMargin=15 * mm, bottomMargin=15 * mm,
        )
        story = [
            Paragraph("Extrato de empréstimos", styles["Title"]),
            Paragraph(
                f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} · "
                f"{len(loans)} registro(s)", styles["Normal"],
            ),
            Spacer(1, 6 * mm),
            table,
        ]
        try:
            document.build(story)
        except (OSError, ValueError) as error:
            messagebox.showerror("Falha ao exportar", str(error), parent=self)
            return
        self.status_message(f"Extrato salvo em {path}")

    def _refresh_item_category_choices(self) -> None:
        if not hasattr(self, "item_category_combo"):
            return
        categories = self.repository.list_categories()
        self.item_category_ids = {row["name"]: row["id"] for row in categories}
        category_names = list(self.item_category_ids)
        self.item_category_combo.configure(values=category_names)
        if self.item_category_var.get() not in self.item_category_ids:
            self.item_category_var.set(category_names[0] if category_names else "")
        self.item_filter_ids = {
            f"{row['name']} · #{row['id']}": row["id"] for row in categories
        }
        filter_values = ["Todas as categorias", *self.item_filter_ids]
        self.item_filter_combo.configure(values=filter_values)
        if self.item_filter_var.get() not in filter_values:
            self.item_filter_var.set(filter_values[0])
        self.item_category_count_value.configure(text=str(len(categories)))

    def _selected_item_category_id(self) -> int | None:
        return self.item_filter_ids.get(self.item_filter_var.get())

    def refresh_items(self) -> None:
        if not hasattr(self, "item_table"):
            return
        rows = self.repository.list_items(
            self.item_search_entry.get(), self._selected_item_category_id(),
        )
        self.visible_item_rows = rows
        visible_ids = {row["id"] for row in rows}
        if self.selected_item_id not in visible_ids:
            self.selected_item_id = None
        for item in self.item_table.get_children():
            self.item_table.delete(item)
        for row in rows:
            tag = f"item-status-{row['status']}"
            status_background = {
                "Disponível": "#E9F3EC",
                "Emprestado": "#FBF1DE",
                "Em manutenção": "#F8E8E7",
            }.get(row["status"], COLORS["surface"])
            self.item_table.tag_configure(tag, background=status_background, foreground=COLORS["text"])
            self.item_table.insert(
                "", "end", iid=str(row["id"]),
                values=(
                    row["number"], row["name"], f"●  {row['category_name']}",
                    self._item_status_icon(str(row["status"])),
                ),
                tags=(tag,),
            )
        self.item_total_value.configure(text=str(self.repository.count_items()))
        self.item_visible_value.configure(text=str(len(rows)))
        if self.selected_item_id is not None and self.item_table.exists(str(self.selected_item_id)):
            self.item_table.selection_set(str(self.selected_item_id))
        self._render_item_grid(rows)
        if hasattr(self, "loan_items_table"):
            self._refresh_loan_item_choices()
        self.refresh_dashboard()

    def _render_item_grid(self, rows: list[sqlite3.Row]) -> None:
        if not hasattr(self, "item_grid_inner"):
            return
        for child in self.item_grid_inner.winfo_children():
            child.destroy()
        self.item_grid_cards = {}
        width = max(self.item_grid_canvas.winfo_width(), 240)
        columns = max(1, width // 238)
        for column in range(columns):
            self.item_grid_inner.grid_columnconfigure(column, weight=1, uniform="item-tile")
        for index, item in enumerate(rows):
            item_id = int(item["id"])
            card_background = COLORS["surface"]
            card = tk.Frame(
                self.item_grid_inner, bg=card_background, highlightthickness=1,
                highlightbackground=COLORS["accent"] if item_id == self.selected_item_id else COLORS["border"],
                padx=14, pady=12,
            )
            card.grid(
                row=index // columns, column=index % columns, sticky="nsew",
                padx=7, pady=7,
            )
            card.grid_columnconfigure(1, weight=1)
            stripe = tk.Frame(card, bg=COLORS["accent"], width=5)
            stripe.grid(row=0, column=0, rowspan=4, sticky="ns", padx=(0, 10))
            status_color = {
                "Disponível": "#2D6A4F",
                "Emprestado": "#A36B16",
                "Em manutenção": COLORS["danger"],
            }.get(item["status"], COLORS["muted"])
            labels = [
                tk.Label(
                    card, text=f"Nº {item['number']}", bg=card_background,
                    fg=COLORS["accent"], font=("Segoe UI Semibold", 10), anchor="w",
                ),
                tk.Label(
                    card, text=item["name"], bg=card_background, fg=COLORS["text"],
                    font=("Segoe UI Semibold", 12), anchor="w", justify="left",
                    wraplength=190,
                ),
                tk.Label(
                    card, text=item["category_name"], bg=card_background,
                    fg=COLORS["muted"], font=("Segoe UI", 9), anchor="w",
                    wraplength=190,
                ),
                tk.Label(
                    card, text=item["status"], bg=card_background,
                    fg=status_color, font=("Segoe UI Semibold", 9), anchor="w",
                ),
            ]
            labels[0].grid(row=0, column=1, sticky="ew", pady=(0, 8))
            labels[1].grid(row=1, column=1, sticky="ew", pady=(0, 6))
            labels[2].grid(row=2, column=1, sticky="ew")
            labels[3].grid(row=3, column=1, sticky="ew", pady=(8, 0))
            self.item_grid_cards[item_id] = card
            for widget in (card, stripe, *labels):
                widget.bind(
                    "<Button-1>",
                    lambda _event, selected_id=item_id: self._select_grid_item(selected_id),
                )
        self.item_grid_canvas.configure(scrollregion=self.item_grid_canvas.bbox("all"))

    def _on_item_grid_resize(self, event) -> None:
        self.item_grid_canvas.itemconfigure(self.item_grid_window, width=event.width)
        if hasattr(self, "visible_item_rows"):
            self.after_idle(lambda: self._render_item_grid(self.visible_item_rows))

    def _select_grid_item(self, item_id: int) -> None:
        self.selected_item_id = item_id
        for card_id, card in self.item_grid_cards.items():
            card.configure(
                highlightbackground=COLORS["accent"] if card_id == item_id else COLORS["border"]
            )

    def _on_item_table_select(self, _event=None) -> None:
        selection = self.item_table.selection()
        self.selected_item_id = int(selection[0]) if selection else None

    def set_item_view(self) -> None:
        if not hasattr(self, "item_list_frame"):
            return
        if self.item_view_var.get() == "grid":
            self.item_grid_frame.tkraise()
        else:
            self.item_list_frame.tkraise()

    def _show_items(self) -> None:
        self._set_active_navigation("items")
        self.items_page.tkraise()
        self._refresh_item_category_choices()
        self.refresh_items()

    def _show_users(self) -> None:
        self._set_active_navigation("users")
        self.users_page.tkraise()
        self.refresh_users()
        self.user_name_entry.focus_set()

    def refresh_users(self) -> None:
        if not hasattr(self, "user_table"):
            return
        user_type = self.user_type_filter_var.get()
        if user_type == "Todos os tipos":
            user_type = None
        rows = self.repository.list_users(self.user_search_entry.get(), user_type)
        visible_ids = {row["id"] for row in rows}
        if self.selected_user_id not in visible_ids:
            self.selected_user_id = None
        for item_id in self.user_table.get_children():
            self.user_table.delete(item_id)
        for user in rows:
            self.user_table.insert(
                "", "end", iid=str(user["id"]),
                values=(user["matricula"], user["name"], user["email"], user["user_type"]),
            )
        if self.selected_user_id is not None:
            self.user_table.selection_set(str(self.selected_user_id))
        self.user_total_value.configure(text=str(self.repository.count_users()))
        self.user_server_value.configure(text=str(self.repository.count_users("Servidor")))
        self.user_other_value.configure(
            text=str(self.repository.count_users("Externo") + self.repository.count_users("Colaborador")),
        )
        if hasattr(self, "loan_user_select_button"):
            self._refresh_loan_user_choices()
        self.refresh_dashboard()

    def save_user(self) -> None:
        try:
            if self.selected_user_id is None:
                self.repository.create_user(
                    self.user_name_entry.get(), self.user_registration_entry.get(),
                    self.user_email_entry.get(), self.user_type_var.get(),
                )
                status_message = "Usuário adicionado."
            else:
                self.repository.update_user(
                    self.selected_user_id, self.user_name_entry.get(),
                    self.user_registration_entry.get(), self.user_email_entry.get(),
                    self.user_type_var.get(),
                )
                status_message = "Usuário atualizado."
        except ValueError as error:
            messagebox.showwarning("Verifique os dados", str(error), parent=self)
            return
        except sqlite3.IntegrityError:
            messagebox.showwarning(
                "Matrícula já cadastrada", "Já existe um usuário com essa matrícula.", parent=self,
            )
            return
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        self.clear_user_form()
        self.refresh_users()
        self.status_message(status_message)

    def edit_selected_user(self) -> None:
        selection = self.user_table.selection()
        if not selection:
            messagebox.showinfo("Editar usuário", "Selecione um usuário na lista.", parent=self)
            return
        user_id = int(selection[0])
        user = self.repository.get_user(user_id)
        if user is None:
            self.refresh_users()
            return
        self.selected_user_id = user_id
        self.user_name_entry.delete(0, "end")
        self.user_name_entry.insert(0, user["name"])
        self.user_registration_entry.delete(0, "end")
        self.user_registration_entry.insert(0, user["matricula"])
        self.user_email_entry.delete(0, "end")
        self.user_email_entry.insert(0, user["email"])
        self.user_type_var.set(user["user_type"])
        self.user_save_button.configure(text="Salvar alterações")
        self.user_cancel_button.grid(row=2, column=2, sticky="e", pady=(10, 0), padx=(0, 8))
        self.user_name_entry.focus_set()

    def clear_user_form(self) -> None:
        self.selected_user_id = None
        self.user_name_entry.delete(0, "end")
        self.user_registration_entry.delete(0, "end")
        self.user_email_entry.delete(0, "end")
        self.user_type_var.set("Servidor")
        self.user_save_button.configure(text="＋  Adicionar usuário")
        self.user_cancel_button.grid_forget()
        if hasattr(self, "user_table"):
            self.user_table.selection_remove(self.user_table.selection())
        self.user_name_entry.focus_set()

    def delete_selected_user(self) -> None:
        selection = self.user_table.selection()
        if not selection:
            messagebox.showinfo("Excluir usuário", "Selecione um usuário na lista.", parent=self)
            return
        user_id = int(selection[0])
        values = self.user_table.item(selection[0], "values")
        if not messagebox.askyesno(
            "Confirmar exclusão", f'Deseja excluir "{values[1]}" ({values[0]})?', parent=self,
        ):
            return
        try:
            self.repository.delete_user(user_id)
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        if self.selected_user_id == user_id:
            self.clear_user_form()
        self.refresh_users()
        self.status_message("Usuário excluído.")

    def import_servers_from_file(self) -> None:
        selected_path = filedialog.askopenfilename(
            parent=self, title="Selecionar planilha de servidores",
            filetypes=[("Planilhas Excel", "*.xlsx *.xls"), ("Todos os arquivos", "*.*")],
        )
        if not selected_path:
            return
        try:
            records = read_server_spreadsheet(Path(selected_path))
        except Exception as error:
            messagebox.showerror("Planilha inválida", str(error), parent=self)
            return
        if not messagebox.askyesno(
            "Atualizar servidores",
            f"A planilha contém {len(records)} registros. Os servidores serão sincronizados "
            "pela matrícula; quando ela mudar, o sistema tentará localizar a mesma pessoa por "
            "e-mail e depois por nome único. Os vínculos de empréstimos e e-mails serão "
            "preservados. A posição/ordem da linha não é usada como identificação. Deseja continuar?",
            parent=self,
        ):
            return
        try:
            result = self.repository.sync_servers(records)
        except ValueError as error:
            messagebox.showwarning("Importação cancelada", str(error), parent=self)
            return
        except sqlite3.IntegrityError:
            messagebox.showerror(
                "Matrícula em conflito",
                "Uma matrícula da planilha já pertence a um usuário Externo ou Colaborador. "
                "Nenhum servidor atual foi removido.",
                parent=self,
            )
            return
        except sqlite3.Error as error:
            messagebox.showerror("Falha na importação", str(error), parent=self)
            return
        if self.selected_user_id is not None:
            edited_user = self.repository.get_user(self.selected_user_id)
            if edited_user is None:
                self.clear_user_form()
        self.refresh_users()
        self.status_message(
            f'{result["processed"]} registros sincronizados; '
            f'{result["added"]} novos e {result["updated"]} atualizados.'
        )
        messagebox.showinfo(
            "Importação concluída",
            f'{result["processed"]} registros processados.\n'
            f'{result["updated"]} servidores atualizados e vínculos preservados.\n'
            f'{result["added"]} novos servidores adicionados.\n'
            f'{result["retained_with_history"]} servidores ausentes na planilha mantidos '
            "por possuírem empréstimos ou e-mails no histórico.\n"
            f'{result["removed"]} servidores sem histórico removidos.',
            parent=self,
        )

    def _filtered_users_for_export(self) -> list[sqlite3.Row]:
        selected_type = self.user_type_filter_var.get()
        user_type = None if selected_type == "Todos os tipos" else selected_type
        return self.repository.list_users(self.user_search_entry.get(), user_type)

    def export_users_excel(self) -> None:
        users = self._filtered_users_for_export()
        if not users:
            messagebox.showinfo("Relatório Excel", "Não há usuários para exportar.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Salvar relatório de usuários", defaultextension=".xlsx",
            filetypes=[("Planilha Excel", "*.xlsx")], initialfile="relatorio_usuarios.xlsx",
        )
        if not path:
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Usuários"
        sheet.append(["Matrícula", "Nome", "E-mail", "Tipo"])
        for user in users:
            sheet.append([user["matricula"], user["name"], user["email"], user["user_type"]])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183B3A")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 24
        for column, width in {"A": 18, "B": 34, "C": 38, "D": 20}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        try:
            workbook.save(path)
        except OSError as error:
            messagebox.showerror("Falha ao exportar", str(error), parent=self)
            return
        self.status_message(f"Relatório de usuários salvo em {path}")

    def export_users_pdf(self) -> None:
        users = self._filtered_users_for_export()
        if not users:
            messagebox.showinfo("Relatório PDF", "Não há usuários para exportar.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Salvar relatório de usuários", defaultextension=".pdf",
            filetypes=[("Documento PDF", "*.pdf")], initialfile="relatorio_usuarios.pdf",
        )
        if not path:
            return
        styles = getSampleStyleSheet()
        styles["Title"].textColor = colors.HexColor(COLORS["sidebar"])
        styles["Title"].fontName = "Helvetica-Bold"
        body_style = styles["BodyText"]
        data = [["Matrícula", "Nome", "E-mail", "Tipo"]]
        data.extend([
            [
                Paragraph(escape(user["matricula"]), body_style),
                Paragraph(escape(user["name"]), body_style),
                Paragraph(escape(user["email"] or ""), body_style),
                user["user_type"],
            ]
            for user in users
        ])
        table = Table(data, colWidths=[35 * mm, 75 * mm, 100 * mm, 42 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COLORS["sidebar"])),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F0")]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(COLORS["border"])),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        document = SimpleDocTemplate(
            path, pagesize=landscape(A4), rightMargin=20 * mm, leftMargin=20 * mm,
            topMargin=20 * mm, bottomMargin=20 * mm,
        )
        story = [
            Paragraph("Relatório de usuários", styles["Title"]),
            Paragraph(
                f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} · "
                f"{len(users)} usuário(s)", styles["Normal"],
            ),
            Spacer(1, 8 * mm),
            table,
        ]
        try:
            document.build(story)
        except (OSError, ValueError) as error:
            messagebox.showerror("Falha ao exportar", str(error), parent=self)
            return
        self.status_message(f"Relatório de usuários salvo em {path}")

    def save_item(self) -> None:
        category_id = self.item_category_ids.get(self.item_category_var.get())
        if category_id is None:
            messagebox.showwarning(
                "Categoria necessária", "Cadastre uma categoria antes de adicionar itens.", parent=self,
            )
            return
        try:
            if self.selected_item_id is None:
                self.repository.create_item(
                    self.item_number_entry.get(), self.item_name_entry.get(), category_id,
                    self.item_status_var.get(),
                )
                message = "Item cadastrado com sucesso."
            else:
                self.repository.update_item(
                    self.selected_item_id, self.item_number_entry.get(),
                    self.item_name_entry.get(), category_id, self.item_status_var.get(),
                )
                message = "Item atualizado com sucesso."
        except ValueError as error:
            messagebox.showwarning("Verifique os dados", str(error), parent=self)
            return
        except sqlite3.IntegrityError:
            messagebox.showwarning(
                "Número já utilizado",
                "Já existe um item com esse número nessa categoria.", parent=self,
            )
            return
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        self.clear_item_form()
        self.refresh_items()
        self.status_message(message)

    def edit_selected_item(self) -> None:
        if self.selected_item_id is None:
            messagebox.showinfo("Editar item", "Selecione um item na lista ou na grade.", parent=self)
            return
        item = self.repository.get_item(self.selected_item_id)
        if item is None:
            self.refresh_items()
            return
        self.item_number_entry.delete(0, "end")
        self.item_number_entry.insert(0, str(item["number"]))
        self.item_name_entry.delete(0, "end")
        self.item_name_entry.insert(0, item["name"])
        self.item_category_var.set(item["category_name"])
        self.item_status_var.set(item["status"])
        self.item_save_button.configure(text="✓  Salvar alterações")
        self.item_cancel_button.grid(row=2, column=3, sticky="e", pady=(8, 0))
        self.item_number_entry.focus_set()

    def clear_item_form(self) -> None:
        self.selected_item_id = None
        self.item_number_entry.delete(0, "end")
        self.item_name_entry.delete(0, "end")
        self.item_status_var.set(DEFAULT_ITEM_STATUS)
        self.item_save_button.configure(text="＋  Adicionar item")
        self.item_cancel_button.grid_forget()
        self.item_number_entry.focus_set()
        self.refresh_items()

    def delete_selected_item(self) -> None:
        if self.selected_item_id is None:
            messagebox.showinfo("Excluir item", "Selecione um item na lista ou na grade.", parent=self)
            return
        item = self.repository.get_item(self.selected_item_id)
        if item is None:
            self.refresh_items()
            return
        if not messagebox.askyesno(
            "Confirmar exclusão",
            f'Deseja excluir o item Nº {item["number"]} - {item["name"]}?', parent=self,
        ):
            return
        try:
            self.repository.delete_item(self.selected_item_id)
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        self.clear_item_form()
        self.status_message("Item excluído.")

    @staticmethod
    def _sidebar_label(parent: tk.Widget, text: str) -> None:
        tk.Label(
            parent, text=text, bg=COLORS["sidebar"], fg="#8FA69D",
            font=("Segoe UI Semibold", 8), anchor="w",
        ).pack(fill="x", padx=24, pady=(0, 10))

    @staticmethod
    def _sidebar_button(
        parent: tk.Widget, text: str, command, active: bool = False
    ) -> tk.Button:
        button = tk.Button(
            parent, text=text, command=command, anchor="w", relief="flat", bd=0,
            padx=22, pady=11, cursor="hand2", font=("Segoe UI Semibold", 10),
            bg=COLORS["sidebar_hover"] if active else COLORS["sidebar"],
            fg="#FFFFFF" if active else "#C6D5CD",
            activebackground=COLORS["sidebar_hover"], activeforeground="#FFFFFF",
        )
        button.pack(fill="x", padx=10, pady=1)
        return button

    @staticmethod
    def _stat_card(parent: ttk.Frame, column: int, title: str, value: str) -> ttk.Label:
        card = ttk.Frame(parent, style="Surface.TFrame", padding=(17, 13))
        card.grid(row=0, column=column, sticky="ew", padx=(0, 10) if column < 2 else 0)
        ttk.Label(card, text=title, style="CardTitle.TLabel").pack(anchor="w")
        value_label = ttk.Label(card, text=value, style="CardValue.TLabel")
        value_label.pack(anchor="w", pady=(6, 0))
        return value_label

    def _show_categories(self) -> None:
        self._set_active_navigation("categories")
        self.content.tkraise()
        self.refresh_categories()
        self.name_entry.focus_set()

    def _show_dashboard(self) -> None:
        self._set_active_navigation("overview")
        self.dashboard_page.tkraise()
        self.refresh_dashboard()

    def _set_active_navigation(self, section: str) -> None:
        self.active_navigation = section
        navigation_buttons = {
            "overview": self.overview_nav_button, "categories": self.category_nav_button,
            "items": self.item_nav_button, "users": self.user_nav_button,
            "loans": self.loan_nav_button, "about": self.about_nav_button,
        }
        for button_section, button in navigation_buttons.items():
            active = section == button_section
            button.configure(
                bg=COLORS["sidebar_hover"] if active else COLORS["sidebar"],
                fg="#FFFFFF" if active else "#C6D5CD",
                activebackground=COLORS["sidebar_hover"],
            )

    def _apply_theme(self, theme_name: str) -> None:
        if theme_name not in THEMES:
            raise ValueError("Selecione um tema disponível.")
        try:
            self._save_theme(theme_name)
        except OSError as error:
            messagebox.showerror("Não foi possível salvar o tema", str(error), parent=self)
            return
        executable = [sys.executable]
        if not getattr(sys, "frozen", False):
            executable.append(str(Path(__file__).resolve()))
        working_directory = (
            APPLICATION_DIRECTORY if IS_FROZEN else Path(__file__).resolve().parent
        )
        restart_environment = os.environ.copy()
        if IS_FROZEN:
            restart_environment["PYINSTALLER_RESET_ENVIRONMENT"] = "1"
        try:
            subprocess.Popen(
                executable,
                cwd=str(working_directory),
                env=restart_environment,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except (OSError, ValueError) as error:
            messagebox.showerror(
                "Tema salvo, mas não foi possível reiniciar",
                f"O tema {theme_name} foi salvo, mas a aplicação continua aberta. "
                f"Feche e abra o sistema para aplicá-lo.\n\n{error}",
                parent=self,
            )
            return
        self.destroy()

    def refresh_categories(self) -> None:
        if not hasattr(self, "category_table"):
            return
        rows = self.repository.list_categories(self.search_entry.get())
        for item in self.category_table.get_children():
            self.category_table.delete(item)
        for row in rows:
            self.category_table.insert(
                "", "end", iid=str(row["id"]),
                values=(row["id"], row["name"], row["created_at"]),
            )
        self.total_value.configure(text=str(self.repository.count_categories()))
        self.filtered_value.configure(text=str(len(rows)))
        if hasattr(self, "item_category_combo"):
            self._refresh_item_category_choices()
            self.refresh_items()
        if hasattr(self, "loan_category_combo"):
            self._refresh_loan_category_choices()
        self.refresh_dashboard()

    def save_category(self) -> None:
        name = self.name_entry.get()
        try:
            if self.selected_category_id is None:
                self.repository.create_category(name)
                message = "Categoria cadastrada com sucesso."
            else:
                self.repository.update_category(self.selected_category_id, name)
                message = "Categoria atualizada com sucesso."
        except ValueError as error:
            messagebox.showwarning("Verifique os dados", str(error), parent=self)
            return
        except sqlite3.IntegrityError:
            messagebox.showwarning(
                "Categoria já existe", "Já existe uma categoria com esse nome.", parent=self
            )
            return
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        self.clear_form()
        self.refresh_categories()
        self.status_message(message)

    def edit_selected(self) -> None:
        selection = self.category_table.selection()
        if not selection:
            messagebox.showinfo("Editar categoria", "Selecione uma categoria na lista.", parent=self)
            return
        values = self.category_table.item(selection[0], "values")
        self.selected_category_id = int(values[0])
        self.name_entry.delete(0, "end")
        self.name_entry.insert(0, values[1])
        self.save_button.configure(text="✓  Salvar alterações")
        self.cancel_edit_button.grid(row=1, column=2, sticky="e", padx=(8, 0))
        self.name_entry.focus_set()

    def clear_form(self) -> None:
        self.selected_category_id = None
        self.name_entry.delete(0, "end")
        self.save_button.configure(text="＋  Adicionar categoria")
        self.cancel_edit_button.grid_forget()
        self.name_entry.focus_set()

    def delete_selected(self) -> None:
        selection = self.category_table.selection()
        if not selection:
            messagebox.showinfo("Excluir categoria", "Selecione uma categoria na lista.", parent=self)
            return
        category_id = int(selection[0])
        name = self.category_table.item(selection[0], "values")[1]
        if not messagebox.askyesno(
            "Confirmar exclusão", f'Deseja excluir a categoria "{name}"?', parent=self
        ):
            return
        try:
            self.repository.delete_category(category_id)
        except sqlite3.IntegrityError:
            messagebox.showwarning(
                "Categoria em uso",
                "Esta categoria possui itens cadastrados. Exclua ou transfira os itens antes de removê-la.",
                parent=self,
            )
            return
        except sqlite3.Error as error:
            messagebox.showerror("Erro no banco de dados", str(error), parent=self)
            return
        if self.selected_category_id == category_id:
            self.clear_form()
        self.refresh_categories()
        self.status_message("Categoria excluída.")

    def open_settings(self) -> None:
        dialog = tk.Toplevel(self)
        dialog.title("Configurações")
        dialog.configure(background=COLORS["background"])
        dialog.resizable(True, True)
        dialog.transient(self)
        dialog.grab_set()
        dialog.geometry("860x680")
        dialog.minsize(800, 620)
        frame = ttk.Frame(dialog, padding=22)
        frame.pack(fill="both", expand=True)
        settings_header = ttk.Frame(frame)
        settings_header.pack(fill="x", pady=(0, 14))
        ttk.Label(
            settings_header, text="Configurações", style="Title.TLabel",
        ).pack(side="left", anchor="w")
        ttk.Button(
            settings_header, text="Voltar", command=dialog.destroy,
        ).pack(side="right", ipadx=14, ipady=5)
        notebook = ttk.Notebook(frame)
        notebook.pack(fill="both", expand=True)
        database_tab = ttk.Frame(notebook, padding=20)
        backup_tab = ttk.Frame(notebook, padding=20)
        email_tab = ttk.Frame(notebook, padding=20)
        colors_tab = ttk.Frame(notebook, padding=20)
        notebook.add(database_tab, text="Banco de dados")
        notebook.add(backup_tab, text="Backup")
        notebook.add(email_tab, text="E-mail")
        notebook.add(colors_tab, text="Cores")

        ttk.Label(database_tab, text="Local do banco SQLite", style="CardTitle.TLabel").pack(anchor="w")
        ttk.Label(
            database_tab,
            text=(
                "O SQLite fica no computador host. Para acesso de outros computadores, "
                "use uma sessão remota no host; não abra o arquivo diretamente por pasta de rede."
            ),
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(4, 18))
        path_row = ttk.Frame(database_tab)
        path_row.pack(fill="x")
        path_entry = PlaceholderEntry(
            path_row, font=("Segoe UI", 9), placeholder="C:\\Dados\\emprestimo.db",
        )
        path_entry.delete(0, "end")
        path_entry.insert(0, str(self.database_path))
        path_entry.pack(side="left", fill="x", expand=True, ipady=5)

        def browse() -> None:
            current_path = path_entry.get() or str(self.database_path)
            selected = filedialog.asksaveasfilename(
                parent=dialog, title="Escolha onde salvar o banco SQLite",
                initialdir=str(Path(current_path).parent),
                initialfile=Path(current_path).name,
                defaultextension=".db",
                filetypes=[("Banco SQLite", "*.db *.sqlite *.sqlite3"), ("Todos os arquivos", "*.*")],
            )
            if selected:
                path_entry.delete(0, "end")
                path_entry.insert(0, selected)

        ttk.Button(path_row, text="Procurar...", command=browse).pack(side="left", padx=(8, 0))

        def apply_settings() -> None:
            path = Path(path_entry.get().strip()).expanduser()
            if not path.name:
                messagebox.showwarning("Caminho inválido", "Informe o arquivo do banco.", parent=dialog)
                return
            try:
                repository = CategoryRepository(path)
                self._save_database_path(repository.database_path)
            except (OSError, sqlite3.Error, ValueError) as error:
                messagebox.showerror("Não foi possível abrir o banco", str(error), parent=dialog)
                return
            self.database_path = repository.database_path
            self.repository = repository
            self.refresh_categories()
            self.refresh_users()
            self.refresh_loans()
            self.status_message("Banco de dados atualizado.")

        ttk.Button(
            database_tab, text="Aplicar caminho", style="Accent.TButton", command=apply_settings,
        ).pack(anchor="e", pady=(18, 0))

        backup_tab.columnconfigure(1, weight=1)
        ttk.Label(
            backup_tab, text="Cópia completa e rotina automática",
            style="CardTitle.TLabel",
        ).grid(row=0, column=0, columnspan=3, sticky="w")
        ttk.Label(
            backup_tab,
            text=(
                "Será criada uma subpasta chamada backup dentro do local escolhido. "
                "A rotina diária/semanal roda enquanto o sistema estiver aberto; se estiver "
                "fechado no dia previsto, faz a cópia vencida na próxima abertura."
            ),
            style="Subtitle.TLabel", wraplength=700, justify="left",
        ).grid(row=1, column=0, columnspan=3, sticky="w", pady=(5, 12))
        ttk.Label(
            backup_tab, text="Local onde criar a pasta backup", style="CardTitle.TLabel",
        ).grid(row=2, column=0, sticky="w", padx=(0, 12), pady=8)
        backup_path_entry = PlaceholderEntry(
            backup_tab, font=("Segoe UI", 9),
            placeholder="Selecione o local onde será criada a pasta backup",
        )
        configured_root = self.app_config.get("backup_directory")
        default_backup_root = (
            Path(str(configured_root)).expanduser().resolve()
            if configured_root is not None
            else self.database_path.parent.resolve()
        )
        try:
            self._ensure_backup_directory()
        except OSError as error:
            messagebox.showwarning(
                "Pasta de backup",
                f"Não foi possível criar a pasta padrão de backup:\n"
                f"{self._backup_directory()}\n\n{error}",
                parent=dialog,
            )
        backup_path_entry.delete(0, "end")
        backup_path_entry.insert(0, str(default_backup_root))
        backup_path_entry.grid(row=2, column=1, sticky="ew", pady=8, ipady=4)
        backup_directory_label = ttk.Label(
            backup_tab,
            text=f"Os arquivos serão salvos em: {self._backup_directory()}",
            style="Subtitle.TLabel", wraplength=700, justify="left",
        )
        backup_directory_label.grid(
            row=3, column=0, columnspan=3, sticky="w", pady=(0, 12),
        )

        def browse_backup_directory() -> None:
            selected = filedialog.askdirectory(
                parent=dialog, title="Escolher local para a pasta backup",
                initialdir=backup_path_entry.get() or str(self.database_path.parent),
            )
            if selected:
                backup_path_entry.delete(0, "end")
                backup_path_entry.insert(0, selected)
                selected_root = Path(selected).resolve()
                preview_directory = (
                    selected_root if selected_root.name.casefold() == "backup"
                    else selected_root / "backup"
                )
                backup_directory_label.configure(
                    text=f"Os arquivos serão salvos em: {preview_directory}",
                )

        ttk.Button(
            backup_tab, text="Procurar...", command=browse_backup_directory,
        ).grid(row=2, column=2, sticky="e", padx=(8, 0), pady=8)
        ttk.Label(
            backup_tab, text="Frequência", style="CardTitle.TLabel",
        ).grid(row=4, column=0, sticky="w", padx=(0, 12), pady=8)
        backup_frequency_var = tk.StringVar(
            value=str(self.app_config.get("backup_frequency", "Desativado")),
        )
        backup_frequency_combo = ttk.Combobox(
            backup_tab, textvariable=backup_frequency_var,
            values=("Desativado", "Diário", "Semanal"), state="readonly", width=20,
        )
        backup_frequency_combo.grid(row=4, column=1, sticky="w", pady=8)
        last_backup = str(self.app_config.get("backup_last_success", ""))
        ttk.Label(
            backup_tab,
            text=f"Último backup: {last_backup or 'ainda não realizado'}",
            style="Subtitle.TLabel",
        ).grid(row=5, column=0, columnspan=3, sticky="w", pady=(8, 18))

        def save_backup_settings() -> None:
            directory_text = backup_path_entry.get().strip()
            if not directory_text:
                messagebox.showwarning(
                    "Pasta necessária", "Selecione onde os backups serão armazenados.",
                    parent=dialog,
                )
                return
            directory = Path(directory_text).expanduser()
            if directory.exists() and not directory.is_dir():
                messagebox.showwarning(
                    "Pasta inválida", "O caminho informado não é uma pasta.", parent=dialog,
                )
                return
            if backup_frequency_var.get() not in ("Desativado", "Diário", "Semanal"):
                messagebox.showwarning(
                    "Frequência inválida", "Selecione uma frequência disponível.", parent=dialog,
                )
                return
            previous_values = {
                key: self.app_config.get(key)
                for key in (
                    "backup_directory", "backup_frequency", "backup_last_attempt",
                )
            }
            backup_root = directory.resolve()
            backup_target = (
                backup_root if backup_root.name.casefold() == "backup"
                else backup_root / "backup"
            )
            if (
                self._backup_directory() != backup_target
                or self.app_config.get("backup_frequency", "Desativado")
                != backup_frequency_var.get()
            ):
                self.app_config.pop("backup_last_attempt", None)
            self.app_config["backup_directory"] = str(backup_root)
            self.app_config["backup_frequency"] = backup_frequency_var.get()
            try:
                backup_target.mkdir(parents=True, exist_ok=True)
                self._save_app_config()
            except OSError as error:
                for key, value in previous_values.items():
                    if value is None:
                        self.app_config.pop(key, None)
                    else:
                        self.app_config[key] = value
                messagebox.showerror(
                    "Não foi possível salvar a rotina de backup", str(error), parent=dialog,
                )
                return
            self.status_message("Configuração de backup salva.")
            backup_directory_label.configure(
                text=f"Os arquivos serão salvos em: {backup_target}",
            )
            messagebox.showinfo(
                "Configuração salva",
                f"A rotina de backup foi atualizada.\n\nPasta criada:\n{backup_target}\n\n"
                "O próximo backup será feito ao iniciar o sistema, caso já esteja vencido.",
                parent=dialog,
            )

        backup_actions = ttk.Frame(backup_tab)
        backup_actions.grid(row=6, column=0, columnspan=3, sticky="e", pady=(16, 0))
        ttk.Button(
            backup_actions, text="Importar banco existente...",
            command=self._import_database_backup,
        ).pack(side="right", padx=(8, 0))
        ttk.Button(
            backup_actions, text="Fazer backup agora", command=self._run_database_backup,
        ).pack(side="right", padx=(8, 0))
        ttk.Button(
            backup_actions, text="Salvar rotina", style="Accent.TButton",
            command=save_backup_settings,
        ).pack(side="right")
        ttk.Label(
            backup_tab,
            text=(
                "Importar banco substitui os dados que estão abertos no sistema. O arquivo "
                "selecionado é validado, mas o banco atual não é copiado automaticamente. "
                "Faça um backup manual antes se quiser guardar a versão atual."
            ),
            style="Subtitle.TLabel", wraplength=700, justify="left",
        ).grid(row=7, column=0, columnspan=3, sticky="w", pady=(22, 0))

        ttk.Label(
            colors_tab, text="Escolha uma paleta para a interface",
            style="CardTitle.TLabel",
        ).pack(anchor="w")
        ttk.Label(
            colors_tab, text="A alteração é aplicada imediatamente e fica salva neste computador.",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(5, 16))
        saved_theme = str(self.app_config.get("theme", "Verde"))
        theme_var = tk.StringVar(value=saved_theme if saved_theme in THEMES else "Verde")
        theme_combo = ttk.Combobox(
            colors_tab, textvariable=theme_var, values=tuple(THEMES),
            state="readonly", width=24,
        )
        theme_combo.pack(anchor="w", ipady=4)
        ttk.Label(
            colors_tab, text=f"Temas disponíveis: {', '.join(THEMES)}.",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(12, 0))
        ttk.Button(
            colors_tab, text="Aplicar tema", style="Accent.TButton",
            command=lambda: self._apply_theme(theme_var.get()),
        ).pack(anchor="w", pady=(22, 0))

        try:
            email_settings = self.repository.get_email_settings()
        except (OSError, RuntimeError, sqlite3.Error) as error:
            email_settings = {}
            messagebox.showwarning(
                "Credencial indisponível",
                f"Informe novamente a senha SMTP para compartilhar a configuração no host.\n\n{error}",
                parent=dialog,
            )

        email_tab.columnconfigure(1, weight=1)
        email_entries: dict[str, PlaceholderEntry] = {}

        def add_email_field(
            key: str, label: str, row: int, placeholder: str,
            show: str | None = None, width: int | None = None,
        ) -> PlaceholderEntry:
            ttk.Label(email_tab, text=label, style="CardTitle.TLabel").grid(
                row=row, column=0, sticky="w", padx=(0, 16), pady=7,
            )
            initial_value = str(email_settings.get(key, ""))
            entry = PlaceholderEntry(
                email_tab, show=show or "", width=width,
                placeholder=placeholder if not initial_value else "",
            )
            if initial_value:
                entry.insert(0, initial_value)
            email_entries[key] = entry
            entry.grid(row=row, column=1, sticky="ew", pady=7, ipady=4)
            return entry

        add_email_field("institution", "Instituição", 0, "Ex.: Secretaria Municipal")
        add_email_field("email", "E-mail de envio", 1, "remetente@instituicao.gov.br")
        add_email_field("app_password", "Senha de app", 2, "Senha de app SMTP", show="*")
        add_email_field("sender_name", "Nome do remetente", 3, "Ex.: Setor de Patrimônio")
        add_email_field("smtp_host", "Servidor SMTP", 4, "smtp.gmail.com")
        add_email_field("smtp_port", "Porta SMTP", 5, "587", width=10)
        tls_var = tk.BooleanVar(value=bool(email_settings.get("use_tls", True)))
        ttk.Checkbutton(email_tab, text="Usar conexão segura TLS", variable=tls_var).grid(
            row=6, column=1, sticky="w", pady=(8, 4),
        )
        ttk.Label(email_tab, text="Assinatura padrão", style="CardTitle.TLabel").grid(
            row=7, column=0, sticky="nw", padx=(0, 16), pady=7,
        )
        signature_text = tk.Text(
            email_tab, height=4, wrap="word", font=("Segoe UI", 9),
            relief="solid", bd=1, padx=7, pady=6,
        )
        signature_text.grid(row=7, column=1, sticky="ew", pady=7)
        signature_text.insert(
            "1.0", str(email_settings.get("signature", DEFAULT_EMAIL_SIGNATURE)),
        )
        ttk.Label(
            email_tab,
            text="A senha de envio fica armazenada no banco do host para ser usada pela aplicação.",
            style="Subtitle.TLabel",
        ).grid(row=8, column=0, columnspan=2, sticky="w", pady=(8, 0))

        def save_email_settings() -> None:
            try:
                port_text = email_entries["smtp_port"].get().strip()
                port = int(port_text or "587")
                self.repository.save_email_settings(
                    email_entries["email"].get(),
                    email_entries["app_password"].get(),
                    email_entries["institution"].get(),
                    email_entries["sender_name"].get(),
                    email_entries["smtp_host"].get(),
                    port,
                    tls_var.get(),
                    signature_text.get("1.0", "end-1c"),
                )
            except ValueError as error:
                messagebox.showwarning("Verifique os dados", str(error), parent=dialog)
                return
            except (OSError, RuntimeError, sqlite3.Error) as error:
                messagebox.showerror("Falha ao salvar credenciais", str(error), parent=dialog)
                return
            self.status_message("Configurações de e-mail salvas.")
            messagebox.showinfo("Configurações salvas", "As credenciais de e-mail foram salvas.", parent=dialog)

        ttk.Button(
            email_tab, text="Salvar credenciais", style="Accent.TButton",
            command=save_email_settings,
        ).grid(row=9, column=1, sticky="e", pady=(16, 0))
        self.update_idletasks()
        dialog.update_idletasks()
        x = self.winfo_rootx() + (self.winfo_width() - dialog.winfo_width()) // 2
        y = self.winfo_rooty() + (self.winfo_height() - dialog.winfo_height()) // 2
        dialog.geometry(f"+{max(0, x)}+{max(0, y)}")
        dialog.lift()

    def export_excel(self) -> None:
        rows = self.repository.list_categories(self.search_entry.get())
        if not rows:
            messagebox.showinfo("Relatório Excel", "Não há categorias para exportar.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Salvar relatório Excel", defaultextension=".xlsx",
            filetypes=[("Planilha Excel", "*.xlsx")], initialfile="relatorio_categorias.xlsx",
        )
        if not path:
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Categorias"
        sheet.append(["ID", "Categoria", "Cadastrada em"])
        for row in rows:
            sheet.append([row["id"], row["name"], row["created_at"]])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183B3A")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 24
        sheet.column_dimensions["A"].width = 10
        sheet.column_dimensions["B"].width = 38
        sheet.column_dimensions["C"].width = 24
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        try:
            workbook.save(path)
        except OSError as error:
            messagebox.showerror("Falha ao exportar", str(error), parent=self)
            return
        self.status_message(f"Planilha salva em {path}")

    def export_pdf(self) -> None:
        rows = self.repository.list_categories(self.search_entry.get())
        if not rows:
            messagebox.showinfo("Relatório PDF", "Não há categorias para exportar.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Salvar relatório PDF", defaultextension=".pdf",
            filetypes=[("Documento PDF", "*.pdf")], initialfile="relatorio_categorias.pdf",
        )
        if not path:
            return
        styles = getSampleStyleSheet()
        styles["Title"].textColor = colors.HexColor(COLORS["sidebar"])
        styles["Title"].fontName = "Helvetica-Bold"
        data = [["ID", "Categoria", "Cadastrada em"]]
        data.extend([
            [
                str(row["id"]), Paragraph(escape(row["name"]), styles["BodyText"]), row["created_at"],
            ]
            for row in rows
        ])
        table = Table(data, colWidths=[18 * mm, 92 * mm, 48 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COLORS["sidebar"])),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F0")]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(COLORS["border"])),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        document = SimpleDocTemplate(
            path, pagesize=landscape(A4), rightMargin=20 * mm, leftMargin=20 * mm,
            topMargin=20 * mm, bottomMargin=20 * mm,
        )
        story = [
            Paragraph("Relatório de categorias", styles["Title"]),
            Paragraph(
                f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} · "
                f"{len(rows)} categoria(s)", styles["Normal"],
            ),
            Spacer(1, 8 * mm),
            table,
        ]
        try:
            document.build(story)
        except (OSError, ValueError) as error:
            messagebox.showerror("Falha ao exportar", str(error), parent=self)
            return
        self.status_message(f"PDF salvo em {path}")

    def export_items_excel(self) -> None:
        rows = self.repository.list_items(
            self.item_search_entry.get(), self._selected_item_category_id(),
        )
        if not rows:
            messagebox.showinfo("Relatório Excel", "Não há itens para exportar.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Salvar relatório de itens", defaultextension=".xlsx",
            filetypes=[("Planilha Excel", "*.xlsx")], initialfile="relatorio_itens.xlsx",
        )
        if not path:
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Itens"
        sheet.append(["Número", "Item", "Categoria", "Status", "Cadastrado em"])
        for row in rows:
            sheet.append([
                row["number"], row["name"], row["category_name"], row["status"], row["created_at"],
            ])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183B3A")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 24
        for column, width in {"A": 12, "B": 36, "C": 28, "D": 20, "E": 24}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        try:
            workbook.save(path)
        except OSError as error:
            messagebox.showerror("Falha ao exportar", str(error), parent=self)
            return
        self.status_message(f"Planilha de itens salva em {path}")

    def export_items_pdf(self) -> None:
        rows = self.repository.list_items(
            self.item_search_entry.get(), self._selected_item_category_id(),
        )
        if not rows:
            messagebox.showinfo("Relatório PDF", "Não há itens para exportar.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Salvar relatório de itens", defaultextension=".pdf",
            filetypes=[("Documento PDF", "*.pdf")], initialfile="relatorio_itens.pdf",
        )
        if not path:
            return
        styles = getSampleStyleSheet()
        styles["Title"].textColor = colors.HexColor(COLORS["sidebar"])
        styles["Title"].fontName = "Helvetica-Bold"
        body_style = styles["BodyText"]
        body_style.wordWrap = "CJK"
        data = [["Número", "Item", "Categoria", "Status", "Cadastrado em"]]
        data.extend([
            [
                str(row["number"]),
                Paragraph(escape(row["name"]), body_style),
                Paragraph(escape(row["category_name"]), body_style),
            row["status"],
                row["created_at"],
            ]
            for row in rows
        ])
        table = Table(
            data, colWidths=[20 * mm, 86 * mm, 58 * mm, 32 * mm, 40 * mm],
            repeatRows=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COLORS["sidebar"])),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F0")]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(COLORS["border"])),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        document = SimpleDocTemplate(
            path, pagesize=landscape(A4), rightMargin=20 * mm, leftMargin=20 * mm,
            topMargin=20 * mm, bottomMargin=20 * mm,
        )
        story = [
            Paragraph("Relatório de itens", styles["Title"]),
            Paragraph(
                f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} · "
                f"{len(rows)} item(ns)", styles["Normal"],
            ),
            Spacer(1, 8 * mm),
            table,
        ]
        try:
            document.build(story)
        except (OSError, ValueError) as error:
            messagebox.showerror("Falha ao exportar", str(error), parent=self)
            return
        self.status_message(f"PDF de itens salvo em {path}")

    def status_message(self, text: str) -> None:
        self.title(f"Emprestimo | Administração | {text}")
        self.after(3500, lambda: self.title("Emprestimo | Administração"))


def main() -> None:
    app = LoanAdminApp()
    app.mainloop()


if __name__ == "__main__":
    main()
